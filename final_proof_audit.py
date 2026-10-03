import openpyxl

file_paths = [
    r"C:/Users/LENOVO/Downloads/test/HDFCBANK_Valuation_Model (1).xlsx",
    r"c:/Users/LENOVO/Downloads/Advance Financial Project/HDFCBANK_Valuation_Model.xlsx"
]

for p in file_paths:
    print(f"\n================ AUDITING {p} ================")
    wb = openpyxl.load_workbook(p, data_only=False)
    wb_v = openpyxl.load_workbook(p, data_only=True)

    errors = []
    verdicts = []
    for sname in wb.sheetnames:
        ws = wb[sname]
        ws_v = wb_v[sname]
        for r in range(1, min(ws.max_row+1, 100)):
            for c in range(1, min(ws.max_column+1, 35)):
                f = str(ws.cell(r, c).value or '')
                v = str(ws_v.cell(r, c).value or '')
                col_l = openpyxl.utils.get_column_letter(c)
                
                # Check Excel errors
                if any(err in v for err in ['#VALUE!', '#REF!', '#DIV/0!', '#NAME?']) or v == '#N/A' or v == '#':
                    errors.append((sname, f"{col_l}{r}", f, v))
                
                # Check Investment verdicts
                if any(kw == v for kw in ['Undervalued', 'Overvalued', 'BUY', 'SELL', 'HOLD', 'Strong Buy', 'Strong Sell']):
                    verdicts.append((sname, f"{col_l}{r}", f, v))

    print(f"Total Excel Errors found: {len(errors)}")
    for e in errors:
        print(f"  Error: {e[0]}!{e[1]} form={e[2]} val={e[3]}")
    
    print(f"Total Automatic Verdicts found: {len(verdicts)}")
    for vd in verdicts:
        print(f"  Verdict: {vd[0]}!{vd[1]} form={vd[2]} val={vd[3]}")

    # Check Key Metrics
    print("\n--- Key Metric Reconciliations ---")
    print("Price (DS B8):", wb_v['Data Sheet']['B8'].value)
    print("Shares (DS B6):", wb_v['Data Sheet']['B6'].value)
    print("Market Cap (DS B9):", wb_v['Data Sheet']['B9'].value)
    print("Beta (WACC K28):", wb_v['WACC']['K28'].value)
    print("Ke (WACC K29):", wb_v['WACC']['K29'].value)
    print("AI Summary E5 (Discount Rate):", wb_v['AI Valuation Summary']['E5'].value)
    print("AI Summary B10 (Pillar 3 narrative):", wb_v['AI Valuation Summary']['B10'].value)
    print("Residual Income D15 (Ke):", wb_v['AI Valuation Summary']['D15'].value)
    print("Residual Income C15 (DuPont ROE):", wb_v['AI Valuation Summary']['C15'].value)
    print("Comp Valuation Q39:", wb_v['Comp_Valuation']['Q39'].value)
    print("Intrinsic Valuation H16:H18:", wb_v['Intrinsic Valuation']['H16'].value, wb_v['Intrinsic Valuation']['H17'].value, wb_v['Intrinsic Valuation']['H18'].value)
