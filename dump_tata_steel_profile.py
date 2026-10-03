import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\LENOVO\Downloads\Advance Financial Project\Tata Steel Ltd.xlsx', data_only=False)
wb_v = openpyxl.load_workbook(r'c:\Users\LENOVO\Downloads\Advance Financial Project\Tata Steel Ltd.xlsx', data_only=True)

with open(r'c:\Users\LENOVO\Downloads\Advance Financial Project\tata_steel_profile.txt', 'w', encoding='utf-8') as out:
    out.write(f"SHEETS: {wb.sheetnames}\n\n")

    # AI Valuation Summary
    out.write("=== AI VALUATION SUMMARY ===\n")
    ws = wb['AI Valuation Summary']
    ws_v = wb_v['AI Valuation Summary']
    for r in range(1, 35):
        row_str = []
        for c in range(1, 8):
            cl = openpyxl.utils.get_column_letter(c)
            f = ws.cell(r, c).value
            v = ws_v.cell(r, c).value
            if f is not None:
                row_str.append(f"{cl}{r}: form=[{f}] val=[{v}]")
        if row_str:
            out.write(" | ".join(row_str) + "\n")

    # Data Sheet
    out.write("\n=== DATA SHEET KEY CELLS ===\n")
    ws_ds = wb['Data Sheet']
    ws_ds_v = wb_v['Data Sheet']
    for cell in ['B1', 'B6', 'B8', 'B9', 'B16', 'B17', 'B23', 'B28', 'B29', 'B30', 'B32', 'B61', 'K17', 'K30', 'K57', 'K59', 'K69', 'K70', 'K72']:
        out.write(f"{cell}: form=[{ws_ds[cell].value}] val=[{ws_ds_v[cell].value}]\n")

    # WACC
    out.write("\n=== WACC KEY CELLS ===\n")
    ws_w = wb['WACC']
    ws_w_v = wb_v['WACC']
    for cell in ['B2', 'C34', 'C35', 'E26', 'E27', 'E28', 'E34', 'E35', 'E38', 'K26', 'K27', 'K28', 'K29', 'K33', 'K34', 'K35', 'K36', 'K40', 'K41', 'K43', 'K44', 'K46']:
        out.write(f"{cell}: form=[{ws_w[cell].value}] val=[{ws_w_v[cell].value}]\n")

    # DCF
    out.write("\n=== DCF KEY CELLS ===\n")
    ws_d = wb['DCF']
    ws_d_v = wb_v['DCF']
    for cell in ['B3', 'B4', 'D18', 'D19', 'D20', 'D21', 'D33', 'D34', 'D35', 'D37', 'D38', 'D39', 'D40', 'D42', 'D44', 'D45']:
        out.write(f"{cell}: form=[{ws_d[cell].value}] val=[{ws_d_v[cell].value}]\n")

    # Comp_Valuation
    out.write("\n=== COMP_VALUATION KEY CELLS ===\n")
    ws_c = wb['Comp_Valuation']
    ws_c_v = wb_v['Comp_Valuation']
    for r in range(10, 42):
        row_str = []
        for col_l in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'O', 'P', 'Q']:
            f = ws_c[f'{col_l}{r}'].value
            v = ws_c_v[f'{col_l}{r}'].value
            if f is not None:
                row_str.append(f"{col_l}{r}: form=[{f}] val=[{v}]")
        if row_str:
            out.write(" | ".join(row_str) + "\n")

    # Intrinsic Valuation
    out.write("\n=== INTRINSIC VALUATION KEY CELLS ===\n")
    ws_iv = wb['Intrinsic Valuation']
    ws_iv_v = wb_v['Intrinsic Valuation']
    for cell in ['A5', 'A21', 'B3', 'B4', 'B67', 'L67', 'B68', 'L68', 'B69', 'L69', 'B70', 'L70', 'B71', 'L71', 'B72', 'L72', 'B73', 'L73']:
        out.write(f"{cell}: form=[{ws_iv[cell].value}] val=[{ws_iv_v[cell].value}]\n")

print("Profile written successfully!")
