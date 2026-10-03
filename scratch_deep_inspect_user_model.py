import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\LENOVO\Downloads\test\HDFCBANK_Valuation_Model (1).xlsx'
wb_f = openpyxl.load_workbook(path, data_only=False)
wb_v = openpyxl.load_workbook(path, data_only=True)

print("=== DEEP INSPECTION OF HDFCBANK_Valuation_Model (1).xlsx ===")

print("\n1. WACC Sheet:")
if 'WACC' in wb_f.sheetnames:
    ws_f = wb_f['WACC']
    ws_v = wb_v['WACC']
    for cell in ['C34', 'C35', 'E34', 'E35', 'E38', 'K26', 'K27', 'K28', 'K29', 'K33', 'K36', 'K40', 'K41', 'K44', 'K46']:
        print(f"  {cell}: val={ws_v[cell].value}, formula={ws_f[cell].value}")

print("\n2. AI Valuation Summary:")
if 'AI Valuation Summary' in wb_f.sheetnames:
    ws_f = wb_f['AI Valuation Summary']
    ws_v = wb_v['AI Valuation Summary']
    for cell in ['A5', 'B5', 'C5', 'D5', 'E5', 'F5', 'G5']:
        print(f"  {cell}: val={ws_v[cell].value}, formula={ws_f[cell].value}")
    for r in range(8, 12):
        print(f"  Row {r}: A='{ws_v.cell(r,1).value}', B='{ws_v.cell(r,2).value}'")
    print("  Excess Return Table:")
    for r in range(15, 20):
        print(f"    Row {r}: B={ws_v.cell(r,2).value} (f={ws_f.cell(r,2).value}), C={ws_v.cell(r,3).value} (f={ws_f.cell(r,3).value}), D={ws_v.cell(r,4).value} (f={ws_f.cell(r,4).value})")
    print("  Bridge:")
    for r in [22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32]:
        print(f"    Row {r}: {ws_v.cell(r,1).value} = {ws_v.cell(r,2).value} (f={ws_f.cell(r,2).value})")

print("\n3. Data Sheet:")
if 'Data Sheet' in wb_f.sheetnames:
    ws_f = wb_f['Data Sheet']
    ws_v = wb_v['Data Sheet']
    for cell in ['B6', 'B8', 'B9', 'K57', 'K58', 'K70']:
        print(f"  {cell}: val={ws_v[cell].value}, formula={ws_f[cell].value}")

print("\n4. Intrinsic Valuation:")
if 'Intrinsic Valuation' in wb_f.sheetnames:
    ws_f = wb_f['Intrinsic Valuation']
    ws_v = wb_v['Intrinsic Valuation']
    print(f"  A21: val={ws_v['A21'].value}, formula={ws_f['A21'].value}")
    print(f"  H16: val={ws_v['H16'].value}, formula={ws_f['H16'].value}")
    print(f"  H17: val={ws_v['H17'].value}, formula={ws_f['H17'].value}")
    print(f"  H18: val={ws_v['H18'].value}, formula={ws_f['H18'].value}")
    print(f"  H37: val={ws_v['H37'].value}, formula={ws_f['H37'].value}")

print("\n5. Common Size & Historical FS:")
if 'Common Size Statement' in wb_f.sheetnames:
    print(f"  Common Size L21: val={wb_v['Common Size Statement']['L21'].value}, formula={wb_f['Common Size Statement']['L21'].value}")
if 'Historical FS' in wb_f.sheetnames:
    print(f"  Historical FS A5: val={wb_v['Historical FS']['A5'].value}, formula={wb_f['Historical FS']['A5'].value}")

print("\n6. Comp_Valuation:")
if 'Comp_Valuation' in wb_f.sheetnames:
    ws_f = wb_f['Comp_Valuation']
    ws_v = wb_v['Comp_Valuation']
    print(f"  Q39: val={ws_v['Q39'].value}, formula={ws_f['Q39'].value}")
