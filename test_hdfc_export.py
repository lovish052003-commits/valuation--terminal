import os
import openpyxl
from excel_exporter import export_valuation_model
from screener_client import fetch_complete_financial_data

print("Testing export for HDFCBANK using current engine...")
data = fetch_complete_financial_data("HDFCBANK")
print(f"Fetched data for {data.get('company_name', 'HDFCBANK')}, price: {data.get('current_price')}")

out_path = export_valuation_model("HDFCBANK", data)
print(f"Generated: {out_path}")

wb = openpyxl.load_workbook(out_path, data_only=False)
ws_dcf = wb['DCF']
print("\n--- DCF SHEET IN GENERATED HDFCBANK ---")
for r in [8, 12, 16, 20, 33, 34, 35, 37, 38, 39, 40, 42, 44, 45]:
    cell_b = ws_dcf.cell(r, 2).value
    cell_d = ws_dcf.cell(r, 4).value
    cell_i = ws_dcf.cell(r, 9).value
    print(f"Row {r:2d} | B: {str(cell_b):<30} | D: {str(cell_d):<25} | I: {str(cell_i)}")

wb_data = openpyxl.load_workbook(out_path, data_only=True)
ws_dcf_val = wb_data['DCF']
print("\n--- DCF EVALUATED VALUES ---")
for r in [20, 33, 34, 35, 37, 38, 39, 40, 42, 44, 45]:
    val_b = ws_dcf_val.cell(r, 2).value
    val_d = ws_dcf_val.cell(r, 4).value
    print(f"Row {r:2d} | B: {str(val_b):<30} | D: {str(val_d)}")
