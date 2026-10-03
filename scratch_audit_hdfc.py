import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

file_path = r'exports\HDFCBANK_Valuation_Model.xlsx'
if not os.path.exists(file_path):
    file_path = r'C:\Users\LENOVO\Downloads\test\HDFCBANK_Valuation_Model.xlsx'

wb_f = openpyxl.load_workbook(file_path, data_only=False)
wb_v = openpyxl.load_workbook(file_path, data_only=True)

print("=== AUDITING EXISTING HDFCBANK WORKBOOK ===")
print("File:", file_path)

# 1. Shares & Price in Data Sheet
ds_f = wb_f['Data Sheet']
ds_v = wb_v['Data Sheet']
print(f"Data Sheet B8 (CMP): f={ds_f['B8'].value}, v={ds_v['B8'].value}")
print(f"Data Sheet B6 (Shares): f={ds_f['B6'].value}, v={ds_v['B6'].value}")
print(f"Data Sheet B9 (MktCap): f={ds_f['B9'].value}, v={ds_v['B9'].value}")
print(f"Data Sheet K70 (Raw shares): f={ds_f['K70'].value}, v={ds_v['K70'].value}")
k70_v = float(ds_v['K70'].value or 0)
print(f"Data Sheet K70/1e7: {k70_v/1e7} Cr")

# 2. WACC vs Residual Income Cost of Equity
if 'WACC' in wb_f.sheetnames:
    ws_wacc_f = wb_f['WACC']
    ws_wacc_v = wb_v['WACC']
    print(f"WACC K29 (Ke): f={ws_wacc_f['K29'].value}, v={ws_wacc_v['K29'].value}")
    print(f"WACC K46 (WACC): f={ws_wacc_f['K46'].value}, v={ws_wacc_v['K46'].value}")
    print(f"WACC peer betas (J14:J18): {[ws_wacc_v[f'J{r}'].value for r in range(14, 19)]}")

# 3. AI Valuation Summary
if 'AI Valuation Summary' in wb_f.sheetnames:
    ws_ai_f = wb_f['AI Valuation Summary']
    ws_ai_v = wb_v['AI Valuation Summary']
    print(f"AI Summary Headline: A5={ws_ai_v['A5'].value}, B5={ws_ai_v['B5'].value}, C5={ws_ai_v['C5'].value}, D5={ws_ai_v['D5'].value}, E5={ws_ai_v['E5'].value}")
    print(f"AI Summary B10 (Pillar 3): {ws_ai_v['B10'].value}")
    print(f"AI Summary B11 (Pillar 4): {ws_ai_v['B11'].value}")
    print("AI Summary Excess Return Table (C15:D19):")
    for r in range(15, 20):
        print(f"  Row {r}: B={ws_ai_v[f'B{r}'].value}, C={ws_ai_v[f'C{r}'].value} (f={ws_ai_f[f'C{r}'].value}), D={ws_ai_v[f'D{r}'].value} (f={ws_ai_f[f'D{r}'].value})")
    print(f"AI Summary B22: {ws_ai_v['B22'].value}, B29: {ws_ai_v['B29'].value}, B30: {ws_ai_v['B30'].value}, B32: {ws_ai_v['B32'].value}")

# 4. Intrinsic Valuation
if 'Intrinsic Valuation' in wb_f.sheetnames:
    ws_iv_f = wb_f['Intrinsic Valuation']
    ws_iv_v = wb_v['Intrinsic Valuation']
    print(f"Intrinsic Valuation A21: f={ws_iv_f['A21'].value}, v={ws_iv_v['A21'].value}")
    print(f"Intrinsic Valuation H16: f={ws_iv_f['H16'].value}, H17: f={ws_iv_f['H17'].value}, H18: f={ws_iv_f['H18'].value}")
    print(f"Intrinsic Valuation H38: f={ws_iv_f['H38'].value}, v={ws_iv_v['H38'].value}")

# 5. Historical FS & Common Size
if 'Historical FS' in wb_f.sheetnames:
    print(f"Historical FS A5: f={wb_f['Historical FS']['A5'].value}")
if 'Common Size Statement' in wb_f.sheetnames:
    print(f"Common Size Statement L21: f={wb_f['Common Size Statement']['L21'].value}, v={wb_v['Common Size Statement']['L21'].value}")

# 6. Comp_Valuation
if 'Comp_Valuation' in wb_f.sheetnames:
    ws_cv_f = wb_f['Comp_Valuation']
    ws_cv_v = wb_v['Comp_Valuation']
    print("Comp_Valuation Row 12 (Peer 1):")
    for col in ['B', 'C', 'D', 'E', 'F', 'H', 'I', 'J', 'O', 'P', 'Q']:
        print(f"  Col {col}: f={ws_cv_f[f'{col}12'].value}, v={ws_cv_v[f'{col}12'].value}")
