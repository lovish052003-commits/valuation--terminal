import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\test\ADANIENT_Valuation_Model (8).xlsx', data_only=False)
ws = wb['Comp_Valuation']
for r in range(10, 42):
    row_vals = [f'{c}{r}: {ws[f"{c}{r}"].value}' for c in ['B','C','I','J','O','P','Q'] if ws[f'{c}{r}'].value is not None]
    if row_vals:
        print(f"Row {r}: " + " | ".join(row_vals))

print("\n--- AI Valuation Summary ---")
ws_ai = wb['AI Valuation Summary']
for r in range(1, 40):
    row_vals = [f'{c}{r}: {ws_ai[f"{c}{r}"].value}' for c in ['A','B','C','D','E','F','G'] if ws_ai[f'{c}{r}'].value is not None]
    if row_vals:
        print(f"Row {r}: " + " | ".join(row_vals))
