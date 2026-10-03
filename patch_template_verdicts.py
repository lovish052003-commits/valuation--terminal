import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx')

if 'Comp_Valuation' in wb.sheetnames:
    ws_comp = wb['Comp_Valuation']
    form_o = '=IF(O37="N/A","N/A",IF(O37>$D$12,TEXT((O37-$D$12)/$D$12,"0.0%")&" Discount",TEXT(($D$12-O37)/$D$12,"0.0%")&" Premium"))'
    ws_comp['O39'] = form_o
    ws_comp['P39'] = form_o.replace('O37', 'P37')
    ws_comp['Q39'] = form_o.replace('O37', 'Q37')

if 'AI Valuation Summary' in wb.sheetnames:
    ws_ai = wb['AI Valuation Summary']
    ws_ai['D5'] = '=IF(C5>=0, TEXT(C5,"0.0%") & " Discount", TEXT(ABS(C5),"0.0%") & " Premium")'

wb.save('ITC Model.xlsx')
wb.close()
print('Successfully patched ITC Model.xlsx base template!')
