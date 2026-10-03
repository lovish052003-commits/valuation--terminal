import os
import shutil
import openpyxl
from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

print("1. Fetching HDFCBANK data from Screener...")
screener_data = fetch_company_data("HDFCBANK")

print("2. Calculating valuation via universal valuation engine...")
val_res = calculate_valuation(screener_data)

print(f"   Beta: {val_res.get('beta')}")
print(f"   Cost of Equity: {val_res.get('cost_of_equity')}")
print(f"   WACC: {val_res.get('wacc')}")
print(f"   Shares: {val_res.get('shares_cr')}")
print(f"   Market Cap: {val_res.get('market_cap_cr')}")
print(f"   Intrinsic Value: {val_res.get('intrinsic_value_per_share')}")

print("3. Exporting valuation model via Excel Exporter...")
exported_file = export_valuation_model(screener_data, val_res)
print(f"   Exported to: {exported_file}")

# Synchronize across all test locations
target_locations = [
    r"c:\Users\LENOVO\Downloads\Advance Financial Project\HDFCBANK_Valuation_Model.xlsx",
    r"c:\Users\LENOVO\Downloads\Advance Financial Project\exports\HDFCBANK_Valuation_Model.xlsx",
    r"C:\Users\LENOVO\Downloads\test\HDFCBANK_Valuation_Model (1).xlsx",
    r"C:\Users\LENOVO\Downloads\test\HDFCBANK_Valuation_Model.xlsx"
]

for loc in target_locations:
    try:
        os.makedirs(os.path.dirname(loc), exist_ok=True)
        shutil.copyfile(exported_file, loc)
        print(f"   Synchronized to: {loc}")
    except Exception as e:
        print(f"   Notice copying to {loc}: {e}")

print("4. Auditing regenerated model...")
wb = openpyxl.load_workbook(exported_file, data_only=False)
wb_v = openpyxl.load_workbook(exported_file, data_only=True)

checks = {}

# 1. Cost of Equity synchronization
wacc_k29 = wb_v['WACC']['K29'].value
wacc_k40 = wb_v['WACC']['K40'].value
wacc_k46 = wb_v['WACC']['K46'].value
ai_e5 = wb_v['AI Valuation Summary']['E5'].value
res_d15 = wb_v['AI Valuation Summary']['D15'].value
ai_b10 = str(wb_v['AI Valuation Summary']['B10'].value or '')

checks['Ke_WACC_equals_AI_E5'] = (abs(float(wacc_k29) - float(ai_e5)) < 1e-4)
checks['Ke_Residual_Income_equals_AI_E5'] = (abs(float(res_d15) - float(ai_e5)) < 1e-4)
checks['Ke_in_AI_B10_matches'] = (f"{float(ai_e5)*100:.2f}%" in ai_b10 or f"{float(wacc_k29)*100:.2f}%" in ai_b10)

# 2. Beta synchronization
wacc_k28 = float(wb_v['WACC']['K28'].value)
checks['Beta_in_AI_B10_matches'] = (f"{wacc_k28:.2f}" in ai_b10)

# 3. Share count reconciliation
ds_b6 = float(wb_v['Data Sheet']['B6'].value)
ds_k70 = float(wb_v['Data Sheet']['K70'].value)
ai_b29 = float(wb_v['AI Valuation Summary']['B29'].value)
dcf_d40 = float(wb_v['DCF']['D40'].value) if 'DCF' in wb.sheetnames else ds_b6

checks['Shares_DS_B6_equals_K70'] = (abs(ds_b6 - (ds_k70 / 1e7)) < 1e-4)
checks['Shares_AI_B29_equals_DS_B6'] = (abs(ai_b29 - ds_b6) < 1e-4)
checks['Shares_DCF_D40_equals_DS_B6'] = (abs(dcf_d40 - ds_b6) < 1e-4)

# 4. Market Cap = Price * Shares
cmp_val = float(wb_v['Data Sheet']['B8'].value)
mcap_val = float(wb_v['Data Sheet']['B9'].value)
calc_mcap = round(cmp_val * ds_b6, 2)
checks['Market_Cap_Reconciliation'] = (abs(calc_mcap - mcap_val) / max(mcap_val, 1) < 0.001)

# 5. Forecast ROE dynamic
res_c15_form = str(wb['AI Valuation Summary']['C15'].value or '')
res_c16_form = str(wb['AI Valuation Summary']['C16'].value or '')
checks['Forecast_ROE_Dynamic'] = ('=G5' in res_c15_form or '=Dupont' in res_c15_form) and ('=C15' in res_c16_form or '=G5' in res_c16_form)
checks['Forecast_ROE_Not_22_pct'] = (float(wb_v['AI Valuation Summary']['C15'].value) != 0.22)

# 6. Cost of equity in residual income dynamic
res_d15_form = str(wb['AI Valuation Summary']['D15'].value or '')
checks['Residual_Income_Ke_Dynamic'] = ('=E$5' in res_d15_form or '=WACC' in res_d15_form)

# 7. Intrinsic Valuation 40/40/20 removed
if 'Intrinsic Valuation' in wb.sheetnames:
    h16 = str(wb['Intrinsic Valuation']['H16'].value or '')
    h17 = str(wb['Intrinsic Valuation']['H17'].value or '')
    h18 = str(wb['Intrinsic Valuation']['H18'].value or '')
    checks['40_40_20_allocation_removed'] = ('0.4' not in h16 and '0.4' not in h17 and '0.2' not in h18 and h16 == 'N/A')
else:
    checks['40_40_20_allocation_removed'] = True

# 8. Automatic verdict removed
if 'Comp_Valuation' in wb.sheetnames:
    q39_form = str(wb['Comp_Valuation']['Q39'].value or '')
    q39_val = str(wb_v['Comp_Valuation']['Q39'].value or '')
    checks['Verdict_Removed_from_Formula'] = not any(w in q39_form for w in ['"Undervalued"', '"Overvalued"', '"BUY"', '"SELL"', '"HOLD"'])
    checks['Verdict_Removed_from_Value'] = not any(w == q39_val for w in ['Undervalued', 'Overvalued', 'BUY', 'SELL', 'HOLD'])
else:
    checks['Verdict_Removed_from_Formula'] = True
    checks['Verdict_Removed_from_Value'] = True

# 9. Excel errors eliminated
error_tokens = ['#VALUE!', '#REF!', '#DIV/0!', '#NAME?', '#N/A']
found_errors = []
for sname in wb_v.sheetnames:
    ws = wb_v[sname]
    for r in range(1, min(ws.max_row+1, 120)):
        for c in range(1, min(ws.max_column+1, 35)):
            val = str(ws.cell(r, c).value or '')
            if any(err == val for err in error_tokens):
                found_errors.append(f"{sname}!{openpyxl.utils.get_column_letter(c)}{r} = {val}")

checks['Zero_Excel_Errors'] = (len(found_errors) == 0)
if found_errors:
    print("Found unexpected errors:", found_errors[:10])

# Check stray '#' tokens
hfs_a5 = str(wb['Historical FS']['A5'].value or '')
iv_a21 = str(wb['Intrinsic Valuation']['A21'].value or '') if 'Intrinsic Valuation' in wb.sheetnames else 'S.No.'
checks['Historical_FS_A5_No_Stray_Hash'] = (hfs_a5 != '#' and hfs_a5 == 'S.No.')
checks['Intrinsic_Valuation_A21_No_Stray_Hash'] = (iv_a21 != '#' and iv_a21 == 'S.No.')

print("\n================ FINAL REPAIR VALIDATION REPORT ================")
for k, v in checks.items():
    print(f"  {k:<45} : {'PASS' if v else 'FAIL'}")
print("================================================================")
