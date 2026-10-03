import openpyxl
import os

target_path = r'C:\Users\LENOVO\Downloads\test\SUNPHARMA_Valuation_Model (1).xlsx'
wb = openpyxl.load_workbook(target_path, data_only=False)
wb_data = openpyxl.load_workbook(target_path, data_only=True)

out_file = r'C:\Users\LENOVO\Downloads\Advance Financial Project\scratch\sunpharma_audit_report.txt'

with open(out_file, 'w', encoding='utf-8') as f:
    f.write("=== SHEET NAMES ===\n")
    f.write(str(wb.sheetnames) + "\n\n")

    f.write("=== 1. HISTORICAL FS AUDIT ===\n")
    if 'Historical FS' in wb.sheetnames:
        ws = wb['Historical FS']
        ws_d = wb_data['Historical FS']
        for r in range(1, 80):
            lbl = ws.cell(row=r, column=2).value
            if lbl:
                f.write(f"Row {r:2d} | Label: {str(lbl):<40}\n")
                f.write(f"   Formula C..G: {[ws.cell(row=r, column=c).value for c in range(3, 8)]}\n")
                f.write(f"   Values  C..G: {[ws_d.cell(row=r, column=c).value for c in range(3, 8)]}\n")

    f.write("\n=== 2. DATA SHEET & RAW FS AUDIT (TAX & OPERATING METRICS) ===\n")
    if 'Data Sheet' in wb.sheetnames:
        ws = wb['Data Sheet']
        ws_d = wb_data['Data Sheet']
        for r in range(1, 100):
            lbl = ws.cell(row=r, column=2).value
            if lbl:
                lbl_str = str(lbl).strip()
                if any(k in lbl_str.lower() for k in ['tax', 'ebit', 'operating', 'depreciation', 'sales', 'revenue', 'debt', 'cash', 'shares', 'price']):
                    f.write(f"Data Sheet Row {r:2d} | B: {lbl_str:<35} | K (Latest Form): {ws.cell(row=r, column=11).value} | K (Val): {ws_d.cell(row=r, column=11).value}\n")

    if 'Raw FS' in wb.sheetnames:
        ws = wb['Raw FS']
        ws_d = wb_data['Raw FS']
        for r in range(1, 70):
            b_val = ws.cell(row=r, column=2).value
            if b_val:
                f.write(f"Raw FS Row {r:2d} | B: {str(b_val):<30} | M: {ws_d.cell(row=r, column=13).value} | N: {ws_d.cell(row=r, column=14).value} | AS: {ws_d.cell(row=r, column=45).value} | AT: {ws_d.cell(row=r, column=46).value} | AU: {ws_d.cell(row=r, column=47).value} | AV: {ws_d.cell(row=r, column=48).value} | AW: {ws_d.cell(row=r, column=49).value}\n")

    f.write("\n=== 3. DCF AUDIT ===\n")
    if 'DCF' in wb.sheetnames:
        ws = wb['DCF']
        ws_d = wb_data['DCF']
        for r in range(1, 55):
            b_val = ws.cell(row=r, column=2).value
            c_val = ws.cell(row=r, column=3).value
            d_val = ws.cell(row=r, column=4).value
            d_val_d = ws_d.cell(row=r, column=4).value
            i_val = ws.cell(row=r, column=9).value
            m_val = ws.cell(row=r, column=13).value
            if b_val or c_val or d_val:
                f.write(f"DCF Row {r:2d} | B: {str(b_val):<30} | C: {str(c_val):<20} | D Form: {str(d_val):<25} | D Val: {str(d_val_d):<15} | I: {str(i_val):<20} | M: {str(m_val):<20}\n")

    f.write("\n=== 4. WACC AUDIT ===\n")
    if 'WACC' in wb.sheetnames:
        ws = wb['WACC']
        ws_d = wb_data['WACC']
        for r in range(1, 50):
            b_val = ws.cell(row=r, column=2).value
            k_val = ws.cell(row=r, column=11).value
            k_val_d = ws_d.cell(row=r, column=11).value
            if b_val or k_val:
                f.write(f"WACC Row {r:2d} | B: {str(b_val):<30} | K Form: {str(k_val):<30} | K Val: {str(k_val_d)}\n")

    f.write("\n=== 5. BETA-REGRESSION AUDIT (Around row 256) ===\n")
    if 'Beta-Regression' in wb.sheetnames:
        ws = wb['Beta-Regression']
        ws_d = wb_data['Beta-Regression']
        for r in range(250, 265):
            row_f = [ws.cell(row=r, column=c).value for c in range(1, 8)]
            row_v = [ws_d.cell(row=r, column=c).value for c in range(1, 8)]
            if any(row_f):
                f.write(f"Beta Row {r:3d} | Form: {row_f}\n")
                f.write(f"          | Val : {row_v}\n")
        f.write("Beta Stats:\n")
        for r in range(1, 15):
            f.write(f"Stats Row {r:2d} | G: {ws.cell(row=r, column=7).value} (Val: {ws_d.cell(row=r, column=7).value}) | H: {ws.cell(row=r, column=8).value} (Val: {ws_d.cell(row=r, column=8).value}) | I: {ws.cell(row=r, column=9).value} (Val: {ws_d.cell(row=r, column=9).value})\n")

    f.write("\n=== 6. COMP_VALUATION AUDIT ===\n")
    if 'Comp_Valuation' in wb.sheetnames:
        ws = wb['Comp_Valuation']
        ws_d = wb_data['Comp_Valuation']
        for r in range(10, 40):
            b_v = ws_d.cell(row=r, column=2).value
            c_v = ws_d.cell(row=r, column=3).value
            d_v = ws_d.cell(row=r, column=4).value
            f_v = ws_d.cell(row=r, column=6).value
            g_v = ws_d.cell(row=r, column=7).value
            h_v = ws_d.cell(row=r, column=8).value
            k_v = ws_d.cell(row=r, column=11).value
            l_v = ws_d.cell(row=r, column=12).value
            m_v = ws_d.cell(row=r, column=13).value
            n_v = ws_d.cell(row=r, column=14).value
            o_v = ws_d.cell(row=r, column=15).value
            p_v = ws_d.cell(row=r, column=16).value
            q_v = ws_d.cell(row=r, column=17).value
            f.write(f"Comp Row {r:2d} | B:{str(b_v):<20} | C:{str(c_v):<10} | D:{str(d_v):<8} | F:{str(f_v):<10} | G:{str(g_v):<10} | H:{str(h_v):<10} | K:{str(k_v):<10} | L:{str(l_v):<10} | M:{str(m_v):<10} | N:{str(n_v):<8} | O:{str(o_v):<8} | P:{str(p_v):<8} | Q:{str(q_v):<8}\n")

    f.write("\n=== 7. AI VALUATION SUMMARY AUDIT ===\n")
    if 'AI Valuation Summary' in wb.sheetnames:
        ws = wb['AI Valuation Summary']
        ws_d = wb_data['AI Valuation Summary']
        for r in range(1, 45):
            b_f = ws.cell(row=r, column=2).value
            c_f = ws.cell(row=r, column=3).value
            b_v = ws_d.cell(row=r, column=2).value
            c_v = ws_d.cell(row=r, column=3).value
            if b_f or c_f:
                f.write(f"AI Row {r:2d} | B Form: {str(b_f):<30} | B Val: {str(b_v):<30} | C Form: {str(c_f):<30} | C Val: {str(c_v)}\n")

print(f"Audit completed and written to {out_file}")
