import openpyxl

target_path = r'C:\Users\LENOVO\Downloads\test\SUNPHARMA_Valuation_Model (1).xlsx'
wb_f = openpyxl.load_workbook(target_path, data_only=False)
wb_d = openpyxl.load_workbook(target_path, data_only=True)

out_file = r'C:\Users\LENOVO\Downloads\Advance Financial Project\scratch\wacc_intrinsic_audit.txt'

with open(out_file, 'w', encoding='utf-8') as f:
    if 'Intrinsic Valuation' in wb_f.sheetnames:
        ws_f = wb_f['Intrinsic Valuation']
        ws_d = wb_d['Intrinsic Valuation']
        f.write("=== INTRINSIC VALUATION SHEET ===\n")
        for r in range(1, 65):
            b_val = ws_f.cell(row=r, column=2).value
            l_form = ws_f.cell(row=r, column=12).value
            l_val = ws_d.cell(row=r, column=12).value
            if b_val or l_form:
                f.write(f"Row {r:2d} | B: {str(b_val):<35} | L Form: {str(l_form):<25} | L Val: {str(l_val)}\n")

    if 'WACC' in wb_f.sheetnames:
        ws_f = wb_f['WACC']
        ws_d = wb_d['WACC']
        f.write("\n=== WACC SHEET (Cols B to K) ===\n")
        for r in range(1, 50):
            row_f = {openpyxl.utils.get_column_letter(c): ws_f.cell(row=r, column=c).value for c in range(2, 12)}
            row_v = {openpyxl.utils.get_column_letter(c): ws_d.cell(row=r, column=c).value for c in range(2, 12)}
            if any(row_f.values()):
                f.write(f"Row {r:2d} Form: {row_f}\n")
                f.write(f"       Val : {row_v}\n")

    if 'Beta-Regression' in wb_f.sheetnames:
        ws_f = wb_f['Beta-Regression']
        ws_d = wb_d['Beta-Regression']
        f.write("\n=== BETA REGRESSION STATS & FORMULAS ===\n")
        for r in range(1, 20):
            row_f = {openpyxl.utils.get_column_letter(c): ws_f.cell(row=r, column=c).value for c in range(7, 15)}
            row_v = {openpyxl.utils.get_column_letter(c): ws_d.cell(row=r, column=c).value for c in range(7, 15)}
            f.write(f"Stats Row {r:2d} Form: {row_f}\n")
            f.write(f"            Val : {row_v}\n")

print("WACC & Intrinsic audit written to scratch/wacc_intrinsic_audit.txt")
