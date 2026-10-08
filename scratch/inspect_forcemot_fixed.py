import openpyxl
import os

fixed_path = r"FORCEMOT_Valuation_Model_FIXED.xlsx"
orig_path = r"exports\FORCEMOT_Valuation_Model.xlsx"

print("Fixed exists:", os.path.exists(fixed_path))
print("Orig exists:", os.path.exists(orig_path))

wb_fixed = openpyxl.load_workbook(fixed_path, data_only=False)
sheets_fixed = wb_fixed.sheetnames
print("Fixed sheets:", sheets_fixed)

if os.path.exists(orig_path):
    wb_orig = openpyxl.load_workbook(orig_path, data_only=False)
    sheets_orig = wb_orig.sheetnames
    print("Orig sheets:", sheets_orig)

    # Compare sheet by sheet
    diffs = []
    for s in sheets_fixed:
        if s not in sheets_orig:
            print(f"Sheet {s} is NEW in FIXED!")
            continue
        ws_f = wb_fixed[s]
        ws_o = wb_orig[s]
        
        max_r = max(ws_f.max_row, ws_o.max_row)
        max_c = max(ws_f.max_column, ws_o.max_column)
        
        for r in range(1, min(max_r + 1, 100)):
            for c in range(1, min(max_c + 1, 40)):
                vf = ws_f.cell(r, c).value
                vo = ws_o.cell(r, c).value
                if vf != vo:
                    diffs.append((s, openpyxl.utils.get_column_letter(c) + str(r), vo, vf))
                    
    print(f"\nTotal cell differences found: {len(diffs)}")
    print("\n--- SAMPLE DIFFERENCES (first 50) ---")
    for s, cell, vo, vf in diffs[:50]:
        print(f"[{s}!{cell}] ORIG: {vo}  -->  FIXED: {vf}")
        
    if len(diffs) > 50:
        print(f"\n--- NEXT 50 DIFFERENCES (50 to 100) ---")
        for s, cell, vo, vf in diffs[50:100]:
            print(f"[{s}!{cell}] ORIG: {vo}  -->  FIXED: {vf}")
            
    # Group diffs by sheet
    from collections import Counter
    sheet_counts = Counter(d[0] for d in diffs)
    print("\nDiffs per sheet:", sheet_counts)
else:
    # Just inspect what's inside fixed
    for s in sheets_fixed:
        print(f"Sheet: {s}, rows: {wb_fixed[s].max_row}, cols: {wb_fixed[s].max_column}")
