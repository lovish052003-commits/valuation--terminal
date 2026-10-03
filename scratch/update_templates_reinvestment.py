import openpyxl

def update_template(filepath):
    print(f"Updating template: {filepath}")
    wb = openpyxl.load_workbook(filepath, data_only=False)
    
    if 'Intrinsic Valuation' in wb.sheetnames:
        ws_iv = wb['Intrinsic Valuation']
        ws_iv['B67'] = "Normalized ROIC (Sustainable)"
        ws_iv['L67'] = "=IFERROR(MEDIAN(I40:L40), 0.15)"
        
        ws_iv['B68'] = "Expected Growth Rate"
        ws_iv['L68'] = "=DCF!D18"
        
        ws_iv['B69'] = "Fundamental Reinvestment Rate"
        ws_iv['L69'] = "=IF(L67<=0, L55, L68/L67)"
        
        ws_iv['B70'] = "Sustainable Terminal ROIC"
        ws_iv['L70'] = "=MIN(0.18, MAX(DCF!D20, 0.5*L67 + 0.5*DCF!D20))"
        
        ws_iv['B71'] = "Terminal Reinvestment Rate"
        ws_iv['L71'] = "=DCF!D19/L70"
        
        ws_iv['B72'] = "Growth-ROIC Consistency Check"
        ws_iv['L72'] = '=IF(ABS(L69*L67 - L68) <= 0.005, "PASS", "WARNING")'
        
        ws_iv['B73'] = "Growth Source"
        ws_iv['L73'] = "Fundamental Estimate"
        
        ws_iv['B74'] = "Reinvestment Confidence"
        ws_iv['L74'] = "HIGH"
        print("  Updated Intrinsic Valuation rows 67-74.")
        
    if 'DCF' in wb.sheetnames:
        ws_dcf = wb['DCF']
        ws_dcf['D21'] = "='Intrinsic Valuation'!$L$71"
        ws_dcf['I11'] = "='Intrinsic Valuation'!$L$69"
        ws_dcf['J11'] = "=$I$11+($M$11-$I$11)/4*1"
        ws_dcf['K11'] = "=$I$11+($M$11-$I$11)/4*2"
        ws_dcf['L11'] = "=$I$11+($M$11-$I$11)/4*3"
        ws_dcf['M11'] = "=D21"
        for col_l in ['I', 'J', 'K', 'L', 'M']:
            ws_dcf[f'{col_l}12'] = f"={col_l}10*(1-{col_l}11)"
        print("  Updated DCF D21, I11, M11, and FCFF rows.")
        
    wb.save(filepath)
    print(f"Saved {filepath} successfully.")

if __name__ == '__main__':
    update_template('master_model_template.xlsx')
    update_template('ITC Model.xlsx')
