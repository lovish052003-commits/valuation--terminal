import os, sys, gc
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

def test_britannia():
    print("Testing Britannia Export...")
    data = fetch_company_data('BRITANNIA')
    print("Company Name:", data.get('company_name'))
    print("CMP:", data.get('current_price'))
    print("Market Cap:", data.get('market_cap_cr'))
    print("Shares in Cr:", data.get('shares_in_cr'))

    val = calculate_valuation(data)
    out_file = export_valuation_model(data, val, "### Britannia Institutional Valuation Report")
    print(f"\nExported to: {out_file}")

    wb = openpyxl.load_workbook(out_file, read_only=True)
    print(f"Total Sheet Count: {len(wb.sheetnames)}")
    print("Sheet Names:")
    for s in wb.sheetnames:
        print(f"  - {s}")
    wb.close()

    # Also inspect Data Sheet in detail
    wb2 = openpyxl.load_workbook(out_file, data_only=True)
    ws = wb2['Data Sheet']
    print(f"\nData Sheet Inspection:")
    print(f"  B1 (Company Name): {ws['B1'].value}")
    print(f"  B93:K93 (Adj Shares): {[ws.cell(93, c).value for c in range(2, 12)]}")
    print(f"  B70 (Shares): {ws['B70'].value}")
    print(f"  B67 (Receivables): {ws['B67'].value}")
    print(f"  B68 (Inventory): {ws['B68'].value}")
    print(f"  B69 (Cash & Bank): {ws['B69'].value}")
    print(f"  B90 (Price): {ws['B90'].value}")
    wb2.close()

if __name__ == '__main__':
    test_britannia()
