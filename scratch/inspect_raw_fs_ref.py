import openpyxl

wb = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=True)
ws = wb['Raw FS']

print("=== RAW FS in Nestle India Model.xlsx ===")
print("Row 3 (BS Periods):", [ws.cell(3, c).value for c in range(2, 16)])
for r in [5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 18, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 33, 35, 36, 38, 39, 40, 41, 42, 43, 45]:
    lbl = ws.cell(r, 2).value
    vals = [ws.cell(r, c).value for c in range(3, 15)]
    print(f"Row {r:2d} | {str(lbl):<28} | {vals[-5:]}")

print("\nRow 3 (PL Periods):", [ws.cell(3, c).value for c in range(19, 33)])
for r in [4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32]:
    lbl = ws.cell(r, 18).value
    vals = [ws.cell(r, c).value for c in range(19, 33)]
    print(f"Row {r:2d} | {str(lbl):<28} | {vals[-5:]}")
