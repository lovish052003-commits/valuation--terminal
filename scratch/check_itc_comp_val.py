import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['Comp_Valuation']

print("--- Comp_Valuation in ITC Model.xlsx ---")
for r in range(10, 40):
    b = ws.cell(r, 2).value
    c = ws.cell(r, 3).value
    d = ws.cell(r, 4).value
    n = ws.cell(r, 14).value
    o = ws.cell(r, 15).value
    p = ws.cell(r, 16).value
    q = ws.cell(r, 17).value
    print(f"Row {r:2d}: B={b} | C={c} | D={d} | N={n} | O={o} | P={p} | Q={q}")
