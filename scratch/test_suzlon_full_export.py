import sys, os
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__) + '/..'))
import screener_client, valuation_engine, excel_exporter
import openpyxl

print("1. Fetching Suzlon data...")
d = screener_client.fetch_company_data('Suzlon')
val = valuation_engine.calculate_valuation(d)

print("2. Exporting Suzlon model...")
dest = excel_exporter.export_valuation_model(d, val, "")
print("Exported to:", dest)

wb = openpyxl.load_workbook(dest, data_only=True)
print("\nSheets:", wb.sheetnames)

if 'AI Valuation Summary' in wb.sheetnames:
    ws = wb['AI Valuation Summary']
    print("\n--- AI Valuation Summary ---")
    for r in range(1, 35):
        row_vals = [str(ws.cell(r, c).value).encode('ascii', 'replace').decode() for c in range(1, 8)]
        if any(v != 'None' for v in row_vals):
            print(f"Row {r:2d}: " + " | ".join(f"{openpyxl.utils.get_column_letter(c)}={row_vals[c-1]}" for c in range(1, 8)))

if 'Comp_Valuation' in wb.sheetnames:
    ws = wb['Comp_Valuation']
    print("\n--- Comp_Valuation Rows 10-25 ---")
    for r in range(10, 26):
        b = str(ws.cell(r, 2).value).encode('ascii', 'replace').decode()
        n = str(ws.cell(r, 14).value).encode('ascii', 'replace').decode()
        o = str(ws.cell(r, 15).value).encode('ascii', 'replace').decode()
        p = str(ws.cell(r, 16).value).encode('ascii', 'replace').decode()
        q = str(ws.cell(r, 17).value).encode('ascii', 'replace').decode()
        print(f"Row {r:2d}: B={b} | N={n} | O={o} | P={p} | Q={q}")

if 'Data Sheet' in wb.sheetnames:
    ws = wb['Data Sheet']
    print("\n--- Data Sheet ---")
    print("B1:", ws['B1'].value)
    print("B8 (Price):", ws['B8'].value)
    print("K69 (Cash):", ws['K69'].value)
    print("K70 (Shares):", ws['K70'].value)
    print("K59 (Debt):", ws['K59'].value)

if 'DCF' in wb.sheetnames:
    ws = wb['DCF']
    print("\n--- DCF ---")
    print("D37 (Cash):", ws['D37'].value)
    print("D38 (Debt):", ws['D38'].value)
    print("D40 (Shares):", ws['D40'].value)
    print("D42 (IV):", ws['D42'].value)
    print("D44 (CMP):", ws['D44'].value)
