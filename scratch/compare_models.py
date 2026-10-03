import openpyxl
import os

def examine_model(filepath, label):
    print(f"\n==================================================")
    print(f"FILE: {label} -> {os.path.basename(filepath)}")
    print(f"==================================================")
    if not os.path.exists(filepath):
        print(f"File not found: {filepath}")
        return

    wb = openpyxl.load_workbook(filepath, data_only=False)
    wb_val = openpyxl.load_workbook(filepath, data_only=True)
    
    # Find Comp_Valuation sheet
    sheet_name = None
    for s in wb.sheetnames:
        if 'comp' in s.lower():
            sheet_name = s
            break
    
    if not sheet_name:
        print("No Comp Valuation sheet found!")
        return

    ws = wb[sheet_name]
    ws_val = wb_val[sheet_name]
    
    print(f"Sheet Name: '{sheet_name}'")
    print("\n--- Header (Rows 9-10) ---")
    headers = [ws.cell(10, c).value for c in range(2, 14)]
    print("Columns (B to M):", headers)
    
    print("\n--- Peer Rows (Rows 12 to 21) ---")
    for r in range(12, 22):
        company = ws_val.cell(r, 2).value
        ticker = ws_val.cell(r, 3).value
        price = ws_val.cell(r, 4).value
        shares = ws_val.cell(r, 5).value
        eq_val_f = ws.cell(r, 6).value
        eq_val = ws_val.cell(r, 6).value
        ev_f = ws.cell(r, 8).value
        ev = ws_val.cell(r, 8).value
        ev_rev_f = ws.cell(r, 9).value
        ev_rev = ws_val.cell(r, 9).value
        ev_ebitda_f = ws.cell(r, 10).value
        ev_ebitda = ws_val.cell(r, 10).value
        
        if company or ticker:
            print(f"Row {r:02d}: Company='{company}' | Ticker='{ticker}' | Price={price} | EV/Rev(Formula)='{ev_rev_f}' Val={ev_rev} | EV/EBITDA(Formula)='{ev_ebitda_f}' Val={ev_ebitda}")

    print("\n--- Summary Multiples (Rows 23 to 28) ---")
    for r in range(23, 29):
        metric = ws.cell(r, 2).value
        ev_rev_f = ws.cell(r, 9).value
        ev_rev_v = ws_val.cell(r, 9).value
        ev_eb_f = ws.cell(r, 10).value
        ev_eb_v = ws_val.cell(r, 10).value
        print(f"Row {r}: {metric:15s} | EV/Rev: Formula='{ev_rev_f}' Val={ev_rev_v} | EV/EBITDA: Formula='{ev_eb_f}' Val={ev_eb_v}")

    print("\n--- Valuation / Implied Target Price (Rows 30 to 39) ---")
    for r in range(30, 40):
        label_cell = ws.cell(r, 2).value
        c9_f = ws.cell(r, 9).value
        c9_v = ws_val.cell(r, 9).value
        c10_f = ws.cell(r, 10).value
        c10_v = ws_val.cell(r, 10).value
        if label_cell or c9_f or c10_f:
            print(f"Row {r:02d}: {str(label_cell):35s} | Col I: {str(c9_f):25s} (Val={c9_v}) | Col J: {str(c10_f):25s} (Val={c10_v})")

# User model
p1 = r"C:\Users\LENOVO\Downloads\Advance Financial Project\Varun Beverages model.xlsx"
# Generated model (4)
p2 = r"C:\Users\LENOVO\Downloads\Advance Financial Project\VBL_Valuation_Model (4).xlsx"
if not os.path.exists(p2):
    p2 = r"C:\Users\LENOVO\Downloads\VBL_Valuation_Model (4).xlsx"

examine_model(p1, "User's Model (Varun Beverages model.xlsx)")
examine_model(p2, "Generated Model (VBL_Valuation_Model (4).xlsx)")
