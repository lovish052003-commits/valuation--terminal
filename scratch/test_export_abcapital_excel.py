import sys, os
sys.path.insert(0, os.path.abspath('.'))

from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model
import openpyxl

print("Fetching ABCAPITAL data...")
screener_data = fetch_company_data('ABCAPITAL')
print("Calculating valuation...")
val_result = calculate_valuation(screener_data)

dest_file = os.path.abspath("test_abcapital_model.xlsx")
if os.path.exists(dest_file):
    try:
        os.remove(dest_file)
    except:
        pass

print(f"Exporting valuation model to {dest_file}...")
out_path = export_valuation_model(screener_data, val_result, "Report")
print(f"Export completed: {out_path}")

wb = openpyxl.load_workbook(out_path, data_only=False)
ws_comp = wb['Comp_Valuation']
ws_raw = wb['Raw FS']

print("\n--- Comp_Valuation Rows 12 to 21 ---")
for r in range(12, 22):
    c_name_val = ws_comp.cell(row=r, column=2).value
    cmp_val = ws_comp.cell(row=r, column=4).value
    shares_val = ws_comp.cell(row=r, column=5).value
    ev_rev = ws_comp.cell(row=r, column=15).value
    ev_ebitda = ws_comp.cell(row=r, column=16).value
    pe = ws_comp.cell(row=r, column=17).value
    print(f"Row {r}: B={c_name_val} | D={cmp_val} | E={shares_val} | EV/Rev={ev_rev} | EV/EBITDA={ev_ebitda} | P/E={pe}")

print("\n--- Comp_Valuation Rows 23 to 28 (Benchmark Stats) ---")
for r in range(23, 29):
    lbl = ws_comp.cell(row=r, column=2).value
    o_val = ws_comp.cell(row=r, column=15).value
    p_val = ws_comp.cell(row=r, column=16).value
    q_val = ws_comp.cell(row=r, column=17).value
    print(f"Row {r} ({lbl}): O={o_val} | P={p_val} | Q={q_val}")

print("\n--- Comp_Valuation Rows 30 to 39 (Target Valuation) ---")
for r in range(30, 40):
    lbl = ws_comp.cell(row=r, column=2).value
    o_val = ws_comp.cell(row=r, column=15).value
    p_val = ws_comp.cell(row=r, column=16).value
    q_val = ws_comp.cell(row=r, column=17).value
    print(f"Row {r} ({lbl}): O={o_val} | P={p_val} | Q={q_val}")

print("\n--- Raw FS Rows 56 to 66 ---")
for r in range(56, 67):
    l_name = ws_raw.cell(row=r, column=12).value
    cmp_v = ws_raw.cell(row=r, column=13).value
    sh_v = ws_raw.cell(row=r, column=14).value
    print(f"Raw FS Row {r}: L={l_name} | M={cmp_v} | N={sh_v}")

# Assertions
# In Comp_Valuation, Row 13 connects to Raw FS L58 (Target Company)
assert "Raw FS'!L58" in str(ws_comp.cell(row=13, column=2).value), "Row 13 B cell must connect to Raw FS L58"
assert "Aditya Birla" in str(ws_raw.cell(row=58, column=12).value), f"Raw FS Row 58 must be Aditya Birla, got {ws_raw.cell(row=58, column=12).value}"

# Benchmark stats formulas must exclude Row 13
for col_idx in [15, 16, 17]:
    f26 = str(ws_comp.cell(row=26, column=col_idx).value)
    assert "14:O21" in f26 or "14:P21" in f26 or "14:Q21" in f26, f"Average formula must exclude Row 13, got {f26}"

# Target valuation formulas must read from Row 13
assert "=K13*O26" == str(ws_comp.cell(row=32, column=15).value), "Row 32 Col O must be =K13*O26"
assert "=$G$13" == str(ws_comp.cell(row=33, column=15).value), "Row 33 Net debt must be =$G$13"
assert "=$E$13" == str(ws_comp.cell(row=35, column=15).value), "Row 35 Shares must be =$E$13"

print("\nALL EXCEL EXPORT VERIFICATIONS PASSED!")
