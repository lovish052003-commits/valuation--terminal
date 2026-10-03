import openpyxl

wb_path = r'C:\Users\LENOVO\Downloads\test\ADANIPOWER_Valuation_Model (2).xlsx'
wb_v = openpyxl.load_workbook(wb_path, data_only=True)
wb_f = openpyxl.load_workbook(wb_path, data_only=False)

print('--- DATA SHEET ROWS 45-75 ---')
ds_v = wb_v['Data Sheet']
ds_f = wb_f['Data Sheet']
for r in range(45, 75):
    row_vals = [ds_v.cell(row=r, column=c).value for c in range(1, 10)]
    forms = [ds_f.cell(row=r, column=c).value for c in range(1, 10)]
    if any(x is not None for x in row_vals):
        print(f'DS Row {r}: vals={row_vals[:5]}, forms={forms[:5]}')

print('\n--- ALTMAN Z ROWS 60-95 ---')
az_sheet = "Altman's Z Score"
az_v = wb_v[az_sheet]
az_f = wb_f[az_sheet]
for r in range(60, 95):
    row_vals = [az_v.cell(row=r, column=c).value for c in range(1, 12)]
    forms = [az_f.cell(row=r, column=c).value for c in range(1, 12)]
    if any(x is not None for x in row_vals):
        print(f'AZ Row {r}: vals={row_vals[:6]}, forms={forms[:6]}')

print('\n--- DUPONT ROWS 55-85 ---')
dp_sheet = 'Dupont Analysis'
dp_v = wb_v[dp_sheet]
dp_f = wb_f[dp_sheet]
for r in range(55, 85):
    row_vals = [dp_v.cell(row=r, column=c).value for c in range(1, 12)]
    forms = [dp_f.cell(row=r, column=c).value for c in range(1, 12)]
    if any(x is not None for x in row_vals):
        print(f'DP Row {r}: vals={row_vals[:6]}, forms={forms[:6]}')
