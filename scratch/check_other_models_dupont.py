import openpyxl

for fn in ['Nestle India Model.xlsx', 'Tata Steel.xlsx', 'Varun Beverages model.xlsx']:
    wb = openpyxl.load_workbook(fn, data_only=False)
    for sn in ['Dupont Analysis', "Altman's Z Score"]:
        if sn in wb.sheetnames:
            ws = wb[sn]
            r_about = 8
            r_upd = 37 if sn == 'Dupont Analysis' else 36
            print(f"{fn} -> {sn}:")
            print("  About:", repr(str(ws.cell(r_about, 2).value)[:80]))
            print("  Upd1:", repr(str(ws.cell(r_upd, 2).value)[:80]))
