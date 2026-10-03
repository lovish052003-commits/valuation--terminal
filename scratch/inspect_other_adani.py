import openpyxl

for fn in ['exports/ADANIENT_Valuation_Model_1789737880.xlsx', 'exports/ADANIENT_Valuation_Model_reconciled.xlsx']:
    wb = openpyxl.load_workbook(fn, data_only=True)
    print(f"\n=== {fn} ===")
    print("Sheets:", wb.sheetnames[:5])
    if 'Data Sheet' in wb.sheetnames:
        print("B1:", wb['Data Sheet']['B1'].value)
        print("K17:", wb['Data Sheet']['K17'].value)
        print("K30:", wb['Data Sheet']['K30'].value)
    if 'DCF' in wb.sheetnames:
        print("D42 (IV):", wb['DCF']['D42'].value)
        print("D44 (CMP):", wb['DCF']['D44'].value)
