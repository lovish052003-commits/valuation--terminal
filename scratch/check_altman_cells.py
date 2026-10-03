import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=True)
ws = wb["Altman's Z Score"]
for r in range(45, 96):
    for c in range(1, 10):
        v = ws.cell(row=r, column=c).value
        if v is not None and ('altman' in str(v).lower() or 'score' in str(v).lower() or 'zone' in str(v).lower()):
            col_letter = openpyxl.utils.get_column_letter(c)
            print(f"{col_letter}{r}: {str(v)[:40]}")
