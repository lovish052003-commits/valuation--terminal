import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import screener_client
import excel_exporter
import openpyxl

print("==================================================")
print("TESTING COMPLETE VALUATION EXPORT FOR WIPRO & BRITANNIA")
print("==================================================")

for sym in ['WIPRO', 'BRITANNIA']:
    print(f"\n>>> PROCESSING {sym}...")
    screener_data = screener_client.fetch_company_data(sym)
    val_result = {
        'intrinsic_value': 500.0,
        'cmp': screener_data.get('current_price', 100.0),
        'wacc': 0.11,
        'terminal_growth': 0.05
    }
    
    # Export valuation model
    exported_path = excel_exporter.export_valuation_model(screener_data, val_result)
    print(f"Exported to: {exported_path}")
    
    # Verify workbook
    wb = openpyxl.load_workbook(exported_path, data_only=True)
    
    # 1. Data Sheet P&L verification
    ws_data = wb['Data Sheet']
    print(f"\n--- {sym} Data Sheet Rows 17-24 ---")
    sales_row = [ws_data.cell(row=17, column=c).value for c in range(2, 12)]
    raw_mat_row = [ws_data.cell(row=18, column=c).value for c in range(2, 12)]
    emp_cost_row = [ws_data.cell(row=22, column=c).value for c in range(2, 12)]
    oth_exp_row = [ws_data.cell(row=24, column=c).value for c in range(2, 12)]
    print(f"Sales:       {sales_row}")
    print(f"Raw Mat:     {raw_mat_row}")
    print(f"Emp Cost:    {emp_cost_row}")
    print(f"Other Exp:   {oth_exp_row}")
    
    # Verify non-zero
    assert any(v and v > 0 for v in emp_cost_row), f"FAIL: {sym} Employee cost is all zero!"
    assert any(v and v > 0 for v in oth_exp_row), f"FAIL: {sym} Other expenses is all zero!"
    print(f"SUCCESS: {sym} Data Sheet P&L is non-zero and populated!")
    
    # 2. Cash Flow Statement verification
    ws_cfs = wb['Cash Flow Statement']
    print(f"\n--- {sym} Cash Flow Statement Sub-Schedules ---")
    cfo_summary = [ws_cfs.cell(row=6, column=c).value for c in range(3, 15)]
    op_profit = [ws_cfs.cell(row=7, column=c).value for c in range(3, 15)]
    receivables = [ws_cfs.cell(row=8, column=c).value for c in range(3, 15)]
    direct_tax = [ws_cfs.cell(row=12, column=c).value for c in range(3, 15)]
    cfi_summary = [ws_cfs.cell(row=13, column=c).value for c in range(3, 15)]
    fixed_assets_purchased = [ws_cfs.cell(row=14, column=c).value for c in range(3, 15)]
    cff_summary = [ws_cfs.cell(row=24, column=c).value for c in range(3, 15)]
    borrowings = [ws_cfs.cell(row=26, column=c).value for c in range(3, 15)]
    
    print(f"CFO:            {cfo_summary}")
    print(f"Profit from Op: {op_profit}")
    print(f"Direct Taxes:   {direct_tax}")
    print(f"CFI:            {cfi_summary}")
    print(f"Capex (FA pur): {fixed_assets_purchased}")
    print(f"CFF:            {cff_summary}")
    print(f"Borrowings:     {borrowings}")
    
    assert any(v and v != 0 for v in op_profit), f"FAIL: {sym} Profit from operations is all zero!"
    assert any(v and v != 0 for v in fixed_assets_purchased), f"FAIL: {sym} Fixed assets purchased is all zero!"
    print(f"SUCCESS: {sym} Cash Flow Statement sub-schedules are non-zero and populated!")

print("\nALL VERIFICATIONS PASSED PERFECTLY!")
