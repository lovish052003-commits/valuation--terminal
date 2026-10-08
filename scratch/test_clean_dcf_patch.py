import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

fn = r'c:\Users\LENOVO\Downloads\test\RELIANCE_Valuation_Model (5).xlsx'
wb = openpyxl.load_workbook(fn, data_only=False)
ws = wb['DCF']

border_none = Border()
border_ev = Border(top=Side(style='thin'), bottom=Side(style='medium'))
font_bold = Font(name='Calibri', size=11, bold=True)
font_regular = Font(name='Calibri', size=11, bold=False)
currency_fmt = '"\u20b9"\\ #,##0.00;\\("\u20b9"\\ #,##0.00\\);\\-'

# Row 39
ws['B39'] = 'Less: Minority Interest'
ws['D39'] = "='Data Sheet'!K73"
ws['B39'].font = font_bold
ws['D39'].font = font_bold
ws['D39'].number_format = currency_fmt

# Row 40
ws['B40'] = 'Equity Value'
ws['D40'] = '=D35+D37-D38-D39'
ws['B40'].font = font_bold
ws['D40'].font = font_bold
ws['D40'].number_format = currency_fmt

# Row 41
ws['B41'] = 'No. of Shares'
ws['D41'] = "='Data Sheet'!K70/10000000"
ws['B41'].font = font_regular
ws['D41'].font = font_regular
ws['D41'].number_format = '#,##0.00'

# Row 42 (clean blank row)
for col in ['B', 'C', 'D']:
    c = ws[f'{col}42']
    c.value = None
    c.border = border_none
    c.font = font_regular
ws['D42'].number_format = 'General'

# Row 43 (Equity Value per Share)
ws['B43'] = 'Equity Value per Share'
ws['B43'].font = font_bold
ws['B43'].border = border_none
ws['D43'] = '=D40/D41'
ws['D43'].font = font_bold
ws['D43'].border = border_ev
ws['D43'].number_format = currency_fmt

# Row 44 (clean blank row)
for col in ['B', 'C', 'D']:
    c = ws[f'{col}44']
    c.value = None
    c.border = border_none
    c.font = font_regular
ws['D44'].number_format = 'General'

# Row 45 (Share Price)
ws['B45'] = 'Share Price'
ws['B45'].font = font_regular
ws['B45'].border = border_none
ws['D45'] = "='Data Sheet'!B8"
ws['D45'].font = font_regular
ws['D45'].border = border_none
ws['D45'].number_format = currency_fmt

# Row 46 (Upside / Downside)
ws['B46'] = 'Upside / (Downside)'
ws['B46'].font = font_regular
ws['B46'].border = border_none
ws['D46'] = '=D43/D45-1'
ws['D46'].font = font_regular
ws['D46'].border = border_none
ws['D46'].number_format = '0.0%'

# Ensure row 47 is clean
for col in ['B', 'C', 'D']:
    c = ws[f'{col}47']
    c.value = None
    c.border = border_none

print('TEST PATCH PREVIEW:')
for r in range(37, 48):
    b = str(ws.cell(row=r, column=2).value).encode('ascii', 'replace').decode('ascii')
    d = str(ws.cell(row=r, column=4).value).encode('ascii', 'replace').decode('ascii')
    fmt = str(ws.cell(row=r, column=4).number_format).encode('ascii', 'replace').decode('ascii')
    top = str(ws.cell(row=r, column=4).border.top.style if (ws.cell(row=r, column=4).border and ws.cell(row=r, column=4).border.top) else None)
    bot = str(ws.cell(row=r, column=4).border.bottom.style if (ws.cell(row=r, column=4).border and ws.cell(row=r, column=4).border.bottom) else None)
    b_bold = str(ws.cell(row=r, column=2).font.bold if ws.cell(row=r, column=2).font else None)
    print(f'  Row {r:2d}: B={b:28s} D={d:25s} fmt={fmt:20s} top={top:6s} bot={bot:6s} bold={b_bold}')
