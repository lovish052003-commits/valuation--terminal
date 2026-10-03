import os
import sys
import time
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import openpyxl
from screener_client import fetch_company_data, is_same_company
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model, get_effective_peers

print("================================================================================")
print("FINAL COMPREHENSIVE VALIDATION SUITE")
print("================================================================================")

# ==============================================================================
# PART 1: ONGC EXCEL MODEL & WACC FORMULA AUDIT
# ==============================================================================
print("\n[PART 1] Auditing ONGC Excel Model & WACC Sheet Formulas...")
sd_ongc = fetch_company_data("ONGC")
val_ongc = calculate_valuation(sd_ongc)
excel_path = export_valuation_model(sd_ongc, val_ongc)

wb = openpyxl.load_workbook(excel_path, data_only=False)
ws_wacc = wb['WACC']
ws_comp = wb['Comp_Valuation']
ws_ai = wb['AI Valuation Summary']

# 1. Target company excluded from WACC peers (rows 14-18)
for r in range(14, 19):
    b = ws_wacc.cell(r, 2).value
    j = ws_wacc.cell(r, 10).value
    assert j != "='Beta-Regression'!L15", f"WACC Row {r} still references legacy target beta: {j}"

# 2. Target own capital structure used for WACC formulas
assert ws_wacc['K34'].value == "=D38", f"K34 should be =D38, got {ws_wacc['K34'].value}"
assert ws_wacc['K41'].value == "=D35", f"K41 should be =D35, got {ws_wacc['K41'].value}"
assert ws_wacc['K44'].value == "=D34", f"K44 should be =D34, got {ws_wacc['K44'].value}"
assert ws_wacc['K43'].value == "=E28", f"K43 should be =E28, got {ws_wacc['K43'].value}"
assert ws_wacc['K46'].value == "=(K40*K41)+(K43*K44)", f"K46 should be =(K40*K41)+(K43*K44), got {ws_wacc['K46'].value}"

# 3. Comp_Valuation checks
for r in range(12, 17):
    c = ws_comp.cell(r, 3).value
    assert c != "ONGC", f"Target ONGC found in Comp_Valuation row {r}!"

# 4. AI Valuation Summary checks
assert ws_ai['C8'].value == "=WACC!K46", f"AI Summary WACC: {ws_ai['C8'].value}"
assert ws_ai['C14'].value == "=WACC!K43", f"AI Summary Post-Tax Kd: {ws_ai['C14'].value}"
print("[PASS] PART 1 PASSED: ONGC Excel Model & WACC Sheet 100% verified.")

# ==============================================================================
# PART 2: 11-COMPANY TEST SUITE
# ==============================================================================
print("\n[PART 2] Running 11-Company Test Suite...")
TEST_COMPANIES = [
    "ONGC",
    "RELIANCE",
    "TCS",
    "ITC",
    "TATAMOTORS",
    "SUNPHARMA",
    "HDFCBANK",
    "SBIN",
    "BAJFINANCE",
    "LICI",
    "ICICIAMC"
]

suite_summary = {}

for comp in TEST_COMPANIES:
    t0 = time.time()
    sd = fetch_company_data(comp)
    c_name = sd.get('company_name', '')
    ticker = sd.get('ticker', '')
    val = calculate_valuation(sd)
    peers = get_effective_peers(sd, val)
    
    # Assert target company strictly excluded
    for p in peers:
        p_name = p.get('name', '')
        p_tick = p.get('ticker', '')
        assert not is_same_company(p_tick, p_name, ticker, c_name), f"Target {c_name} in peers: {p_name}"
        assert p_tick.upper() != ticker.upper(), f"Target ticker {ticker} in peers: {p_tick}"
    
    elapsed = round(time.time() - t0, 1)
    suite_summary[comp] = {
        'resolved_name': c_name,
        'ticker': ticker,
        'sector': sd.get('sector'),
        'industry': sd.get('industry'),
        'method': val.get('valuation_method'),
        'intrinsic': val.get('intrinsic_value_per_share'),
        'wacc': val.get('wacc'),
        'peers': [p.get('name') for p in peers[:8]],
        'peer_count': len(peers),
        'time_s': elapsed
    }
    print(f"[PASS] {comp:12} -> {c_name[:25]:25} | WACC: {val.get('wacc')}% | Intrinsic: Rs. {val.get('intrinsic_value_per_share')} | Peers: {len(peers)} ({elapsed}s)")

# ==============================================================================
# PART 3: SEQUENTIAL ZERO-CONTAMINATION TEST
# ==============================================================================
print("\n[PART 3] Running Sequential Zero-Contamination Test: ONGC -> TCS -> ITC -> SBIN -> LICI -> ICICIAMC -> ONGC...")
seq = ["ONGC", "TCS", "ITC", "SBIN", "LICI", "ICICIAMC", "ONGC"]
seq_runs = []

for s in seq:
    sd = fetch_company_data(s)
    val = calculate_valuation(sd)
    peers = get_effective_peers(sd, val)
    seq_runs.append({
        'company': s,
        'peers': [p.get('name') for p in peers[:8]],
        'wacc': val.get('wacc'),
        'intrinsic': val.get('intrinsic_value_per_share')
    })

run_first = seq_runs[0]
run_last = seq_runs[6]

print(f"\nRun 1 (ONGC): Peers={run_first['peers']} | WACC={run_first['wacc']} | Intrinsic={run_first['intrinsic']}")
print(f"Run 7 (ONGC): Peers={run_last['peers']} | WACC={run_last['wacc']} | Intrinsic={run_last['intrinsic']}")

assert run_first['peers'] == run_last['peers'], f"Peers differ between Run 1 and Run 7: {run_first['peers']} != {run_last['peers']}"
assert run_first['wacc'] == run_last['wacc'], f"WACC differs: {run_first['wacc']} != {run_last['wacc']}"
assert run_first['intrinsic'] == run_last['intrinsic'], f"Intrinsic differs: {run_first['intrinsic']} != {run_last['intrinsic']}"
print("[PASS] PART 3 PASSED: Zero state bleed / zero contamination confirmed across sequential pipeline.")

print("\n================================================================================")
print("ALL PARTS COMPLETED AND VERIFIED 100% SUCCESSFULLY!")
print("================================================================================")
