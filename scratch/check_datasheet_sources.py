import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=False)
ws = wb['Data Sheet']

print("--- Data Sheet Formulas / Values in Nestle India (2).xlsx ---")
for r in range(1, 95):
    lbl = ws.cell(r, 1).value
    c_vals = [ws.cell(r, c).value for c in range(2, 12)]
    # Check if any cell has a formula or value
    has_formula = any(isinstance(v, str) and v.startswith('=') for v in c_vals)
    has_val = any(v is not None for v in c_vals)
    if has_val or lbl:
        print(f"Row {r:2d} | Label: {str(lbl):25s} | Formula? {has_formula} | B={str(c_vals[0])[:20]} | K={str(c_vals[-1])[:20]}")
