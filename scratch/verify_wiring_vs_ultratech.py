import os
import sys
import openpyxl

sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from institutional_control_checks import apply_institutional_wiring

# Load UltraTech_fixed_v2 as the gold standard
ref_path = "UltraTech_fixed_v2.xlsx"
wb_ref = openpyxl.load_workbook(ref_path, data_only=False)

# Make a temporary copy of master_model_template.xlsx and apply wiring
test_path = "scratch/test_template_wired.xlsx"
wb_test = openpyxl.load_workbook("master_model_template.xlsx", data_only=False)
apply_institutional_wiring(wb_test)
wb_test.save(test_path)
wb_test.close()

wb_test = openpyxl.load_workbook(test_path, data_only=False)

print("=== VERIFYING WIRING AGAINST ULTRATECH_FIXED_V2 ===")

# 1. Control Sheet: Section 4
ws_c_ref = wb_ref['Control']
ws_c_test = wb_test['Control']
print("\n--- Control Sheet Rows 32-36 ---")
for r in range(32, 37):
    for c in ['B', 'C', 'D']:
        ref_v = ws_c_ref[f'{c}{r}'].value
        test_v = ws_c_test[f'{c}{r}'].value
        match = "MATCH" if str(ref_v) == str(test_v) else f"DIFF (ref='{ref_v}' vs test='{test_v}')"
        print(f"  {c}{r}: {match}")

# 2. Checks Sheet: Check 8 (Row 13) and Advisory Checks (Rows 17-21)
ws_k_ref = wb_ref['Checks']
ws_k_test = wb_test['Checks']
print("\n--- Checks Sheet Row 13 & Rows 17-21 ---")
for r in [13] + list(range(17, 22)):
    for col in ['B', 'C', 'D', 'E', 'F']:
        ref_v = ws_k_ref[f'{col}{r}'].value
        test_v = ws_k_test[f'{col}{r}'].value
        match = "MATCH" if str(ref_v) == str(test_v) else f"DIFF (ref='{ref_v}' vs test='{test_v}')"
        if match != "MATCH":
            print(f"  {col}{r}: {match}")
        else:
            print(f"  {col}{r}: MATCH ({test_v})")

# 3. DCF Sheet: Alternate 10Y DCF (Rows 58-83)
ws_d_ref = wb_ref['DCF']
ws_d_test = wb_test['DCF']
print("\n--- DCF Sheet Rows 58-83 ---")
dcf_diffs = []
for r in range(58, 84):
    for c_idx in range(2, 14):
        cl = openpyxl.utils.get_column_letter(c_idx)
        ref_v = ws_d_ref[f'{cl}{r}'].value
        test_v = ws_d_test[f'{cl}{r}'].value
        if str(ref_v) != str(test_v):
            dcf_diffs.append((f'{cl}{r}', ref_v, test_v))

if not dcf_diffs:
    print("  ALL CELLS MATCH in DCF Rows 58-83!")
else:
    print(f"  Found {len(dcf_diffs)} diffs in DCF:")
    for cell, ref_v, test_v in dcf_diffs[:10]:
        print(f"    {cell}: ref='{ref_v}' vs test='{test_v}'")

# 4. Intrinsic Valuation Sheet: Rows 67-71 and J55/J65
ws_iv_ref = wb_ref['Intrinsic Valuation']
ws_iv_test = wb_test['Intrinsic Valuation']
print("\n--- Intrinsic Valuation Rows 67-71 & J55/J65 ---")
iv_diffs = []
for cell in ['J55', 'J65'] + [f'{cl}{r}' for r in range(67, 72) for cl in ['B', 'H', 'I', 'J', 'K', 'L', 'M']]:
    ref_v = ws_iv_ref[cell].value
    test_v = ws_iv_test[cell].value
    if str(ref_v) != str(test_v):
        iv_diffs.append((cell, ref_v, test_v))

if not iv_diffs:
    print("  ALL CELLS MATCH in Intrinsic Valuation Organic Cross-Check & Labels!")
else:
    print(f"  Found {len(iv_diffs)} diffs in Intrinsic Valuation:")
    for cell, ref_v, test_v in iv_diffs[:10]:
        print(f"    {cell}: ref='{ref_v}' vs test='{test_v}'")

# 5. AI Valuation Summary
ws_ai_ref = wb_ref['AI Valuation Summary']
ws_ai_test = wb_test['AI Valuation Summary']
print("\n--- AI Valuation Summary key cells ---")
for cell in ['I4', 'I5', 'D10', 'D12', 'B41', 'C41', 'D41', 'B42', 'C42', 'D42']:
    ref_v = ws_ai_ref[cell].value
    test_v = ws_ai_test[cell].value
    match = "MATCH" if str(ref_v) == str(test_v) else f"DIFF (ref='{ref_v}' vs test='{test_v}')"
    print(f"  {cell}: {match}")

print("\nVerification script finished.")
