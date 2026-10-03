import sys, os
sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data, clean_num
import openpyxl

# Load target reference from Nestle India (2).xlsx
wb_ref = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=True)
ws_ref = wb_ref['Data Sheet']

ref_data = {}
for r in [17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 42, 43, 44, 45, 46, 47, 48, 49, 50, 57, 58, 59, 60, 61, 62, 63, 64, 65, 66, 67, 68, 69, 70, 72, 82, 83, 84, 85, 90, 93]:
    lbl = ws_ref.cell(r, 1).value
    vals = [ws_ref.cell(r, c).value for c in range(2, 12)]
    ref_data[r] = (lbl, vals)

print(f"Loaded {len(ref_data)} reference rows from Nestle India (2).xlsx Data Sheet.")
print("Sample - Row 93 (Adjusted Equity Shares in Cr):", ref_data[93])
print("Sample - Row 90 (PRICE):", ref_data[90])
print("Sample - Row 67 (Receivables):", ref_data[67])
print("Sample - Row 68 (Inventory):", ref_data[68])
print("Sample - Row 69 (Cash & Bank):", ref_data[69])
print("Sample - Row 70 (Shares):", ref_data[70])
