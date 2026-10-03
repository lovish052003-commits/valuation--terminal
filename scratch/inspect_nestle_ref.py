import openpyxl

wb = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=False)
wb_v = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=True)

print("=== 1. DATA SHEET META & KEY CELLS ===")
ws = wb['Data Sheet']
ws_v = wb_v['Data Sheet']
for r in [1, 6, 7, 8, 9, 16, 17, 28, 30, 31, 32, 56, 57, 58, 59, 67, 68, 69, 70, 71, 72, 73, 74, 75, 81, 82, 90, 93]:
    lbl = ws.cell(r, 1).value
    val_b = ws_v.cell(r, 2).value
    form_b = ws.cell(r, 2).value
    val_k = ws_v.cell(r, 11).value
    form_k = ws.cell(r, 11).value
    print(f"Row {r:<2} | {str(lbl):<25} | Col B: val={val_b}, form={form_b} | Col K: val={val_k}, form={form_k}")

print("\n=== 2. WACC SHEET ===")
ws = wb['WACC']
ws_v = wb_v['WACC']
for r in range(10, 42):
    lbl = ws.cell(r, 2).value or ws.cell(r, 3).value
    val = ws_v.cell(r, 4).value
    form = ws.cell(r, 4).value
    if val is not None or form is not None:
        print(f"Row {r:<2} | {str(lbl):<30} | D{r}: val={val} | form={form}")

print("\n=== 3. DCF SHEET ===")
ws = wb['DCF']
ws_v = wb_v['DCF']
for r in list(range(5, 28)) + list(range(33, 46)):
    for col_let in ['D', 'E', 'H', 'I']:
        val = ws_v[f"{col_let}{r}"].value
        form = ws[f"{col_let}{r}"].value
        if val is not None or form is not None:
            lbl = ws[f"C{r}"].value or ws[f"G{r}"].value or ''
            print(f"{col_let}{r:<2} | {str(lbl):<30} | val={val} | form={form}")

print("\n=== 4. INTRINSIC VALUATION SHEET ===")
ws = wb['Intrinsic Valuation']
ws_v = wb_v['Intrinsic Valuation']
for r in range(1, 20):
    lbl = ws.cell(r, 1).value
    val = ws_v.cell(r, 2).value
    form = ws.cell(r, 2).value
    if val is not None or form is not None:
        print(f"Row {r:<2} | {str(lbl):<30} | val={val} | form={form}")

print("\n=== 5. COMP_VALUATION SHEET ===")
ws = wb['Comp_Valuation']
ws_v = wb_v['Comp_Valuation']
for r in list(range(11, 22)) + list(range(30, 42)):
    peer_name = ws.cell(r, 2).value
    val_c = ws_v.cell(r, 3).value
    form_c = ws.cell(r, 3).value
    val_f = ws_v.cell(r, 6).value
    val_h = ws_v.cell(r, 8).value
    val_o = ws_v.cell(r, 15).value
    val_p = ws_v.cell(r, 16).value
    val_q = ws_v.cell(r, 17).value
    print(f"Row {r:<2} | {str(peer_name):<25} | C: {form_c} | EV/Rev: {val_o} | EV/EBITDA: {val_p} | P/E: {val_q}")
