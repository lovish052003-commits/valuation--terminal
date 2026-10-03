import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['Comp_Valuation']
for r in range(12, 18):
    print(f'Row {r}: B={ws[f"B{r}"].value} | C={ws[f"C{r}"].value}')

wb_data = openpyxl.load_workbook('ITC Model.xlsx', data_only=True)
ws_data = wb_data['Comp_Valuation']
for r in range(12, 18):
    print(f'Data Row {r}: B={ws_data[f"B{r}"].value} | C={ws_data[f"C{r}"].value}')
