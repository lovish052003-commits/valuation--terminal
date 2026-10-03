import openpyxl

wb = openpyxl.load_workbook('exports/SBIN_Valuation_Model.xlsx', data_only=False)
wb_data = openpyxl.load_workbook('exports/SBIN_Valuation_Model.xlsx', data_only=True)

ws_comp = wb['Comp_Valuation']
ws_comp_data = wb_data['Comp_Valuation']

print("=== COMP_VALUATION INSPECTION ===")
print("Headers in Row 10:")
for col in ['B','C','D','E','F','G','H','I','J','K','L','M','O','P','Q']:
    print(f"  {col}10: {ws_comp[f'{col}10'].value}")

print("\nPeers (Rows 12-20):")
for r in range(12, 21):
    b = ws_comp[f'B{r}'].value
    c = ws_comp[f'C{r}'].value
    if b or c:
        mcap = ws_comp_data[f'F{r}'].value
        pat = ws_comp_data[f'M{r}'].value
        pe_eval = ws_comp_data[f'Q{r}'].value
        pe_form = ws_comp[f'Q{r}'].value
        rev_form = ws_comp[f'I{r}'].value
        ebitda_form = ws_comp[f'J{r}'].value
        print(f"  Row {r}: Peer={c} | Mcap={mcap} | PAT={pat} | PE_val={pe_eval} | PE_form={pe_form}")
        print(f"          EV/Rev_form={rev_form} | EV/EBITDA_form={ebitda_form}")

print("\nBenchmark Stats (Rows 23-28):")
for r in range(23, 29):
    b = ws_comp[f'B{r}'].value
    q_form = ws_comp[f'Q{r}'].value
    q_val = ws_comp_data[f'Q{r}'].value
    print(f"  Row {r}: {b} | Formula={q_form} | Value={q_val}")

print("\nTarget Valuation (Rows 30-39):")
for r in [30, 32, 33, 34, 35, 37, 39]:
    b = ws_comp[f'B{r}'].value
    for col in ['I', 'J', 'O', 'P', 'Q']:
        form = ws_comp[f'{col}{r}'].value
        val = ws_comp_data[f'{col}{r}'].value
        print(f"  Row {r} Col {col}: label={b} | Formula={form} | Value={val}")

print("\nColumn Widths:")
for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'O', 'P', 'Q']:
    print(f"  Col {col}: width={ws_comp.column_dimensions[col].width}")

print("\n=== DATA SHEET INSPECTION ===")
ws_ds = wb['Data Sheet']
ws_ds_data = wb_data['Data Sheet']

years = [ws_ds.cell(row=16, column=c).value for c in range(3, 13)]
print("Years:", years)

net_profits = [ws_ds.cell(row=30, column=c).value for c in range(3, 13)]
print("Row 30 Net Profit:", net_profits)

dividends = [ws_ds.cell(row=31, column=c).value for c in range(3, 13)]
print("Row 31 Dividend Amount:", dividends)

sales = [ws_ds.cell(row=17, column=c).value for c in range(3, 13)]
other_inc = [ws_ds.cell(row=25, column=c).value for c in range(3, 13)]
depr = [ws_ds.cell(row=26, column=c).value for c in range(3, 13)]
interest = [ws_ds.cell(row=27, column=c).value for c in range(3, 13)]
pbt = [ws_ds.cell(row=28, column=c).value for c in range(3, 13)]
tax = [ws_ds.cell(row=29, column=c).value for c in range(3, 13)]

print("\nOperating Expense Breakdown (Rows 18-24):")
for r in range(18, 25):
    line_name = ws_ds.cell(row=r, column=1).value
    row_vals = [ws_ds.cell(row=r, column=c).value for c in range(3, 13)]
    print(f"  Row {r:2d} ({line_name}): {row_vals}")

print("\nExpense Sum vs Implied Total Expense:")
for idx, yr in enumerate(years):
    exp_sum = sum(ws_ds.cell(row=r, column=3+idx).value or 0 for r in range(18, 25))
    s = sales[idx] or 0
    oi = other_inc[idx] or 0
    d = depr[idx] or 0
    i = interest[idx] or 0
    pb = pbt[idx] or 0
    implied_exp = s + oi - d - i - pb
    print(f"  {yr}: Sum(18..24)={exp_sum:,.2f} vs Implied={implied_exp:,.2f} | Diff={abs(exp_sum - implied_exp):.4f}")
