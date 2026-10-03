import os, sys, shutil
import openpyxl
import pandas as pd

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from screener_client import fetch_company_data, clean_num
from scratch.test_chart_prices import get_historical_prices

print("1. Fetching NESTLEIND from Screener...")
data = fetch_company_data('NESTLEIND')

# Copy template to test file
test_dest = 'exports/TEST_NESTLE_DATASHEET.xlsx'
shutil.copyfile('ITC Model.xlsx', test_dest)

wb = openpyxl.load_workbook(test_dest)
ws_data = wb['Data Sheet']

tables = data.get('tables', {})
pl_df = tables.get('profit-loss')
bs_df = tables.get('balance-sheet')
cf_df = tables.get('cash-flow')
qtr_df = tables.get('quarters')
schedules = data.get('schedules', {})

# Meta
ws_data['B1'] = data['company_name']
ws_data['B6'] = "=IF(B9>0, B9/B8, 0)"
ws_data['B7'] = clean_num(data.get('face_value', 1))
ws_data['B8'] = clean_num(data.get('current_price', 0))
ws_data['B9'] = clean_num(data.get('market_cap_cr', 0))

# Get the 10 annual period columns (EXCLUDE TTM!)
bs_cols = [c for c in bs_df.columns if c != 'Metric'][-10:] if bs_df is not None else []
pl_cols = [c for c in pl_df.columns if c not in ('Metric', 'TTM')][-10:] if pl_df is not None else []
cf_cols = [c for c in cf_df.columns if c != 'Metric'][-10:] if cf_df is not None else []

print(f"Annual periods (BS): {bs_cols}")
print(f"Annual periods (P&L): {pl_cols}")

# Helper to write rows
def write_row_by_keywords(row_idx, df, keywords, cols):
    if df is None or df.empty:
        return
    for kw in keywords:
        m = df[df['Metric'].str.contains(kw, case=False, na=False)]
        if not m.empty:
            r = m.iloc[0]
            for offset, col_name in enumerate(cols):
                c_idx = 2 + offset
                val = clean_num(r.get(col_name, 0))
                ws_data.cell(row=row_idx, column=c_idx, value=val)
            break

def write_dates(row_idx, cols):
    for offset, col_name in enumerate(cols):
        c_idx = 2 + offset
        ws_data.cell(row=row_idx, column=c_idx, value=str(col_name))

# 1. Profit & Loss
write_dates(16, pl_cols)
write_row_by_keywords(17, pl_df, ['^Sales', 'Revenue'], pl_cols)
write_row_by_keywords(18, pl_df, ['Raw Material Cost', 'Material'], pl_cols)
write_row_by_keywords(19, pl_df, ['Change in Inventory'], pl_cols)
write_row_by_keywords(20, pl_df, ['Power and Fuel', 'Power'], pl_cols)
write_row_by_keywords(21, pl_df, ['Other Mfr. Exp', 'Manufacturing'], pl_cols)
write_row_by_keywords(22, pl_df, ['Employee Cost'], pl_cols)
write_row_by_keywords(23, pl_df, ['Selling and admin', 'Selling'], pl_cols)
write_row_by_keywords(24, pl_df, ['Other Expenses'], pl_cols)
write_row_by_keywords(25, pl_df, ['Other Income'], pl_cols)
write_row_by_keywords(26, pl_df, ['Depreciation'], pl_cols)
write_row_by_keywords(27, pl_df, ['Interest'], pl_cols)
write_row_by_keywords(28, pl_df, ['Profit before tax', 'PBT'], pl_cols)
write_row_by_keywords(29, pl_df, ['Tax %', 'Tax'], pl_cols)
write_row_by_keywords(30, pl_df, ['Net profit', 'PAT'], pl_cols)
write_row_by_keywords(31, pl_df, ['Dividend Amount', 'Dividend Payout'], pl_cols)

# Row 32 EBITDA formula =SUM(B26:B28)
for offset in range(len(pl_cols)):
    c_idx = 2 + offset
    col_let = openpyxl.utils.get_column_letter(c_idx)
    ws_data.cell(row=32, column=c_idx, value=f"=SUM({col_let}26:{col_let}28)")

# 2. Quarters (Rows 41 to 50)
if qtr_df is not None and not qtr_df.empty:
    qtr_cols = [c for c in qtr_df.columns if c != 'Metric'][-10:]
    write_dates(41, qtr_cols)
    write_row_by_keywords(42, qtr_df, ['Sales', 'Revenue'], qtr_cols)
    write_row_by_keywords(43, qtr_df, ['Expenses'], qtr_cols)
    write_row_by_keywords(44, qtr_df, ['Other Income'], qtr_cols)
    write_row_by_keywords(45, qtr_df, ['Depreciation'], qtr_cols)
    write_row_by_keywords(46, qtr_df, ['Interest'], qtr_cols)
    write_row_by_keywords(47, qtr_df, ['Profit before tax', 'PBT'], qtr_cols)
    write_row_by_keywords(48, qtr_df, ['Tax'], qtr_cols)
    write_row_by_keywords(49, qtr_df, ['Net profit', 'PAT'], qtr_cols)
    write_row_by_keywords(50, qtr_df, ['Operating Profit'], qtr_cols)

# 3. Balance Sheet (Rows 56 to 66)
write_dates(56, bs_cols)
write_row_by_keywords(57, bs_df, ['Equity Capital', 'Share Capital'], bs_cols)
write_row_by_keywords(58, bs_df, ['Reserves'], bs_cols)
write_row_by_keywords(59, bs_df, ['Borrowings', 'Total Debt'], bs_cols)
write_row_by_keywords(60, bs_df, ['Other Liabilities'], bs_cols)
write_row_by_keywords(61, bs_df, ['Total Liabilities'], bs_cols)
write_row_by_keywords(62, bs_df, ['Fixed Assets', 'Net Block'], bs_cols)
write_row_by_keywords(63, bs_df, ['CWIP'], bs_cols)
write_row_by_keywords(64, bs_df, ['Investments'], bs_cols)
write_row_by_keywords(65, bs_df, ['Other Assets'], bs_cols)
write_row_by_keywords(66, bs_df, ['Total Assets'], bs_cols)

# 4. Working Capital Breakdown & Shares (Rows 67 to 74)
oa_sch = schedules.get('Other Assets', {})
fv = clean_num(data.get('face_value', 1.0))
if fv == 0: fv = 1.0

for offset, col_name in enumerate(bs_cols):
    c_idx = 2 + offset
    col_let = openpyxl.utils.get_column_letter(c_idx)
    
    # Row 67: Receivables
    rec_val = clean_num(oa_sch.get('Trade receivables', {}).get(col_name, 0))
    ws_data.cell(row=67, column=c_idx, value=rec_val)
    
    # Row 68: Inventory
    inv_val = clean_num(oa_sch.get('Inventories', {}).get(col_name, 0))
    ws_data.cell(row=68, column=c_idx, value=inv_val)
    
    # Row 69: Cash & Bank
    cash_val = clean_num(oa_sch.get('Cash Equivalents', {}).get(col_name, 0))
    ws_data.cell(row=69, column=c_idx, value=cash_val)
    
    # Row 57 eq capital in Cr
    eq_cap = clean_num(ws_data.cell(row=57, column=c_idx).value or 0)
    
    # Row 72: Face value
    ws_data.cell(row=72, column=c_idx, value=fv)
    
    # Row 70: No. of Equity Shares = (Eq Capital in Cr * 10^7) / Face value
    num_shares = (eq_cap * 10000000.0) / fv if fv > 0 else 0
    ws_data.cell(row=70, column=c_idx, value=round(num_shares))
    
    # Row 71: Bonus shares (None)
    ws_data.cell(row=71, column=c_idx, value=None)
    
    # Row 74: Other Assets = =B65-SUM(B67:B69)
    ws_data.cell(row=74, column=c_idx, value=f"={col_let}65-SUM({col_let}67:{col_let}69)")

# 5. Cash Flow (Rows 81 to 85)
write_dates(81, cf_cols)
write_row_by_keywords(82, cf_df, ['Cash from Operating Activity'], cf_cols)
write_row_by_keywords(83, cf_df, ['Cash from Investing Activity'], cf_cols)
write_row_by_keywords(84, cf_df, ['Cash from Financing Activity'], cf_cols)
write_row_by_keywords(85, cf_df, ['Net Cash Flow'], cf_cols)

# 6. Historical Prices (Row 90) & Adjusted Equity Shares in Cr (Row 93)
cid = data.get('company_id')
chart_prices = get_historical_prices(cid, bs_cols)
cmp = clean_num(data.get('current_price', 0))
shares_cr = clean_num(data.get('shares_in_cr', 0))
if shares_cr == 0 and fv > 0:
    latest_eq = clean_num(ws_data.cell(row=57, column=11).value or 0)
    shares_cr = latest_eq / fv

for offset, col_name in enumerate(bs_cols):
    c_idx = 2 + offset
    # Price
    p_val = chart_prices.get(col_name)
    if not p_val and offset == len(bs_cols) - 1:
        p_val = cmp
    if not p_val:
        p_val = cmp
    ws_data.cell(row=90, column=c_idx, value=p_val)
    
    # Row 93: Adjusted Equity Shares in Cr
    ws_data.cell(row=93, column=c_idx, value=round(shares_cr, 2))

wb.save(test_dest)
print("Saved to", test_dest)

# Now compare with reference
wb_ref = openpyxl.load_workbook(r'c:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=True)
ws_ref = wb_ref['Data Sheet']

wb_test = openpyxl.load_workbook(test_dest, data_only=True)
ws_t = wb_test['Data Sheet']

print("\n=== VERIFICATION AGAINST NESTLE INDIA (2).XLSX ===")
test_rows = [16, 17, 22, 26, 27, 28, 30, 41, 42, 49, 56, 57, 58, 62, 65, 67, 68, 69, 70, 72, 81, 82, 90, 93]
for r in test_rows:
    lbl = ws_ref.cell(r, 1).value or ws_t.cell(r, 1).value
    vr_latest = ws_ref.cell(r, 11).value
    vt_latest = ws_t.cell(r, 11).value
    vr_first = ws_ref.cell(r, 2).value
    vt_first = ws_t.cell(r, 2).value
    print(f"Row {r:2d} ({str(lbl):25s}) | Ref Latest: {str(vr_latest):15s} | Our Latest: {str(vt_latest):15s} | Ref 1st: {str(vr_first):12s} | Our 1st: {str(vt_first)}")
