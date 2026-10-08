import openpyxl

wb_f = openpyxl.load_workbook("FORCEMOT_Valuation_Model_FIXED.xlsx", data_only=False)
wb_o = openpyxl.load_workbook(r"exports\FORCEMOT_Valuation_Model.xlsx", data_only=False)

already_analyzed = {'Control', 'Checks', 'AI Valuation Summary', 'DCF', 'WACC', 'Beta-Regression', 'Ratio Analysis', 'Forecasting', 'Raw Data'}

for s in wb_f.sheetnames:
    if s in already_analyzed:
        continue
    if s not in wb_o.sheetnames:
        print(f"Sheet {s} is ONLY in FIXED!")
        continue
    ws_f = wb_f[s]
    ws_o = wb_o[s]
    diffs = []
    max_r = max(ws_f.max_row, ws_o.max_row)
    max_c = max(ws_f.max_column, ws_o.max_column)
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            vf = ws_f.cell(r, c).value
            vo = ws_o.cell(r, c).value
            if vf != vo:
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                diffs.append((coord, vo, vf))
    if diffs:
        print(f"\nSheet {s} has {len(diffs)} diffs:")
        for coord, vo, vf in diffs[:20]:
            print(f"  {coord}: ORIG={vo!r} -> FIXED={vf!r}")
    else:
        print(f"Sheet {s}: IDENTICAL (0 diffs)")
