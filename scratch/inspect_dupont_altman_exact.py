import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)

for sheet_name in ['Dupont Analysis', "Altman's Z Score"]:
    if sheet_name not in wb.sheetnames:
        continue
    ws = wb[sheet_name]
    print(f"\n==================== {sheet_name} ====================")
    for r in range(1, 46):
        b_val = ws.cell(r, 2).value
        if b_val is not None:
            print(f"Row {r:2d} [Col B]: {repr(b_val)[:120]}")
