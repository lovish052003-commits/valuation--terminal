import openpyxl

paths = [
    r"C:\Users\LENOVO\Downloads\test\SBIN_Valuation_Model (2).xlsx",
    r"C:\Users\LENOVO\Downloads\test\SBIN_Valuation_Model.xlsx",
    r"C:\Users\LENOVO\Downloads\Advance Financial Project\exports\SBIN_Valuation_Model.xlsx"
]

for p in paths:
    print(f"\n=======================================================")
    print(f"INSPECTING: {p}")
    print(f"=======================================================")
    try:
        wb_f = openpyxl.load_workbook(p, data_only=False)
        wb_v = openpyxl.load_workbook(p, data_only=True)
    except Exception as e:
        print(f"Failed to open {p}: {e}")
        continue

    print(f"Sheets: {wb_f.sheetnames}")

    # Inspect Data Sheet around row 32
    if 'Data Sheet' in wb_f.sheetnames:
        ws_f = wb_f['Data Sheet']
        ws_v = wb_v['Data Sheet']
        print("\n--- Data Sheet Rows 20 to 35 ---")
        for r in range(20, 36):
            b_val = ws_v.cell(r, 2).value
            k_formula = ws_f.cell(r, 11).value
            k_val = ws_v.cell(r, 11).value
            print(f"Row {r:2d} | Col B: {str(b_val)[:30]:30s} | Col K Form: {str(k_formula)[:30]:30s} | Col K Val: {k_val}")

        print("\n--- Data Sheet Rows 65 to 75 ---")
        for r in range(65, 76):
            b_val = ws_v.cell(r, 2).value
            k_formula = ws_f.cell(r, 11).value
            k_val = ws_v.cell(r, 11).value
            print(f"Row {r:2d} | Col B: {str(b_val)[:30]:30s} | Col K Form: {str(k_formula)[:30]:30s} | Col K Val: {k_val}")

    # Inspect DCF
    if 'DCF' in wb_f.sheetnames:
        ws_f = wb_f['DCF']
        ws_v = wb_v['DCF']
        print("\n--- DCF Rows 7 to 15 (EBIT / D&A / NOPAT / FCFF) ---")
        for r in range(7, 16):
            a_val = ws_v.cell(r, 1).value
            c_val = ws_v.cell(r, 3).value
            d_val = ws_v.cell(r, 4).value
            h_f = ws_f.cell(r, 8).value
            h_v = ws_v.cell(r, 8).value
            print(f"Row {r:2d} | A: {str(a_val)[:25]:25s} | C: {str(c_val)[:15]} | H Form: {str(h_f)[:35]:35s} | H Val: {h_v}")

        print("\n--- DCF Rows 32 to 46 (Bridge) ---")
        for r in range(32, 47):
            c_val = ws_v.cell(r, 3).value
            d_f = ws_f.cell(r, 4).value
            d_v = ws_v.cell(r, 4).value
            print(f"Row {r:2d} | Label: {str(c_val)[:35]:35s} | Form: {str(d_f)[:35]:35s} | Val: {d_v}")

    # Inspect AI Valuation Summary
    if 'AI Valuation Summary' in wb_f.sheetnames:
        ws_f = wb_f['AI Valuation Summary']
        ws_v = wb_v['AI Valuation Summary']
        print("\n--- AI Valuation Summary Rows 4 to 12 ---")
        for r in range(4, 13):
            row_f = [ws_f.cell(r, c).value for c in range(1, 9)]
            row_v = [ws_v.cell(r, c).value for c in range(1, 9)]
            print(f"Row {r:2d} Forms: {row_f}")
            print(f"Row {r:2d} Vals : {row_v}")

        print("\n--- AI Valuation Summary Rows 15 to 30 ---")
        for r in range(15, 31):
            val_a = ws_v.cell(r, 1).value
            val_b = ws_v.cell(r, 2).value
            val_c = ws_v.cell(r, 3).value
            if any([val_a, val_b, val_c]):
                print(f"Row {r:2d} | A: {str(val_a)[:30]} | B: {str(val_b)[:30]} | C: {str(val_c)[:40]}")

    # Inspect Comp_Valuation
    if 'Comp_Valuation' in wb_f.sheetnames:
        ws_f = wb_f['Comp_Valuation']
        ws_v = wb_v['Comp_Valuation']
        print("\n--- Comp_Valuation Rows 10 to 26 ---")
        for r in range(10, 27):
            d_name = ws_v.cell(r, 4).value
            ev_rev = ws_v.cell(r, 15).value
            ev_ebitda = ws_v.cell(r, 16).value
            pe = ws_v.cell(r, 17).value
            print(f"Row {r:2d} | Peer: {str(d_name)[:20]:20s} | EV/Rev: {ev_rev} | EV/EBITDA: {ev_ebitda} | P/E: {pe}")
        print("\n--- Comp_Valuation Rows 30 to 42 ---")
        for r in range(30, 43):
            n_val = ws_v.cell(r, 14).value
            o_v = ws_v.cell(r, 15).value
            p_v = ws_v.cell(r, 16).value
            q_v = ws_v.cell(r, 17).value
            o_f = ws_f.cell(r, 15).value
            print(f"Row {r:2d} | Label: {str(n_val)[:30]:30s} | O: {o_v} ({str(o_f)[:20]}) | P: {p_v} | Q: {q_v}")
