import sys, os
sys.path.insert(0, os.path.abspath('.'))
import screener_client
import valuation_engine
import excel_exporter
import openpyxl

print("1. Fetching SUNPHARMA data...")
screener_data = screener_client.fetch_company_data('SUNPHARMA')
print("2. Calculating valuation...")
val_result = valuation_engine.calculate_valuation(screener_data)

out_file = os.path.abspath('scratch/test_full_sunpharma_fix.xlsx')
if os.path.exists(out_file):
    os.remove(out_file)

print("3. Exporting model via OpenPyXL / exporter...")
# Force OpenPyXL test directly
shutil_path = 'scratch/test_full_sunpharma_fix.xlsx'
import shutil
shutil.copyfile('ITC Model.xlsx', shutil_path)

wb = openpyxl.load_workbook(shutil_path)
excel_exporter.populate_data_sheet_openpyxl(wb, screener_data)
excel_exporter.populate_raw_fs_sheet_openpyxl(wb, screener_data, val_result)
excel_exporter.populate_cash_flow_statement_sheet_openpyxl(wb, screener_data)
excel_exporter.populate_raw_data_prices_openpyxl(wb, screener_data)
excel_exporter.update_comp_valuation_sheet_openpyxl(wb, screener_data, val_result)
excel_exporter.update_wacc_raw_data_openpyxl(wb, screener_data, val_result)
excel_exporter.populate_dupont_altman_sheets_openpyxl(wb, screener_data)
excel_exporter.populate_ai_summary_sheet_openpyxl(wb, screener_data, val_result)

wb.calculation.fullCalcOnLoad = True
wb.save(shutil_path)
wb.close()

# Apply patch_valuation_workbook
reconciled = excel_exporter.patch_valuation_workbook(shutil_path, screener_data, val_result)
print("Reconciled path:", reconciled)

# Verify Data Sheet rows
wb_check = openpyxl.load_workbook(reconciled, data_only=False)
ws = wb_check['Data Sheet']
print("\n=== DATA SHEET PROFIT & LOSS ===")
for r in range(16, 26):
    lbl = ws.cell(r, 1).value
    vals = [ws.cell(r, c).value for c in range(2, 12)]
    print(f"Row {r:2d} | {str(lbl):25s} | {vals}")

print("\n=== RAW FS ROW 56 (TARGET) ===")
ws_raw = wb_check['Raw FS']
for col_idx in [12, 13, 14, 44, 45, 46, 47, 48, 49, 52, 53]:
    let = openpyxl.utils.get_column_letter(col_idx)
    print(f"Col {let}56: {ws_raw.cell(56, col_idx).value}")

print("\n=== COMP_VALUATION ROWS 32-34 ===")
ws_comp = wb_check['Comp_Valuation']
for r in [32, 33, 34, 35, 37, 39]:
    lbl = ws_comp[f'B{r}'].value
    print(f"Row {r:2d} | {str(lbl):25s} | O={ws_comp[f'O{r}'].value} | P={ws_comp[f'P{r}'].value} | Q={ws_comp[f'Q{r}'].value}")
