import os, sys, gc
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

def test_company_without_download(symbol):
    print(f"\n--- Testing Dynamic Scraping & Export for {symbol} (No local file) ---")
    data = fetch_company_data(symbol)
    print(f"Company: {data.get('company_name')}")
    print(f"CMP: {data.get('current_price')}")
    print(f"Shares in Cr: {data.get('shares_in_cr'):.2f}")
    print(f"Historical price points: {len(data.get('historical_prices', []))}")

    val = calculate_valuation(data)
    out_path = export_valuation_model(data, val, "### Test Dynamic Export")
    print(f"Exported to: {out_path}")

    wb = openpyxl.load_workbook(out_path, data_only=True)
    ws = wb['Data Sheet']

    print("Checking key cells in exported Data Sheet:")
    print("Row 1 (Name):", ws.cell(1, 2).value)
    print("Row 93 (Adj Shares Cr):", [ws.cell(93, c).value for c in range(2, 12)])
    print("Row 70 (No of Shares):", [ws.cell(70, c).value for c in range(2, 12)][:3])
    print("Row 67 (Receivables):", [ws.cell(67, c).value for c in range(2, 12)][:3])
    print("Row 68 (Inventory):", [ws.cell(68, c).value for c in range(2, 12)][:3])
    print("Row 69 (Cash & Bank):", [ws.cell(69, c).value for c in range(2, 12)][:3])
    print("Row 90 (Price):", [ws.cell(90, c).value for c in range(2, 12)][:3])
    wb.close()

if __name__ == '__main__':
    test_company_without_download('HINDUNILVR')
