import openpyxl, os

screener_files = [
    r'c:\Users\LENOVO\Downloads\Nestle India (2).xlsx',
    r'c:\Users\LENOVO\Downloads\Britannia Inds.xlsx',
    r'c:\Users\LENOVO\Downloads\ITC Model.xlsx',
    r'c:\Users\LENOVO\Downloads\Hind. Unilever (1).xlsx',
    r'c:\Users\LENOVO\Downloads\Tata Steel.xlsx'
]

for f in screener_files:
    if not os.path.exists(f):
        continue
    wb = openpyxl.load_workbook(f, data_only=True)
    if 'Data Sheet' in wb.sheetnames:
        ws = wb['Data Sheet']
        cname = ws.cell(1, 2).value
        r57 = [ws.cell(57, c).value for c in range(2, 12)] # equity share capital
        r70 = [ws.cell(70, c).value for c in range(2, 12)] # no of shares
        r72 = [ws.cell(72, c).value for c in range(2, 12)] # face value
        r93 = [ws.cell(93, c).value for c in range(2, 12)] # adjusted shares
        r90 = [ws.cell(90, c).value for c in range(2, 12)] # prices
        print(f"\n==========================================")
        print(f"Company: {cname} ({os.path.basename(f)})")
        print(f"  Eq Capital (Row 57): {r57[-3:]}")
        print(f"  Face Value (Row 72): {r72[-3:]}")
        print(f"  No Shares  (Row 70): {r70[-3:]}")
        print(f"  Adj Shares (Row 93): {r93[-3:]}")
        print(f"  Prices     (Row 90): {r90[-3:]}")
