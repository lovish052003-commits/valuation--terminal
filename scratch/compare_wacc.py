import openpyxl

wb_r = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=False)
wb_g = openpyxl.load_workbook('exports/NESTLEIND_Valuation_Model.xlsx', data_only=False)

ws_r = wb_r['WACC']
ws_g = wb_g['WACC']

wb_rd = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=True)
wb_gd = openpyxl.load_workbook('exports/NESTLEIND_Valuation_Model.xlsx', data_only=True)
ws_rd = wb_rd['WACC']
ws_gd = wb_gd['WACC']

print("=== WACC Sheet Comparison ===")
for r in range(12, 48):
    for c in range(1, 14):
        fr = str(ws_r.cell(r, c).value or '')
        fg = str(ws_g.cell(r, c).value or '')
        vr = str(ws_rd.cell(r, c).value or '')
        vg = str(ws_gd.cell(r, c).value or '')
        if fr != fg or vr != vg:
            coord = f"{openpyxl.utils.get_column_letter(c)}{r}"
            print(f"Cell {coord:<4}: REF={fr} (val={vr[:12]}) | GEN={fg} (val={vg[:12]})")
