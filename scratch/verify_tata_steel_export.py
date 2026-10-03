import os
import sys
sys.path.insert(0, os.path.abspath('.'))
import openpyxl

from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

print("1. Fetching Screener data for TATASTEEL...")
screener_data = fetch_company_data('TATASTEEL')
print("   Company:", screener_data.get('company_name'))

print("2. Running valuation engine...")
val_res = calculate_valuation(screener_data)

print("3. Exporting full Excel model with new Comp_Valuation logic...")
out_path = export_valuation_model(screener_data, val_res, report_markdown="")
print("   Exported path:", out_path)

print("4. Inspecting Comp_Valuation sheet...")
wb = openpyxl.load_workbook(out_path, data_only=False)
ws = wb['Comp_Valuation']

print("\n--- Rows 12 to 21 (Peers Table) ---")
for r in range(12, 22):
    b = ws[f'B{r}'].value
    d = ws[f'D{r}'].value
    e = ws[f'E{r}'].value
    g = ws[f'G{r}'].value
    k = ws[f'K{r}'].value
    o = ws[f'O{r}'].value
    p = ws[f'P{r}'].value
    q = ws[f'Q{r}'].value
    print(f"  Row {r}: Name={b} | Price={d} | Shares={e} | NetDebt={g} | Rev={k} | P={p}")

print("\n--- Rows 23 to 28 (Peer Statistics) ---")
for r in range(23, 29):
    lbl = ws[f'B{r}'].value
    o = ws[f'O{r}'].value
    p = ws[f'P{r}'].value
    q = ws[f'Q{r}'].value
    print(f"  Row {r} ({lbl}): P={p}")

print("\n--- Rows 30 to 39 (Comparable Valuation) ---")
for r in range(30, 40):
    lbl = ws[f'B{r}'].value
    o = ws[f'O{r}'].value
    p = ws[f'P{r}'].value
    q = ws[f'Q{r}'].value
    print(f"  Row {r} ({lbl}): O={o} | P={p} | Q={q}")

wb_eval = openpyxl.load_workbook(out_path, data_only=True)
ws_eval = wb_eval['Comp_Valuation']
print("\n--- Evaluated Values ---")
print("  Row 12 Company:", ws_eval['B12'].value)
print("  Row 13 Company:", ws_eval['B13'].value)
print("  P27 (75th percentile EV/EBITDA):", ws_eval['P27'].value)
print("  P26 (Average EV/EBITDA):", ws_eval['P26'].value)
print("  O32 (Implied EV):", ws_eval['O32'].value)
print("  O33 (Net Debt):", ws_eval['O33'].value)
print("  O35 (Shares):", ws_eval['O35'].value)
print("  O37 (Implied Price/Share):", ws_eval['O37'].value)
print("  P37 (Implied Price/Share):", ws_eval['P37'].value)
print("  Q37 (Implied Price/Share):", ws_eval['Q37'].value)
print("  O39 (Verdict):", ws_eval['O39'].value)

print("\nVerification completed!")
