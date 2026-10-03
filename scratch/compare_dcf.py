import openpyxl

wb_ref = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=False)
wb_gen = openpyxl.load_workbook('exports/NESTLEIND_Valuation_Model.xlsx', data_only=False)

ws_r = wb_ref['DCF']
ws_g = wb_gen['DCF']

wb_ref_d = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=True)
wb_gen_d = openpyxl.load_workbook('exports/NESTLEIND_Valuation_Model.xlsx', data_only=True)
ws_rd = wb_ref_d['DCF']
ws_gd = wb_gen_d['DCF']

print(f"{'Row':<4} | {'REF Formula (D)':<30} | {'REF Val':<15} | {'GEN Formula (D)':<30} | {'GEN Val':<15}")
print('-' * 105)
for r in range(6, 48):
    lb_r = str(ws_r.cell(r, 2).value or '')
    fd_r = str(ws_r.cell(r, 4).value or '')
    vd_r = str(ws_rd.cell(r, 4).value or '')
    
    lb_g = str(ws_g.cell(r, 2).value or '')
    fd_g = str(ws_g.cell(r, 4).value or '')
    vd_g = str(ws_gd.cell(r, 4).value or '')
    
    if lb_r or fd_r or lb_g or fd_g:
        print(f"R{r:02d} ({lb_r[:15]} / {lb_g[:15]}):")
        print(f"     REF: D={fd_r:<25} val={vd_r}")
        print(f"     GEN: D={fd_g:<25} val={vd_g}")
        if ws_r.cell(r, 8).value or ws_g.cell(r, 8).value:
            fh_r = str(ws_r.cell(r, 8).value or '')
            vh_r = str(ws_rd.cell(r, 8).value or '')
            fh_g = str(ws_g.cell(r, 8).value or '')
            vh_g = str(ws_gd.cell(r, 8).value or '')
            print(f"     REF Col H: {fh_r} (val={vh_r})")
            print(f"     GEN Col H: {fh_g} (val={vh_g})")
