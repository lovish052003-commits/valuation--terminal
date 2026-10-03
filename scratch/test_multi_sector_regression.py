import os
import sys
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
sys.stdout.reconfigure(encoding='utf-8')

from screener_client import fetch_company_data
import valuation_engine
from excel_exporter import export_valuation_model

TEST_TICKERS = [
    ("ADANIENT", "Adani Enterprises (Conglomerate / Gold Standard Reference)"),
    ("OIL", "Oil India (Oil & Gas / Energy)"),
    ("SUNPHARMA", "Sun Pharma (Pharmaceuticals)"),
    ("TCS", "Tata Consultancy Services (IT Services)")
]

def run_sector_regression():
    results = {}
    for ticker, desc in TEST_TICKERS:
        print(f"\n{'='*80}")
        print(f"REGRESSION TEST: {ticker} — {desc}")
        print(f"{'='*80}")

        s_data = fetch_company_data(ticker)
        cmp_price = s_data.get('current_price')
        sector = s_data.get('sector')
        print(f"Company: {s_data.get('company_name')} | CMP: Rs. {cmp_price} | Sector: {sector}")

        val_res = valuation_engine.calculate_valuation(s_data)
        print(f"Engine Intrinsic Value: Rs. {val_res.get('intrinsic_value')} | Method: {val_res.get('selected_methodology')}")

        out_path = export_valuation_model(s_data, val_res, report_markdown="")
        print(f"Workbook Generated at: {out_path}")

        # Audit generated workbook
        wb_data = openpyxl.load_workbook(out_path, data_only=True)
        ws_ai = wb_data['AI Valuation Summary']
        ws_dcf = wb_data['DCF']

        ai_cmp = ws_ai['A5'].value
        ai_iv = ws_ai['B5'].value
        ai_mos = ws_ai['C5'].value
        ai_vrd = ws_ai['D5'].value
        ai_wacc = ws_ai['E5'].value
        ai_az = ws_ai['F5'].value
        ai_roe = ws_ai['G5'].value

        dcf_iv = ws_dcf['D42'].value
        dcf_wacc = ws_dcf['D20'].value
        dcf_fcff_y1 = ws_dcf['I16'].value

        # Check for error tokens
        error_tokens = ['#VALUE!', '#REF!', '#DIV/0!', '#NAME?', '#NUM!']
        errors_found = []
        for s_name in ['AI Valuation Summary', 'Comp_Valuation', 'DCF', 'Intrinsic Valuation', 'Data Sheet']:
            if s_name in wb_data.sheetnames:
                ws = wb_data[s_name]
                for row in ws.iter_rows(values_only=True):
                    for cell in row:
                        if isinstance(cell, str) and any(err in cell for err in error_tokens):
                            errors_found.append((s_name, cell))

        wb_data.close()

        print(f"Evaluated KPIs:")
        print(f"  CMP: Rs. {ai_cmp}")
        print(f"  Intrinsic Value: Rs. {ai_iv}")
        print(f"  Margin of Safety: {ai_mos}")
        print(f"  Valuation Gap: {ai_vrd}")
        print(f"  WACC: {ai_wacc}")
        print(f"  Altman Z: {ai_az}")
        print(f"  DuPont ROE: {ai_roe}")
        print(f"  DCF Year 1 PV of FCFF: {dcf_fcff_y1}")
        print(f"  DCF Intrinsic per share: {dcf_iv}")
        print(f"  Error tokens in workbook: {len(errors_found)}")

        if errors_found:
            print(f"  [FAIL] Error tokens found: {errors_found[:5]}")
            results[ticker] = False
        else:
            print(f"  [PASS] ZERO error tokens across all key sheets!")
            results[ticker] = True

    print(f"\n{'='*80}")
    print("FINAL MULTI-SECTOR REGRESSION SUMMARY")
    print(f"{'='*80}")
    for ticker, passed in results.items():
        print(f"{ticker}: {'PASSED' if passed else 'FAILED'}")

if __name__ == '__main__':
    run_sector_regression()
