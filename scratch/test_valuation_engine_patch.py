import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import screener_client

data = screener_client.fetch_company_data('ADANIENT')
ticker = data.get('ticker')
excel_file = f'exports/{ticker}_Valuation_Model.xlsx'

if os.path.exists(excel_file):
    import openpyxl
    wb = openpyxl.load_workbook(excel_file, data_only=True)
    if 'DCF' in wb.sheetnames:
        ws = wb['DCF']
        d42 = ws['D42'].value
        d35 = ws['D35'].value
        d39 = ws['D39'].value
        d37 = ws['D37'].value
        d38 = ws['D38'].value
        d40 = ws['D40'].value
        d20 = ws['D20'].value
        d18 = ws['D18'].value
        d19 = ws['D19'].value
        d44 = ws['D44'].value
        print("Found Excel Model:")
        print(f"  DCF D42 (Intrinsic): {d42}")
        print(f"  DCF D35 (EV): {d35}")
        print(f"  DCF D39 (Equity): {d39}")
        print(f"  DCF D37 (Cash): {d37}")
        print(f"  DCF D38 (Debt): {d38}")
        print(f"  DCF D40 (Shares): {d40}")
        print(f"  DCF D20 (WACC): {d20}")
