import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['WACC']
print("=== WACC Sheet Key Cells ===")
cells = ['C34', 'C35', 'C36', 'D34', 'D35', 'E26', 'E27', 'E28', 'K26', 'K27', 'K28', 'K29', 'K33', 'K34', 'K35', 'K36', 'K40', 'K41', 'K43', 'K44', 'K46']
for c in cells:
    print(f"{c:4s}: {ws[c].value}")
print("\nRows 14 to 18 peer levered betas:")
for r in range(14, 19):
    print(f"Row {r}: Peer={ws[f'A{r}'].value or ws[f'B{r}'].value}, Beta={ws[f'J{r}'].value}, Formula={ws[f'K{r}'].value}")
