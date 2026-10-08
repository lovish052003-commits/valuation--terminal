import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from institutional_control_checks import apply_institutional_wiring
from excel_exporter import strip_calc_chain_from_xlsx

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)
apply_institutional_wiring(wb)
wb.save('master_model_template.xlsx')
wb.close()
strip_calc_chain_from_xlsx('master_model_template.xlsx')
print('Successfully applied institutional wiring to master_model_template.xlsx!')

# Verify
wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)
print('Sheet list:', wb.sheetnames)
print('Control C21 (g):', wb['Control']['C21'].value)
print('Control C26 (LatestCol):', wb['Control']['C26'].value)
print('DCF D19:', wb['DCF']['D19'].value)
print('DCF I8:', wb['DCF']['I8'].value)
print('DCF B49:', repr(wb['DCF']['B49'].value))
print('DCF E51:', wb['DCF']['E51'].value)
print('Checks C3:', wb['Checks']['C3'].value)
print('WACC E26:', wb['WACC']['E26'].value)
print('WACC K26:', wb['WACC']['K26'].value)
print('Raw Data T24:', wb['Raw Data']['T24'].value)
print('AI Summary A1:', wb['AI Valuation Summary']['A1'].value)
print('AI Summary B2:', wb['AI Valuation Summary']['B2'].value)
print('Altman B5:', wb["Altman's Z Score"]['B5'].value)
print('Dupont B5:', wb['Dupont Analysis']['B5'].value)

