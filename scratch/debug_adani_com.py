import sys
import traceback
sys.path.insert(0, '.')
import screener_client
import valuation_engine
import excel_exporter

print("Fetching ADANIENT...")
sd = screener_client.fetch_company_data("ADANIENT")
print("Company Name:", sd.get('company_name'))
print("Ticker:", sd.get('ticker'))
val = valuation_engine.calculate_valuation(sd)

dest_path = "exports/TEST_ADANI_COM.xlsx"
print("Attempting export_via_excel_com...")
try:
    res = excel_exporter.export_via_excel_com(dest_path, sd, val)
    print("export_via_excel_com returned:", res)
except Exception as e:
    print("export_via_excel_com FAILED with exception:")
    traceback.print_exc()

import openpyxl
wb = openpyxl.load_workbook(dest_path, data_only=True)
print("\nVerifying exports/TEST_ADANI_COM.xlsx:")
print("Data Sheet B1:", wb['Data Sheet']['B1'].value)
print("Data Sheet B8 (CMP):", wb['Data Sheet']['B8'].value)
print("Data Sheet K17 (Sales):", wb['Data Sheet']['K17'].value)
print("Data Sheet K30 (PAT):", wb['Data Sheet']['K30'].value)
print("DCF D42 (IV):", wb['DCF']['D42'].value)
print("AI Valuation Summary in sheets:", 'AI Valuation Summary' in wb.sheetnames)
if 'AI Valuation Summary' in wb.sheetnames:
    print("AI Summary A1:", wb['AI Valuation Summary']['A1'].value)
    print("AI Summary B5:", wb['AI Valuation Summary']['B5'].value)
