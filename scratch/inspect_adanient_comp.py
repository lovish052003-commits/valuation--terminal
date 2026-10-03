import openpyxl

p = r"C:\Users\LENOVO\Downloads\test\ADANIENT_Valuation_Model (8).xlsx"
wb_v = openpyxl.load_workbook(p, data_only=True)
ws = wb_v['Comp_Valuation']

print("--- ADANIENT Comp_Valuation Row 10 Headers ---")
headers = [ws.cell(10, c).value for c in range(1, 25)]
print(headers)

print("--- ADANIENT Comp_Valuation Row 11 Headers ---")
headers11 = [ws.cell(11, c).value for c in range(1, 25)]
print(headers11)

print("--- ADANIENT Comp_Valuation Row 12 (First peer) ---")
row12 = [ws.cell(12, c).value for c in range(1, 25)]
print(row12)

print("--- ADANIENT Comp_Valuation Row 25 (Medians) ---")
row25 = [ws.cell(25, c).value for c in range(1, 25)]
print(row25)

print("--- ADANIENT Comp_Valuation Rows 30-39 ---")
for r in range(30, 40):
    vals = [ws.cell(r, c).value for c in [2, 9, 10, 15, 16, 17, 18]]
    print(f"Row {r}: {vals}")
