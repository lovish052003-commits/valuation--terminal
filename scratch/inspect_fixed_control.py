import openpyxl

wb_f = openpyxl.load_workbook("FORCEMOT_Valuation_Model_FIXED.xlsx", data_only=False)
ws_c = wb_f['Control']

print("=== CONTROL SHEET IN FIXED ===")
for r in range(1, 35):
    row_vals = [f"{openpyxl.utils.get_column_letter(c)}{r}: {ws_c.cell(r, c).value!r}" 
                for c in range(1, 15) if ws_c.cell(r, c).value is not None]
    if row_vals:
        print(" | ".join(row_vals))
