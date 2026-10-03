import openpyxl

p = r"C:\Users\LENOVO\Downloads\Advance Financial Project\exports\SBIN_Valuation_Model.xlsx"
wb_v = openpyxl.load_workbook(p, data_only=True)
wb_f = openpyxl.load_workbook(p, data_only=False)

if 'Raw FS' in wb_v.sheetnames:
    ws = wb_v['Raw FS']
    ws_f = wb_f['Raw FS']
    print("--- Raw FS Rows 55 to 65 ---")
    for r in range(55, 66):
        name = ws.cell(r, 4).value
        cmp_v = ws.cell(r, 13).value # Col M
        cmp_f = ws_f.cell(r, 13).value
        shares_v = ws.cell(r, 14).value # Col N
        mcap_v = ws.cell(r, 15).value # Col O
        mcap_block = ws.cell(r, 16).value # Col P
        print(f"Row {r:2d} | Name: {str(name)[:25]:25s} | CMP: {cmp_v} ({cmp_f}) | Shares: {shares_v} | MCap: {mcap_v} | Block MCap: {mcap_block}")
