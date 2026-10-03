import openpyxl, sys

fname = sys.argv[1] if len(sys.argv) > 1 else 'exports/TATASTEEL_Valuation_Model.xlsx'
print(f"Loading {fname}...")
wb = openpyxl.load_workbook(fname, data_only=True)
print(f"Total sheets in workbook: {len(wb.sheetnames)}")

# 1. Data Sheet
ws_ds = wb['Data Sheet']
print("\n[1] Data Sheet:")
print(f"    B1 (Company Name): {repr(ws_ds['B1'].value)}")
print(f"    B8 (Current Price): Rs. {ws_ds['B8'].value}")
print(f"    Sales Row 17: Y1 = {ws_ds['B17'].value} | Latest = {ws_ds['K17'].value}")
print(f"    Expenses Row 24: {ws_ds['K24'].value}")
print(f"    PAT Row 30: {ws_ds['K30'].value}")
print(f"    Equity Cap Row 57: {ws_ds['K57'].value}")
print(f"    Debt Row 59: {ws_ds['K59'].value}")
print(f"    Cash Row 69: {ws_ds['K69'].value}")
print(f"    Shares Row 70: {ws_ds['K70'].value}")

# 2. Raw FS
ws_raw = wb['Raw FS']
print("\n[2] Raw FS Sheet:")
print(f"    Target Name L56: {repr(ws_raw['L56'].value)}")
print(f"    Target Price M56: {ws_raw['M56'].value}")
print(f"    Target Mcap AR56: {ws_raw['AR56'].value}")
print(f"    Sector Peer 1 L57: {repr(ws_raw['L57'].value)}")
print(f"    Sector Peer 2 L58: {repr(ws_raw['L58'].value)}")
print(f"    Sector Peer 3 L59: {repr(ws_raw['L59'].value)}")

# 3. Comp_Valuation
ws_comp = wb['Comp_Valuation']
print("\n[3] Comp_Valuation Sheet:")
print(f"    Title B30: {repr(ws_comp['B30'].value)}")
for r in range(12, 17):
    print(f"    Row {r}: {ws_comp.cell(r, 2).value} | Price: {ws_comp.cell(r, 4).value}")

# 4. DCF
ws_dcf = wb['DCF']
print("\n[4] DCF Sheet:")
print(f"    D44 (Current Price): Rs. {ws_dcf['D44'].value}")
print(f"    D42 (Intrinsic Value): Rs. {ws_dcf['D42'].value}")
print(f"    D35 (Enterprise Value): Rs. {ws_dcf['D35'].value}")
print(f"    D37 (Cash): Rs. {ws_dcf['D37'].value}")
print(f"    D38 (Total Debt): Rs. {ws_dcf['D38'].value}")
print(f"    D40 (Shares Cr): {ws_dcf['D40'].value}")

# 5. WACC
ws_wacc = wb['WACC']
print("\n[5] WACC Sheet:")
print(f"    Target Company Row 15: {repr(ws_wacc['B15'].value)} | Beta: {ws_wacc['J15'].value}")
print(f"    WACC / Ke K29: {ws_wacc['K29'].value}")
for r in range(14, 19):
    print(f"    Row {r} Peer: {ws_wacc.cell(r, 2).value} | Debt: {ws_wacc.cell(r, 5).value} | Mcap: {ws_wacc.cell(r, 6).value}")


# 6. Dupont Analysis & Altman Z Score
ws_dup = wb['Dupont Analysis']
ws_alt = wb["Altman's Z Score"]
print("\n[6] DuPont Analysis & Altman Z-Score Sheets:")
print(f"    52-Week Range B5: {repr(ws_dup['B5'].value)}")
print(f"    Wikipedia Overview B8: {repr(str(ws_dup['B8'].value)[:75])}...")
print(f"    Altman Overview B8: {repr(str(ws_alt['B8'].value)[:75])}...")
print(f"    Recent Updates B37: {repr(str(ws_dup['B37'].value)[:75])}...")

# 7. AI Valuation Summary
if 'AI Valuation Summary' in wb.sheetnames:
    ws_ai = wb['AI Valuation Summary']
    print("\n[7] AI Valuation Summary Sheet:")
    print(f"    Header A1: {repr(ws_ai['A1'].value)}")
    print(f"    A5 (CMP): {ws_ai['A5'].value}")
    print(f"    B5 (Intrinsic Value): {ws_ai['B5'].value}")
    print(f"    C5 (Margin of Safety): {ws_ai['C5'].value}")
    print(f"    D5 (Verdict): {repr(ws_ai['D5'].value)}")
