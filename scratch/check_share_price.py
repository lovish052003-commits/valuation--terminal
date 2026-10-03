import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', read_only=False)
for sname in ['DCF', 'Intrinsic Valuation', 'Comp_Valuation', 'Customization']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        print(f"\n--- Sheet: {sname} ---")
        for r in range(1, 60):
            for c in range(1, 15):
                v = ws.cell(r, c).value
                if v is not None and any(w in str(v).lower() for w in ['share price', 'cmp', 'current market price', '276.5']):
                    col_letter = openpyxl.utils.get_column_letter(c)
                    print(f"  [{col_letter}{r}] = {v}")
                    # Also print neighboring cells
                    for dc in [1, 2, 3]:
                        nc = openpyxl.utils.get_column_letter(c + dc)
                        print(f"    Neighbor [{nc}{r}] = {ws.cell(r, c + dc).value}")
wb.close()
