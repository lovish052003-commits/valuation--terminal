import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx')
ws = wb['Beta-Regression']
max_pts = 247
for i in range(max_pts):
    raw_r = 6 + i
    beta_r = 10 + i
    ws.cell(row=beta_r, column=2, value=f"='Raw Data'!G{raw_r}")
    ws.cell(row=beta_r, column=3, value=f"='Raw Data'!H{raw_r}")
    ws.cell(row=beta_r, column=4, value="-" if i == 0 else f"=C{beta_r-1}/C{beta_r}-1")
    ws.cell(row=beta_r, column=6, value=f"='Raw Data'!G{raw_r}")
    ws.cell(row=beta_r, column=7, value=f"='Raw Data'!J{raw_r}")
    ws.cell(row=beta_r, column=8, value="-" if i == 0 else f"=G{beta_r-1}/G{beta_r}-1")

print('Sample Row 10:', ws['B10'].value, ws['C10'].value, ws['D10'].value, ws['F10'].value, ws['G10'].value, ws['H10'].value)
print('Sample Row 11:', ws['B11'].value, ws['C11'].value, ws['D11'].value, ws['F11'].value, ws['G11'].value, ws['H11'].value)
