import openpyxl

wb = openpyxl.load_workbook('exports/WIPRO_Valuation_Model.xlsx', data_only=True)
print("Sheetnames:", wb.sheetnames)
ws = wb['Data Sheet']
print("Data Sheet B1:", ws['B1'].value)
print("Data Sheet B6:", ws['B6'].value)
print("Data Sheet B8:", ws['B8'].value)
print("Data Sheet B9:", ws['B9'].value)
print("Data Sheet K17:", ws['K17'].value)
print("Data Sheet K69:", ws['K69'].value)
print("Data Sheet K70:", ws['K70'].value)
print("Data Sheet K59:", ws['K59'].value)

if 'AI Valuation Summary' in wb.sheetnames:
    ws_ai = wb['AI Valuation Summary']
    print("\nAI Valuation Summary Rows 1-15:")
    for r in range(1, 16):
        vals = [str(ws_ai.cell(r, c).value) for c in range(1, 6)]
        print(f"Row {r}: " + " | ".join(vals))

if 'DCF' in wb.sheetnames:
    ws_dcf = wb['DCF']
    print("\nDCF:")
    print("D37 (Cash):", ws_dcf['D37'].value)
    print("D38 (Debt):", ws_dcf['D38'].value)
    print("D40 (Shares):", ws_dcf['D40'].value)
    print("D42 (IV):", ws_dcf['D42'].value)
    print("D44 (CMP):", ws_dcf['D44'].value)
