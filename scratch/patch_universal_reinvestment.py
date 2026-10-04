import os
import openpyxl
import zipfile
import shutil
import re

def strip_calc_chain_from_xlsx(file_path: str):
    temp_zip = file_path + ".temp.zip"
    try:
        shutil.copyfile(file_path, temp_zip)
        with zipfile.ZipFile(temp_zip, 'r') as zin:
            with zipfile.ZipFile(file_path, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    if item.filename == "xl/calcChain.xml":
                        continue
                    if item.filename == "[Content_Types].xml":
                        data = zin.read(item.filename)
                        data = re.sub(rb'<Override PartName="/xl/calcChain\.xml"[^>]*/>', b'', data)
                        zout.writestr(item, data)
                        continue
                    if item.filename == "xl/_rels/workbook.xml.rels":
                        data = zin.read(item.filename)
                        data = re.sub(rb'<Relationship[^>]*Target="calcChain\.xml"[^>]*/>', b'', data)
                        zout.writestr(item, data)
                        continue
                    zout.writestr(item, zin.read(item.filename))
        if os.path.exists(temp_zip):
            os.remove(temp_zip)
        print(f"  [OK] Stripped calcChain.xml from {file_path}")
    except Exception as e:
        if os.path.exists(temp_zip):
            try: os.remove(temp_zip)
            except: pass
        print(f"  [Notice] strip_calc_chain: {e}")

def update_workbook(path):
    if not os.path.exists(path):
        return
    print(f"Updating {path}...")
    wb = openpyxl.load_workbook(path, data_only=False)
    if 'Intrinsic Valuation' in wb.sheetnames:
        ws_iv = wb['Intrinsic Valuation']
        if ws_iv['H51'].value == '-':
            ws_iv['H51'] = '=H37-G37'
        ws_iv['H52'] = '=H51/H49'
        ws_iv['I52'] = '=I51/I49'
        ws_iv['J52'] = '=J51/J49'
        ws_iv['K52'] = '=K51/K49'
        ws_iv['L52'] = '=L51/L49'
        ws_iv['H60'] = '=H52'
        ws_iv['H62'] = '=H59*H60'
        print("  [OK] Intrinsic Valuation: H52:L52 = H51:L51 / H49:L49")

    if 'DCF' in wb.sheetnames:
        ws_dcf = wb['DCF']
        ws_dcf['D18'] = "='Intrinsic Valuation'!L62"
        print("  [OK] DCF: D18 = ='Intrinsic Valuation'!L62")

    wb.save(path)
    wb.close()
    strip_calc_chain_from_xlsx(path)

if __name__ == '__main__':
    update_workbook('master_model_template.xlsx')
    update_workbook('ITC Model.xlsx')
    print("Done!")
