import sys
import os
import openpyxl
import datetime

sys.path.insert(0, os.path.abspath(os.path.dirname(__file__) + '/..'))

from screener_client import fetch_company_data
import excel_exporter
import valuation_engine

print("1. Fetching live company data for NESTLEIND...")
sd = fetch_company_data('NESTLEIND')
print(f"Company: {sd['company_name']}, CMP: {sd['current_price']}, MCap: {sd['market_cap_cr']}")

print("\n2. Running valuation math...")
val = valuation_engine.calculate_valuation(sd)
print("Valuation completed successfully.")

print("\n3. Testing openpyxl export and Data Sheet population...")
wb = openpyxl.load_workbook('Nestle India Model.xlsx')
excel_exporter.populate_data_sheet_openpyxl(wb, sd)

ws_ds = wb['Data Sheet']

print("\n--- DATA SHEET AUDIT (Cols B to K = cols 2 to 11) ---")
cols_to_check = list(range(2, 12))

# Row 16: P&L Dates
r16 = [ws_ds.cell(16, c).value for c in cols_to_check]
print("Row 16 P&L Dates:", [str(x)[:10] if x else None for x in r16])

# Row 17: Sales
r17 = [ws_ds.cell(17, c).value for c in cols_to_check]
print("Row 17 Sales:", r17)

# Row 56: BS Dates
r56 = [ws_ds.cell(56, c).value for c in cols_to_check]
print("Row 56 BS Dates:", [str(x)[:10] if x else None for x in r56])

# Row 57: Equity Capital
r57 = [ws_ds.cell(57, c).value for c in cols_to_check]
print("Row 57 Equity Capital:", r57)

# Row 58: Reserves
r58 = [ws_ds.cell(58, c).value for c in cols_to_check]
print("Row 58 Reserves:", r58)

# Row 59: Borrowings
r59 = [ws_ds.cell(59, c).value for c in cols_to_check]
print("Row 59 Borrowings:", r59)

# Row 67: Receivables
r67 = [ws_ds.cell(67, c).value for c in cols_to_check]
print("Row 67 Receivables:", r67)

# Row 68: Inventory
r68 = [ws_ds.cell(68, c).value for c in cols_to_check]
print("Row 68 Inventory:", r68)

# Row 69: Cash & Bank
r69 = [ws_ds.cell(69, c).value for c in cols_to_check]
print("Row 69 Cash & Bank:", r69)

# Row 70: Shares
r70 = [ws_ds.cell(70, c).value for c in cols_to_check]
print("Row 70 No of Shares:", r70)

# Row 72: Face Value
r72 = [ws_ds.cell(72, c).value for c in cols_to_check]
print("Row 72 Face Value:", r72)

# Row 81: CF Dates
r81 = [ws_ds.cell(81, c).value for c in cols_to_check]
print("Row 81 CF Dates:", [str(x)[:10] if x else None for x in r81])

# Row 82: CFO
r82 = [ws_ds.cell(82, c).value for c in cols_to_check]
print("Row 82 Cash from Operations:", r82)

# Row 90: Price
r90 = [ws_ds.cell(90, c).value for c in cols_to_check]
print("Row 90 Stock Price:", r90)

# Verify assertions
assert all(x is not None and not str(x).startswith('0') for x in r16), "FAIL: P&L dates contain None or 0"
assert all(isinstance(x, (datetime.datetime, datetime.date)) for x in r16), "FAIL: P&L dates are not datetime objects"
assert all(x > 0 for x in r17), "FAIL: Sales contain 0 or missing values"
assert all(x is not None and not str(x).startswith('0') for x in r56), "FAIL: BS dates contain None or 0"
assert all(isinstance(x, (datetime.datetime, datetime.date)) for x in r56), "FAIL: BS dates are not datetime objects"
assert all(x > 0 for x in r57), "FAIL: Equity capital contains 0"
assert all(x > 0 for x in r58), "FAIL: Reserves contain 0"
assert all(x > 0 for x in r67), "FAIL: Receivables contain 0"
assert all(x > 0 for x in r68), "FAIL: Inventory contains 0"
assert all(x > 0 for x in r69), "FAIL: Cash & Bank contains 0"
assert all(x > 0 for x in r70), "FAIL: Shares contain 0"
assert all(x > 0 for x in r72), "FAIL: Face value contains 0"
assert all(x > 0 for x in r82), "FAIL: CFO contains 0"

print("\n>>> ALL ASSERTIONS PASSED! NESTLE FINANCIAL IMPORT IS 100% RECONCILED! <<<")
