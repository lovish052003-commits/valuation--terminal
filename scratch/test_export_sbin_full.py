import os
import sys
import openpyxl

# Add project root to sys.path
sys.path.insert(0, os.path.abspath('.'))

from screener_client import fetch_company_data
import valuation_engine
from excel_exporter import export_valuation_model

def test_sbin():
    print("Fetching Screener data for SBIN...")
    s_data = fetch_company_data("SBIN")
    print(f"Company: {s_data.get('company_name')} | CMP: {s_data.get('current_price')} | Sector: {s_data.get('sector')}")

    print("Running Valuation Engine...")
    val_res = valuation_engine.calculate_valuation(s_data)
    print(f"Engine Result Intrinsic Value: {val_res.get('intrinsic_value')} | Method: {val_res.get('selected_methodology')}")

    print("Exporting Valuation Model...")
    out_path = export_valuation_model(s_data, val_res, report_markdown="")
    print(f"Exported Model Path: {out_path}")

    # Inspect the generated model
    wb = openpyxl.load_workbook(out_path, data_only=False)
    print("\n--- FORMULAS INSPECTION ---")
    ws_ai = wb['AI Valuation Summary']
    print(f"AI Valuation Summary A5: {ws_ai['A5'].value}")
    print(f"AI Valuation Summary B5: {ws_ai['B5'].value}")
    print(f"AI Valuation Summary C5: {ws_ai['C5'].value}")
    print(f"AI Valuation Summary D5: {ws_ai['D5'].value}")
    print(f"AI Valuation Summary E5: {ws_ai['E5'].value}")
    print(f"AI Valuation Summary F5: {ws_ai['F5'].value}")
    print(f"AI Valuation Summary G5: {ws_ai['G5'].value}")
    print(f"AI Valuation Summary A13: {ws_ai['A13'].value}")
    print(f"AI Valuation Summary A21: {ws_ai['A21'].value}")
    print(f"AI Valuation Summary B22: {ws_ai['B22'].value}")
    print(f"AI Valuation Summary B28: {ws_ai['B28'].value}")
    print(f"AI Valuation Summary B30: {ws_ai['B30'].value}")

    ws_dcf = wb['DCF']
    print(f"\nDCF B3: {ws_dcf['B3'].value}")
    print(f"DCF H8: {ws_dcf['H8'].value}")
    print(f"DCF D35: {ws_dcf['D35'].value}")
    print(f"DCF D39: {ws_dcf['D39'].value}")
    print(f"DCF D42: {ws_dcf['D42'].value}")
    wb.close()

    # Data only inspection (cached values)
    wb_val = openpyxl.load_workbook(out_path, data_only=True)
    print("\n--- CACHED EVALUATED VALUES INSPECTION ---")
    ws_ai_v = wb_val['AI Valuation Summary']
    print(f"AI Valuation Summary A5 (CMP): {ws_ai_v['A5'].value}")
    print(f"AI Valuation Summary B5 (Intrinsic): {ws_ai_v['B5'].value}")
    print(f"AI Valuation Summary C5 (Margin of Safety): {ws_ai_v['C5'].value}")
    print(f"AI Valuation Summary D5 (Valuation Gap): {ws_ai_v['D5'].value}")
    print(f"AI Valuation Summary E5 (Cost of Equity): {ws_ai_v['E5'].value}")
    print(f"AI Valuation Summary F5 (Altman Z): {ws_ai_v['F5'].value}")
    print(f"AI Valuation Summary G5 (DuPont ROE): {ws_ai_v['G5'].value}")
    print(f"AI Valuation Summary B22 (Book Value): {ws_ai_v['B22'].value}")
    print(f"AI Valuation Summary B28 (Net Equity Value): {ws_ai_v['B28'].value}")
    print(f"AI Valuation Summary B30 (Intrinsic Per Share): {ws_ai_v['B30'].value}")

    # Check for error tokens
    error_tokens = ['#VALUE!', '#REF!', '#DIV/0!', '#NAME?', '#NUM!']
    errors_found = []
    for s_name in ['AI Valuation Summary', 'Comp_Valuation', 'DCF', 'Intrinsic Valuation', 'Data Sheet']:
        if s_name in wb_val.sheetnames:
            ws = wb_val[s_name]
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if isinstance(cell, str) and any(err in cell for err in error_tokens):
                        errors_found.append((s_name, cell))

    wb_val.close()
    if errors_found:
        print(f"\n[FAIL] Found {len(errors_found)} error tokens: {errors_found[:10]}")
    else:
        print("\n[SUCCESS] ZERO ERROR TOKENS in key valuation sheets!")

if __name__ == '__main__':
    test_sbin()
