import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['AI Valuation Summary']
print('=== AI Valuation Summary in ITC Model.xlsx ===')
for r in range(1, ws.max_row + 1):
    row_vals = []
    for c in range(1, 10):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            col_letter = openpyxl.utils.get_column_letter(c)
            v_str = str(v).encode('ascii', errors='replace').decode('ascii')[:50]
            row_vals.append(f'{col_letter}{r}={v_str}')
    if row_vals:
        print(' | '.join(row_vals))
