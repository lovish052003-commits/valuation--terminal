import sys
import os
import openpyxl

sys.path.insert(0, os.path.abspath(os.path.dirname(__file__) + '/..'))

from screener_client import fetch_company_data
import excel_exporter
import valuation_engine

print("Running full end-to-end export_valuation_model for NESTLEIND...")
sd = fetch_company_data('NESTLEIND')
val = valuation_engine.calculate_valuation(sd)
model_path = excel_exporter.export_valuation_model(sd, val)
print("Exported file path:", model_path)

# Verify exported workbook
wb = openpyxl.load_workbook(model_path, data_only=True)
ws_ds = wb['Data Sheet']

cols = list(range(2, 12))
r16 = [ws_ds.cell(16, c).value for c in cols]
r56 = [ws_ds.cell(56, c).value for c in cols]
r81 = [ws_ds.cell(81, c).value for c in cols]

print("Data Sheet Row 16 (P&L Dates):", [str(x)[:10] if x else None for x in r16])
print("Data Sheet Row 56 (BS Dates):", [str(x)[:10] if x else None for x in r56])
print("Data Sheet Row 81 (CF Dates):", [str(x)[:10] if x else None for x in r81])

assert all(x is not None and '1900' not in str(x) and '0000' not in str(x) for x in r16)
assert all(x is not None and '1900' not in str(x) and '0000' not in str(x) for x in r56)
assert all(x is not None and '1900' not in str(x) and '0000' not in str(x) for x in r81)

print("\n>>> SUCCESS: NESTLE EXCEL VALUATION MODEL EXPORTED AND VERIFIED FLAWLESSLY! <<<")
