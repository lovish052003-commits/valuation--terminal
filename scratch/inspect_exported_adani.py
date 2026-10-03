import openpyxl

wb = openpyxl.load_workbook('exports/ADANIENT_Valuation_Model.xlsx', data_only=False)

if 'Comp_Valuation' in wb.sheetnames:
    ws = wb['Comp_Valuation']
    print("=== ADANIENT Comp_Valuation ===")
    for r in range(9, 25):
        vals = [f"{openpyxl.utils.get_column_letter(c)}{r}={ws.cell(r, c).value}" for c in range(2, 18) if ws.cell(r, c).value is not None]
        if vals:
            print(f"Row {r:2d}:", " | ".join(vals[:6]))

if 'Dupont Analysis' in wb.sheetnames:
    ws = wb['Dupont Analysis']
    print("\n=== ADANIENT Dupont Analysis ===")
    print("B5 (52W):", repr(ws['B5'].value))
    print("B8 (About):", repr(str(ws['B8'].value)[:100]))
    print("B37 (Upd1):", repr(str(ws['B37'].value)[:100]))

if "Altman's Z Score" in wb.sheetnames:
    ws = wb["Altman's Z Score"]
    print("\n=== ADANIENT Altman's Z Score ===")
    print("B5 (52W):", repr(ws['B5'].value))
    print("B8 (About):", repr(str(ws['B8'].value)[:100]))
    print("B36 (Upd1):", repr(str(ws['B36'].value)[:100]))
