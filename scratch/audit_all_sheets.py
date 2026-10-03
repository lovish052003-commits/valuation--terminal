import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
print("=== COMPLETE AUDIT OF ITC Model.xlsx (23 SHEETS) ===")
for i, name in enumerate(wb.sheetnames, 1):
    ws = wb[name]
    max_r = ws.max_row
    max_c = ws.max_column
    # Sample top cells
    samples = []
    for r in range(1, min(6, max_r + 1)):
        for c in range(1, min(5, max_c + 1)):
            v = ws.cell(row=r, column=c).value
            if v is not None:
                coord = f"{openpyxl.utils.get_column_letter(c)}{r}"
                v_str = str(v).encode('ascii', errors='replace').decode('ascii')[:30]
                samples.append(f"{coord}={v_str}")
    sample_str = " | ".join(samples[:4])
    print(f"{i:2d}. {name:25s} (max_row={max_r:3d}, max_col={max_c:2d}) -> {sample_str}")
