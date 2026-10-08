import openpyxl

wb_f = openpyxl.load_workbook("FORCEMOT_Valuation_Model_FIXED.xlsx", data_only=False)
ws_s = wb_f['AI Valuation Summary']

print("=== AI Valuation Summary rows 30-47 in FIXED ===")
for r in range(30, 48):
    row_vals = [f"{openpyxl.utils.get_column_letter(c)}{r}: {str(ws_s.cell(r, c).value)!r}" 
                for c in range(1, 9) if ws_s.cell(r, c).value is not None]
    if row_vals:
        print(" | ".join(row_vals).encode('ascii', errors='replace').decode('ascii'))
