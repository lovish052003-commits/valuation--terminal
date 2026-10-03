import openpyxl

for fname in ['master_model_template.xlsx', 'Hind. Unilever Model.xlsx']:
    wb = openpyxl.load_workbook(fname, data_only=False)
    sheet_name = 'AI Valuation Summary' if 'AI Valuation Summary' in wb.sheetnames else ('Summary' if 'Summary' in wb.sheetnames else None)
    if sheet_name:
        ws = wb[sheet_name]
        print(f"=== {fname} {sheet_name} Rows 1-45 ===")
        for r in range(1, 46):
            row_vals = []
            for c in range(1, 10):
                v = ws.cell(r, c).value
                if v is not None:
                    col_let = openpyxl.utils.get_column_letter(c)
                    s_v = str(v).replace('\u20b9', 'Rs.').replace('\u2014', '-').replace('\u2013', '-')
                    row_vals.append(f"{col_let}{r}={s_v}")
            if row_vals:
                print(" | ".join(row_vals))
