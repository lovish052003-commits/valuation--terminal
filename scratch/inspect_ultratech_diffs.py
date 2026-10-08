import openpyxl
import os

f1 = "UltraTech_fixed.xlsx"
f2 = "UltraTech_fixed_v2.xlsx"
orig = r"exports\ULTRACEMCO_Valuation_Model.xlsx"

print("f1 exists:", os.path.exists(f1))
print("f2 exists:", os.path.exists(f2))
print("orig exists:", os.path.exists(orig))

wb_f1 = openpyxl.load_workbook(f1, data_only=False)
wb_f2 = openpyxl.load_workbook(f2, data_only=False)
wb_orig = openpyxl.load_workbook(orig, data_only=False) if os.path.exists(orig) else None

print("\n--- SHEETS in f1 ---")
print(wb_f1.sheetnames)
print("\n--- SHEETS in f2 ---")
print(wb_f2.sheetnames)

def compare_wbs(wb_a, wb_b, label_a, label_b):
    print(f"\n=================== COMPARING {label_a} vs {label_b} ===================")
    for s in wb_b.sheetnames:
        if s not in wb_a.sheetnames:
            print(f"Sheet {s} is ONLY in {label_b}!")
            continue
        ws_a = wb_a[s]
        ws_b = wb_b[s]
        diffs = []
        max_r = max(ws_a.max_row, ws_b.max_row)
        max_c = max(ws_a.max_column, ws_b.max_column)
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                va = ws_a.cell(r, c).value
                vb = ws_b.cell(r, c).value
                if va != vb:
                    coord = openpyxl.utils.get_column_letter(c) + str(r)
                    diffs.append((coord, va, vb))
        if diffs:
            print(f"\nSheet {s}: {len(diffs)} differences:")
            for coord, va, vb in diffs[:30]:
                print(f"  {coord}: {label_a}={va!r} -> {label_b}={vb!r}")
            if len(diffs) > 30:
                print(f"  ... and {len(diffs) - 30} more diffs")
        else:
            print(f"Sheet {s}: IDENTICAL")

# First compare orig vs f1
if wb_orig:
    compare_wbs(wb_orig, wb_f1, "ORIG", "FIXED_V1")

# Then compare f1 vs f2
compare_wbs(wb_f1, wb_f2, "FIXED_V1", "FIXED_V2")
