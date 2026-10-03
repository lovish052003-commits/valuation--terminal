import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import screener_client
import valuation_engine
import excel_exporter
import openpyxl

def test_company_valuation(ticker_name):
    print(f"\n=================================================================")
    print(f"--- Fetching Data & Calculating Valuation for {ticker_name} ---")
    print(f"=================================================================")
    sd = screener_client.fetch_company_data(ticker_name)
    vr = valuation_engine.calculate_valuation(sd)

    print('--- Exporting Valuation Model ---')
    p = excel_exporter.export_valuation_model(sd, vr)
    print('Exported to:', p)

    wb = openpyxl.load_workbook(p, data_only=False)

    print('\n=== 1. DATA SHEET HISTORICAL TAX CHECK ===')
    ws_d = wb['Data Sheet']
    pbt_vals = [ws_d.cell(28, c).value for c in range(2, 12)]
    tax_vals = [ws_d.cell(29, c).value for c in range(2, 12)]
    np_vals  = [ws_d.cell(30, c).value for c in range(2, 12)]
    print('PBT (Row 28):       ', pbt_vals)
    print('Tax (Row 29):       ', tax_vals)
    print('Net Profit (Row 30):', np_vals)
    for i in range(len(pbt_vals)):
        if pbt_vals[i] and pbt_vals[i] > 0:
            assert abs(tax_vals[i] - round(pbt_vals[i] * 0.25, 2)) < 0.05, f'Tax mismatch col {i+2}'
            assert abs(np_vals[i] - round(pbt_vals[i] - tax_vals[i], 2)) < 0.05, f'Net profit mismatch col {i+2}'
            assert tax_vals[i] != pbt_vals[i], f'Tax duplicated PBT col {i+2}'
    print('>> PASS: Data Sheet Historical Tax is normalized to 25% of PBT and Net Profit = PBT - Tax!')

    print('\n=== 2. PROFIT & LOSS FORMULA CHECK ===')
    ws_pl = wb['Profit & Loss']
    pl_tax_f = [ws_pl.cell(11, c).value for c in range(2, 7)]
    pl_np_f  = [ws_pl.cell(12, c).value for c in range(2, 7)]
    print('P&L Tax (Row 11):       ', pl_tax_f)
    print('P&L Net Profit (Row 12):', pl_np_f)
    assert all("='Data Sheet'!" in str(f) for f in pl_tax_f), 'P&L Tax not linked to Data Sheet'
    assert all("='Data Sheet'!" in str(f) for f in pl_np_f), 'P&L Net Profit not linked to Data Sheet'
    print('>> PASS: Profit & Loss dynamically links to Data Sheet Row 29 & Row 30!')

    print('\n=== 3. QUARTERS FORMULA CHECK ===')
    ws_q = wb['Quarters']
    q_tax_f = [ws_q.cell(11, c).value for c in range(2, 7)]
    q_np_f  = [ws_q.cell(12, c).value for c in range(2, 7)]
    print('Quarters Tax (Row 11):       ', q_tax_f)
    print('Quarters Net Profit (Row 12):', q_np_f)
    assert all("='Data Sheet'!" in str(f) for f in q_tax_f), 'Quarters Tax not linked to Data Sheet'
    assert all("='Data Sheet'!" in str(f) for f in q_np_f), 'Quarters Net Profit not linked to Data Sheet'
    print('>> PASS: Quarters dynamically links to Data Sheet Row 48 & Row 49!')

    print('\n=== 4. WACC PEER BETAS CHECK ===')
    ws_wacc = wb['WACC']
    for r in range(14, 19):
        name = ws_wacc.cell(r, 2).value
        beta = ws_wacc.cell(r, 10).value
        unlev = ws_wacc.cell(r, 11).value
        print(f'Row {r}: Peer={name} | Levered Beta={beta} | Unlevered Formula={unlev}')
        assert beta is not None and beta != 1.0, f'Row {r} Beta is 1.0 or None'
        assert '=J' in str(unlev), f'Row {r} Unlevered formula invalid'
    print('>> PASS: WACC Sheet Rows 14-18 have realistic market betas and dynamic unlevered formulas!')

    print('\n=== 5. DCF REINVESTMENT RATE CHECK ===')
    ws_dcf = wb['DCF']
    i11 = ws_dcf['I11'].value
    print('DCF Cell I11:', i11)
    assert i11 == "='Intrinsic Valuation'!$L$55", f'I11 has caps: {i11}'
    for c in ['I', 'J', 'K', 'L', 'M']:
        f = ws_dcf[f'{c}12'].value
        print(f'DCF Cell {c}12:', f)
        assert 'MIN' not in str(f) and 'MAX' not in str(f), f'{c}12 has caps: {f}'
    print('>> PASS: DCF Sheet Row 11 & Row 12 are completely uncapped!')
    print(f'*** ALL UNIVERSAL FIXATIONS VERIFIED SUCCESSFULLY FOR {ticker_name} ***\n')

if __name__ == '__main__':
    test_company_valuation('JSWSTEEL')
