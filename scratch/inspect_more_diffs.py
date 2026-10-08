import openpyxl

wb_f = openpyxl.load_workbook("FORCEMOT_Valuation_Model_FIXED.xlsx", data_only=False)
wb_o = openpyxl.load_workbook(r"exports\FORCEMOT_Valuation_Model.xlsx", data_only=False)

def inspect_all_diffs(name):
    if name not in wb_f.sheetnames or name not in wb_o.sheetnames:
        return
    ws_f = wb_f[name]
    ws_o = wb_o[name]
    diffs = []
    max_r = max(ws_f.max_row, ws_o.max_row)
    max_c = max(ws_f.max_column, ws_o.max_column)
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            vf = ws_f.cell(r, c).value
            vo = ws_o.cell(r, c).value
            if vf != vo:
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                diffs.append((coord, vo, vf))
    if diffs:
        print(f"\n=================== ALL {len(diffs)} DIFFS IN {name} ===================")
        for coord, vo, vf in diffs:
            print(f"{coord}:\n   ORIG : {vo!r}\n   FIXED: {vf!r}")

inspect_all_diffs('DCF')
inspect_all_diffs('Comp_Valuation')
inspect_all_diffs('Intrinsic Valuation')
