import openpyxl

wb = openpyxl.load_workbook('exports/TMCV_Valuation_Model_reconciled.xlsx', data_only=False)

for s in ['AI Valuation Summary', 'DCF', 'Intrinsic Valuation']:
    ws = wb[s]
    print(f'*** {s} ***')
    if s == 'AI Valuation Summary':
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            print(f'{col}4: {ws[col+"4"].value} | {col}5: {ws[col+"5"].value}')
    elif s == 'DCF':
        for r in [6, 8, 10, 11, 12, 14, 16, 18, 19, 20, 21, 22, 25, 29, 33, 34, 35, 37, 38, 40, 41, 43, 45, 46]:
            print(f'Row {r}: B={ws["B"+str(r)].value} | D={ws["D"+str(r)].value} | I={ws["I"+str(r)].value}')
