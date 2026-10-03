import openpyxl

wb = openpyxl.load_workbook('exports/ADANIENT_Valuation_Model.xlsx', data_only=True)
if 'Raw FS' in wb.sheetnames:
    ws = wb['Raw FS']
    print("=== Raw FS Rows 56 to 66 (Values) ===")
    for r in range(56, 67):
        l_name = ws.cell(r, 12).value
        cmp_p = ws.cell(r, 13).value
        mcap = ws.cell(r, 44).value
        print(f"Row {r:2d}: Name={repr(l_name)} | CMP={cmp_p} | Mcap={mcap}")
