import openpyxl

wb = openpyxl.load_workbook('exports/ADANIENT_Valuation_Model_reconciled.xlsx', data_only=False)

print("=== 1. DATA SHEET ===")
ws_data = wb['Data Sheet']
print("K69 (Cash):", ws_data['K69'].value)
print("K70 (Shares):", ws_data['K70'].value)
print("B8 (CMP):", ws_data['B8'].value)

print("\n=== 2. DCF SHEET ===")
ws_dcf = wb['DCF']
print("D37 (Add Cash):", ws_dcf['D37'].value)
print("D40 (Shares in Cr):", ws_dcf['D40'].value)
print("D44 (Live Price):", ws_dcf['D44'].value)
print("I11 (Reinvest Y1):", ws_dcf['I11'].value)
print("M11 (Terminal Reinvest):", ws_dcf['M11'].value)
for c in ['I', 'J', 'K', 'L', 'M']:
    print(f"{c}12 (FCFF formula):", ws_dcf[f'{c}12'].value)

print("\n=== 3. COMP_VALUATION SHEET ===")
ws_comp = wb['Comp_Valuation']
for r in range(10, 16):
    vals = [f"{ws_comp.cell(r, c).coordinate}={ws_comp.cell(r, c).value}" for c in range(2, 11) if ws_comp.cell(r, c).value is not None]
    print(f"R{r}:", " | ".join(vals))
print("O25 (Median EV/Rev):", ws_comp['O25'].value)
print("P25 (Median EV/EBITDA):", ws_comp['P25'].value)

print("\n=== 4. AI VALUATION SUMMARY ===")
ws_sum = wb['AI Valuation Summary']
for r in [4, 5, 6, 7, 8, 10, 11, 12, 15, 16, 18, 19, 22, 23, 26, 27, 28, 29, 30, 31, 32]:
    vals = [f"{ws_sum.cell(r, c).coordinate}={ws_sum.cell(r, c).value}" for c in range(1, 8) if ws_sum.cell(r, c).value is not None]
    if vals:
        print(f"R{r}:", " | ".join(vals))
