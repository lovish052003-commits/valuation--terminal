import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)

# 1. Patch DCF Sheet
if 'DCF' in wb.sheetnames:
    ws_dcf = wb['DCF']
    # Dates: replace +365 with EDATE(prev, 12)
    ws_dcf['I6'] = "=EDATE(H6, 12)"
    ws_dcf['J6'] = "=EDATE(I6, 12)"
    ws_dcf['K6'] = "=EDATE(J6, 12)"
    ws_dcf['L6'] = "=EDATE(K6, 12)"
    ws_dcf['M6'] = "=EDATE(L6, 12)"

    # Fundamental Growth Rate
    ws_dcf['D18'] = "='Intrinsic Valuation'!L65"
    ws_dcf['D18'].number_format = '0.00%'

    # Terminal ROIC and Reinvestment Rate
    ws_dcf['B22'] = "Terminal ROIC (Stable)"
    ws_dcf['D22'] = "=MAX(0.18, MIN(0.25, 'Intrinsic Valuation'!L40))"
    ws_dcf['D22'].number_format = '0.00%'
    ws_dcf['D21'] = "=D19/D22"
    ws_dcf['D21'].number_format = '0.00%'

    # Cash & Liquid Investments Bridge (Cash & Bank K69 + Investments K64)
    ws_dcf['D37'] = "='Data Sheet'!K69+'Data Sheet'!K64"

    # Upside / (Downside)
    ws_dcf['B45'] = "Upside / (Downside)"
    ws_dcf['D45'] = "=D42/D44-1"
    ws_dcf['D45'].number_format = '0.0%'

# 2. Patch Intrinsic Valuation Sheet
if 'Intrinsic Valuation' in wb.sheetnames:
    ws_iv = wb['Intrinsic Valuation']
    # Row 38: EBIT = Operating Profit - Depreciation (Raw FS Row 6 - Row 10)
    ws_iv['I38'] = "='Raw FS'!AA6-'Raw FS'!AA10"
    ws_iv['J38'] = "='Raw FS'!AB6-'Raw FS'!AB10"
    ws_iv['K38'] = "='Raw FS'!AC6-'Raw FS'!AC10"
    ws_iv['L38'] = "='Raw FS'!AD6-'Raw FS'!AD10"

    # Row 44: Net CapEx = Gross CapEx - Depreciation (Raw FS Row 27:28 - Row 10)
    ws_iv['I44'] = "=-SUM('Raw FS'!AA27:AA28)-'Raw FS'!AA10"
    ws_iv['J44'] = "=-SUM('Raw FS'!AB27:AB28)-'Raw FS'!AB10"
    ws_iv['K44'] = "=-SUM('Raw FS'!AC27:AC28)-'Raw FS'!AC10"
    ws_iv['L44'] = "=-SUM('Raw FS'!AD27:AD28)-'Raw FS'!AD10"

# 3. Patch Raw Data Sheet
if 'Raw Data' in wb.sheetnames:
    ws_rd = wb['Raw Data']
    ws_rd['T24'] = 0.2517
    ws_rd['T24'].number_format = '0.00%'

    # Ensure Col Q (No. Eq. Shares Cr.) formula in rows 12 to 21
    for r in range(12, 22):
        ws_rd.cell(row=r, column=17, value=f"=IF(P{r}>0, R{r}/P{r}, 0)")

# 4. Patch WACC Sheet
if 'WACC' in wb.sheetnames:
    ws_wacc = wb['WACC']
    ws_wacc['J16'] = "='Beta-Regression'!O11"
    ws_wacc['E27'] = "='Raw Data'!T24"

# 5. Patch Dupont Analysis & Altman's Z Score Sheets
if 'Dupont Analysis' in wb.sheetnames:
    wb['Dupont Analysis']['B3'] = "='Data Sheet'!B1"
if "Altman's Z Score" in wb.sheetnames:
    wb["Altman's Z Score"]['B3'] = "='Data Sheet'!B1"

# Save updated template
wb.save('master_model_template.xlsx')
print('Successfully patched master_model_template.xlsx!')
