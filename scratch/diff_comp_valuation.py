import openpyxl

wb_tata = openpyxl.load_workbook('Tata Steel Final Model.xlsx', data_only=False)
ws_tata = wb_tata['Comp_Valuation']

wb_tmpl = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)
ws_tmpl = wb_tmpl['Comp_Valuation']

print("Comparing Comp_Valuation between Tata Steel and master_model_template:")
for r in range(9, 41):
    for c in range(2, 18):
        col_let = openpyxl.utils.get_column_letter(c)
        coord = f"{col_let}{r}"
        vt = ws_tata[coord].value
        vm = ws_tmpl[coord].value
        if vt != vm:
            print(f"Diff at {coord:4s}: Tata='{vt}' | Template='{vm}'")
