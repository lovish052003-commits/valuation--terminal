import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
from screener_client import fetch_company_data
import valuation_engine
import excel_exporter

ticker = 'ADANIENT'
print(f"=== 1. Fetching Screener Data for {ticker} ===")
screener_data = fetch_company_data(ticker)
print(f"Company: {screener_data.get('company_name')}")
print(f"Sector: {screener_data.get('sector')} | Industry: {screener_data.get('industry')}")
print(f"1Y Historical Prices count: {len(screener_data.get('historical_prices_1y', []))}")
print(f"Tables available: {list(screener_data.get('tables', {}).keys())}")

print(f"\n=== 2. Calculating Valuation ===")
val_result = valuation_engine.calculate_valuation(screener_data)
print(f"CMP: Rs. {val_result.get('current_price')}")
print(f"Intrinsic Value: Rs. {val_result.get('intrinsic_value_dcf')}")
print(f"Shares: {val_result.get('shares_cr')} Cr | Total Debt: Rs. {val_result.get('total_debt')} Cr")

print(f"\n=== 3. Exporting Valuation Model ===")
export_path = excel_exporter.export_valuation_model(screener_data, val_result, "")
print(f"Exported to: {export_path}")

print(f"\n=== 4. Inspecting Exported Workbook ===")
wb = openpyxl.load_workbook(export_path, data_only=False)

# Check Cash Flow Statement
ws_cf = wb['Cash Flow Statement']
print("\n--- Cash Flow Statement Sheet (Formulas/Labels) ---")
print("B2:", ws_cf['B2'].value)
print("Row 4 Dates:", [ws_cf.cell(4, c).value for c in range(3, 8)])
print("Row 6 (CFO):", [ws_cf.cell(6, c).value for c in range(3, 8)])
print("Row 7 (Profit from ops):", [ws_cf.cell(7, c).value for c in range(3, 8)])
print("Row 8 (Receivables):", [ws_cf.cell(8, c).value for c in range(3, 8)])
print("Row 13 (CFI):", [ws_cf.cell(13, c).value for c in range(3, 8)])
print("Row 24 (CFF):", [ws_cf.cell(24, c).value for c in range(3, 8)])
print("Row 31 (Other financing):", [ws_cf.cell(31, c).value for c in range(3, 8)])
print("Row 32 (Net Cash Flow):", [ws_cf.cell(32, c).value for c in range(3, 8)])
print("Row 32 (Latest periods):", [ws_cf.cell(32, c).value for c in range(ws_cf.max_column - 3, ws_cf.max_column + 1)])

# Check Beta Regression
ws_beta = wb['Beta-Regression']
print("\n--- Beta Regression Sheet ---")
print("B7:", ws_beta['B7'].value, "| F7:", ws_beta['F7'].value)
print("Row 10 Base:", [ws_beta.cell(10, c).value for c in [2, 3, 4, 6, 7, 8]])
print("Row 11 Day 1:", [ws_beta.cell(11, c).value for c in [2, 3, 4, 6, 7, 8]])
print("Row 12 Day 2:", [ws_beta.cell(12, c).value for c in [2, 3, 4, 6, 7, 8]])
print("Cell O11 (Slope):", ws_beta['O11'].value)
print("Cell L9:", ws_beta['L9'].value)
print("Cell L15 (Adjusted Beta):", ws_beta['L15'].value)

# Check WACC and Peers
ws_wacc = wb['WACC']
ws_raw = wb['Raw Data']
print("\n--- Raw Data (Rows 24-28 Peers) ---")
for r in range(24, 29):
    print(f"Row {r}: Name='{ws_raw.cell(r, 15).value}' | Country='{ws_raw.cell(r, 17).value}' | Debt={ws_raw.cell(r, 18).value} | Mcap={ws_raw.cell(r, 19).value}")

print("\n--- WACC Sheet (Rows 14-18 Peer Beta Comps) ---")
for r in range(14, 19):
    print(f"Row {r}: Comp='{ws_wacc.cell(r, 2).value}' | Debt='{ws_wacc.cell(r, 5).value}' | Mcap='{ws_wacc.cell(r, 6).value}' | LevBeta='{ws_wacc.cell(r, 10).value}' | UnlevBeta='{ws_wacc.cell(r, 11).value}'")

print("WACC Median Unlevered Beta (K21):", ws_wacc['K21'].value)
print("Target Total Debt (C34):", ws_wacc['C34'].value)
print("Target Equity (C35):", ws_wacc['C35'].value)
print("WACC Result (K46):", ws_wacc['K46'].value)

# Check DCF
ws_dcf = wb['DCF']
print("\n--- DCF Sheet ---")
print("DCF WACC (D20):", ws_dcf['D20'].value)
print("DCF Cash (D37):", ws_dcf['D37'].value)
print("DCF Debt (D38):", ws_dcf['D38'].value)
print("DCF Shares (D40):", ws_dcf['D40'].value)
print("DCF CMP (D44):", ws_dcf['D44'].value)

print("\n=== ALL INSPECTIONS COMPLETE ===")
