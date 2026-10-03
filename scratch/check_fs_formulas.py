import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
for sname in ['Profit & Loss', 'Balance Sheet', 'Cash Flow', 'Quarters']:
    ws = wb[sname]
    print(f"=== Sheet: {sname} ===")
    for r in range(3, 10):
        row_str = [f"{openpyxl.utils.get_column_letter(c)}{r}: {ws.cell(r, c).value}" for c in range(1, 6)]
        print("  " + " | ".join(row_str))
