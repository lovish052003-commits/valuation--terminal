import openpyxl
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

def apply_universal_changes_to_wb(wb_path):
    print(f"[UNIVERSAL] Processing: {os.path.basename(wb_path)}")
    wb = openpyxl.load_workbook(wb_path, data_only=False)

    # 1. WACC!E26 ='Data Sheet'!K27/AVERAGE('Data Sheet'!J59:K59)
    if 'WACC' in wb.sheetnames:
        ws_wacc = wb['WACC']
        ws_wacc['E26'] = "='Data Sheet'!K27/AVERAGE('Data Sheet'!J59:K59)"
        print("  -> WACC!E26 set to ='Data Sheet'!K27/AVERAGE('Data Sheet'!J59:K59)")

    # 2. DCF!D22 =MAX('Intrinsic Valuation'!L40,D20)
    if 'DCF' in wb.sheetnames:
        ws_dcf = wb['DCF']
        ws_dcf['D22'] = "=MAX('Intrinsic Valuation'!L40,D20)"
        print("  -> DCF!D22 set to =MAX('Intrinsic Valuation'!L40,D20)")

    # 3. Raw Data W12:W21 =IF(X{r}>0,U{r}/X{r},"N/A")
    if 'Raw Data' in wb.sheetnames:
        ws_rd = wb['Raw Data']
        for r in range(12, 22):
            ws_rd.cell(row=r, column=23, value=f'=IF(X{r}>0,U{r}/X{r},"N/A")')
        print("  -> Raw Data W12:W21 replaced with =IF(X{r}>0,U{r}/X{r},\"N/A\")")

    wb.calculation.fullCalcOnLoad = True
    wb.save(wb_path)
    wb.close()

    from patch_valuation_model import strip_calc_chain
    strip_calc_chain(wb_path)
    print(f"  -> Calc chain stripped. Successfully patched {os.path.basename(wb_path)}")

if __name__ == '__main__':
    import glob
    all_files = [f for f in glob.glob('*.xlsx') if not os.path.basename(f).startswith('~$')]
    for t in all_files:
        if os.path.exists(t):
            try:
                apply_universal_changes_to_wb(t)
            except Exception as e:
                print(f"Error processing {t}: {e}")
