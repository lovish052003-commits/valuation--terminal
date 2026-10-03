import openpyxl, sys

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
for name in wb.sheetnames:
    if name == 'List of Stocks':
        continue
    ws = wb[name]
    titles = []
    for r in range(1, 4):
        for c in range(1, 6):
            v = ws.cell(row=r, column=c).value
            if v:
                col_l = openpyxl.utils.get_column_letter(c)
                titles.append(f'{col_l}{r}={str(v)[:40]}')
    if titles:
        sys.stdout.buffer.write(f'[{name}] {" | ".join(titles[:4])}\n'.encode('utf-8'))
