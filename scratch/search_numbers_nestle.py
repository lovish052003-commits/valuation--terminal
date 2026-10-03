import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=True)

targets = {
    'Receivables': 88.97,
    'Inventory': 902.47,
    'Cash': 1457.42,
    'Shares': 96415716,
    'Price': 393.55,
    'AdjShares': 192.8
}

for name, target in targets.items():
    print(f"\nSearching for {name} ({target}):")
    for s in wb.sheetnames:
        ws = wb[s]
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None:
                    try:
                        fv = float(str(v).replace(',', ''))
                        if abs(fv - target) < 0.1:
                            lbl = ws.cell(r, 1).value or ws.cell(r, 2).value
                            print(f"  Found in '{s}' at {ws.cell(r, c).coordinate} (Row label: '{lbl}')")
                    except ValueError:
                        pass
