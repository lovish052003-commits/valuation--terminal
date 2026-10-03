import os
import sys
import json
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
from excel_exporter import export_valuation_model
from valuation_engine import calculate_valuation
from screener_client import load_offline_company_data

def test_company_export(ticker):
    print(f"\n==========================================")
    print(f"Testing End-to-End Export for {ticker}")
    print(f"==========================================")
    sd = load_offline_company_data(ticker)
    if not sd:
        print(f"Company data for {ticker} not loaded. Skipping.")
        return

    # Run valuation
    val_res = calculate_valuation(sd)
    if not val_res:
        val_res = {
            'ticker': ticker,
            'company_name': sd.get('company_name', ticker),
            'current_price': sd.get('current_price', 100),
            'fair_value': 120,
            'recommendation': 'BUY',
            'upside_downside_pct': 20.0
        }

    # Run model export
    exported_file = export_valuation_model(sd, val_res)
    print(f"Export completed: {exported_file}")
    assert os.path.exists(exported_file), f"File not created: {exported_file}"

    # Verify workbook content
    wb = openpyxl.load_workbook(exported_file, data_only=False)
    ws_raw = wb['Raw FS']

    # Read Raw FS Sales (Row 4), EBIT / Op Profit (Row 6), PBT (Row 11), PAT (Row 13)
    p_sales = [ws_raw.cell(row=4, column=c).value for c in range(19, 32)]
    p_ebit = [ws_raw.cell(row=6, column=c).value for c in range(19, 32)]
    p_pbt = [ws_raw.cell(row=11, column=c).value for c in range(19, 32)]
    p_pat = [ws_raw.cell(row=13, column=c).value for c in range(19, 32)]

    print(f"Raw FS Sales (Cols S-AE): {p_sales}")
    print(f"Raw FS Op Profit (Cols S-AE): {p_ebit}")
    print(f"Raw FS PBT (Cols S-AE): {p_pbt}")
    print(f"Raw FS PAT (Cols S-AE): {p_pat}")

    # Column AD (Col 30) is the latest annual fiscal year
    latest_sales = ws_raw.cell(row=4, column=30).value
    latest_pbt = ws_raw.cell(row=11, column=30).value
    latest_pat = ws_raw.cell(row=13, column=30).value

    # Template Maruti had Sales = 183,316, PBT = 18,340, Net Profit = 13,858
    assert latest_sales is not None and latest_sales < 50000, f"Template bleed detected in Sales: {latest_sales}"
    assert latest_pbt is not None and latest_pbt < 10000, f"Template bleed detected in PBT: {latest_pbt}"
    assert latest_pat is not None and latest_pat < 10000, f"Template bleed detected in PAT: {latest_pat}"

    # Also inspect evaluated values via Excel COM
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    excel = win32com.client.Dispatch('Excel.Application')
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb_com = excel.Workbooks.Open(os.path.abspath(exported_file))
        ws_iv = wb_com.Sheets('Intrinsic Valuation')
        val_l38 = ws_iv.Range('L38').Value
        ws_dcf = wb_com.Sheets('DCF')
        val_dcf_b8 = ws_dcf.Range('B8').Value
        val_dcf_c8 = ws_dcf.Range('C8').Value
        print(f"Evaluated Intrinsic Valuation L38 (EBIT): {val_l38}")
        print(f"Evaluated DCF Base EBIT (B8): {val_dcf_b8}, Year 1 Forecast EBIT (C8): {val_dcf_c8}")
        wb_com.Close(False)
        assert val_l38 is not None and val_l38 < 1000, f"Template bleed in Intrinsic Valuation EBIT: {val_l38}"
    finally:
        excel.Quit()

    print(f"SUCCESS: {ticker} passed all template bleed tests!")

if __name__ == '__main__':
    for t in ['UNITEDTEA', 'OISL']:
        test_company_export(t)
