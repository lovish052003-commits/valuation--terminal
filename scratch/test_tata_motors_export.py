import os, sys
sys.path.insert(0, os.path.abspath(os.path.dirname(__file__) + '/..'))
import screener_client, valuation_engine, excel_exporter
import openpyxl

print("1. Fetching Tata Motors data...")
d = screener_client.fetch_company_data('Tata Motors')
val = valuation_engine.calculate_valuation(d)

dest = os.path.abspath('exports/TEST_TATAMOTORS_DEBUG.xlsx')
print("2. Exporting via COM...")
excel_exporter.export_via_excel_com(dest, d, val, "")

wb = openpyxl.load_workbook(dest, data_only=True)
ws_data = wb['Data Sheet']
ws_dcf = wb['DCF']
ws_ai = wb['AI Valuation Summary']
ws_comp = wb['Comp_Valuation']

print("\n--- DATA SHEET DEBUG ---")
print("B1 (Name):", ws_data['B1'].value)
print("B6 (Shares formula/val):", ws_data['B6'].value)
print("B7 (Face value):", ws_data['B7'].value)
print("B8 (Price):", ws_data['B8'].value)
print("B9 (Market cap):", ws_data['B9'].value)
print("K57 (Equity Cap):", ws_data['K57'].value)
print("K69 (Cash):", ws_data['K69'].value)
print("K70 (Shares raw):", ws_data['K70'].value)
print("K72 (Face Value):", ws_data['K72'].value)

print("\n--- DCF SHEET DEBUG ---")
print("D37 (Cash):", ws_dcf['D37'].value)
print("D38 (Debt):", ws_dcf['D38'].value)
print("D40 (Shares Cr):", ws_dcf['D40'].value)
print("D42 (Intrinsic Value):", ws_dcf['D42'].value)
print("D44 (Current Price):", ws_dcf['D44'].value)

print("\n--- COMP VALUATION DEBUG ---")
print("B30 Header:", ws_comp['B30'].value)
for r in range(12, 22):
    print(f"Row {r}: B={ws_comp.cell(r,2).value} | C={ws_comp.cell(r,3).value} | D={ws_comp.cell(r,4).value} | E={ws_comp.cell(r,5).value} | F={ws_comp.cell(r,6).value} | G={ws_comp.cell(r,7).value} | N={ws_comp.cell(r,14).value} | O={ws_comp.cell(r,15).value} | P={ws_comp.cell(r,16).value}")

print("\n--- AI VALUATION SUMMARY DEBUG ---")
for r in range(1, 40):
    vals = [str(ws_ai.cell(r, c).value).encode('ascii', 'replace').decode() for c in range(1, 8)]
    if any(v != 'None' for v in vals):
        print(f"Row {r}: " + " | ".join(f"{openpyxl.utils.get_column_letter(c)}={vals[c-1]}" for c in range(1, 8)))
