import openpyxl
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from screener_client import fetch_company_data
from excel_exporter import export_valuation_model, get_effective_peers

print("1. Fetching Screener data for RELIANCE...")
sd = fetch_company_data('RELIANCE')
print(f"Target: {sd.get('ticker')} ({sd.get('company_name')})")

print("\n2. Checking effective peers for RELIANCE:")
peers = get_effective_peers(sd)
for idx, p in enumerate(peers, 1):
    print(f"  Peer {idx}: {p.get('ticker')} - {p.get('name')} | MCap: {p.get('mcap')} | Segment: {p.get('segment')}")

banned = ['continental', 'gulf oil', 'savita', 'gandhar', 'gp petroleum', 'gp petroleums']
for p in peers:
    for b in banned:
        assert b not in p['name'].lower(), f"Banned microcap {b} found in peer {p['name']}"

print("All banned microcaps successfully dropped!")

print("\n3. Generating Valuation Excel for RELIANCE...")
output_path = export_valuation_model(sd, None)
print(f"Generated: {output_path}")

wb = openpyxl.load_workbook(output_path, data_only=False)

print("\n4. Verifying Formulas in Generated Model:")
print("  WACC!E26:", wb['WACC']['E26'].value)
assert wb['WACC']['E26'].value in ["='Data Sheet'!K27/AVERAGE('Data Sheet'!J59:K59)", "=INDEX('Data Sheet'!$B27:$K27, Control!$C$26)/AVERAGE(INDEX('Data Sheet'!$B59:$K59, Control!$C$26-1), INDEX('Data Sheet'!$B59:$K59, Control!$C$26))"], f"WACC!E26 mismatch: {wb['WACC']['E26'].value}"

print("  DCF!D22:", wb['DCF']['D22'].value)
d22_val = str(wb['DCF']['D22'].value)
assert "MAX(" in d22_val and "L40" in d22_val and "D20" in d22_val, f"DCF!D22 mismatch: {d22_val}"

print("  Raw Data!W12:", wb['Raw Data']['W12'].value)
print("  Raw Data!W20:", wb['Raw Data']['W20'].value)
assert wb['Raw Data']['W20'].value == '=IF(X20>0,U20/X20,"N/A")', f"Raw Data!W20 mismatch: {wb['Raw Data']['W20'].value}"

print("\n5. Verifying Comp_Valuation Sheet:")
ws_comp = wb['Comp_Valuation']
print("  Comp_Valuation B30:", ws_comp['B30'].value)
print("  Comp_Valuation P32:", ws_comp['P32'].value)
print("  Comp_Valuation B42 (SOTP Header):", ws_comp['B42'].value)
print("  Comp_Valuation B44 (O2C):", ws_comp['B44'].value, "| EV Formula:", ws_comp['G44'].value)
print("  Comp_Valuation B45 (Telecom):", ws_comp['B45'].value, "| EV Formula:", ws_comp['G45'].value)
print("  Comp_Valuation B46 (Retail):", ws_comp['B46'].value, "| EV Formula:", ws_comp['G46'].value)
print("  Comp_Valuation G49 (Net SOTP EV):", ws_comp['G49'].value)
print("  Comp_Valuation G53 (SOTP Value per Share):", ws_comp['G53'].value)

for r in range(12, 18):
    print(f"  Row {r} Peer: {ws_comp.cell(row=r, column=2).value} | Ticker: {ws_comp.cell(row=r, column=3).value}")

print("\nALL VERIFICATIONS PASSED PERFECTLY!")
