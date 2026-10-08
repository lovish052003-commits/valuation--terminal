import openpyxl

wb = openpyxl.load_workbook("UltraTech_fixed_v2.xlsx", data_only=False)

def print_row_details(ws, rows, cols):
    print(f"\n--- {ws.title} Details ---")
    for r in rows:
        row_str = []
        for c in cols:
            cell = ws.cell(r, c)
            val = cell.value
            if val is not None:
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                row_str.append(f"{coord}: val={val!r}, num_fmt={cell.number_format!r}")
        if row_str:
            print(" | ".join(row_str).encode('ascii', errors='replace').decode('ascii'))

print_row_details(wb['AI Valuation Summary'], [4, 5, 10, 12, 40, 41, 42], range(1, 11))
print_row_details(wb['Control'], range(31, 38), range(1, 6))
print_row_details(wb['Checks'], [13] + list(range(17, 23)), range(1, 7))
