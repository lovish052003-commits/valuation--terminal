import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=False)
ws = wb['Data Sheet']

print("=== Data Sheet in Nestle India (2).xlsx ===")
for r in range(55, 95):
    row_vals = [f"{ws.cell(r, c).coordinate}: {ws.cell(r, c).value}" for c in range(1, 12) if ws.cell(r, c).value is not None]
    if row_vals:
        print(f"Row {r:2d}: " + " | ".join(row_vals[:5]))
