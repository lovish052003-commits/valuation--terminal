import openpyxl

wb = openpyxl.load_workbook('exports/ADANIENT_Valuation_Model.xlsx', data_only=False)

# 1. Altman's Z Score sheet
alt_name = [s for s in wb.sheetnames if 'Altman' in s][0]
ws_alt = wb[alt_name]
print("=== ALTMAN SHEET ===")
for r in range(1, 45):
    for c in range(1, 15):
        val = str(ws_alt.cell(r, c).value or '')
        if any(w in val.lower() for w in ['z-score', 'z score', 'altman', 'zone', 'score']):
            coord = ws_alt.cell(r, c).coordinate
            print(f"Altman {coord}: {val}")

# 2. Dupont sheet
dup_name = [s for s in wb.sheetnames if 'Dupont' in s or 'DuPont' in s][0]
ws_dup = wb[dup_name]
print("\n=== DUPONT SHEET ===")
for r in range(1, 35):
    for c in range(1, 15):
        val = str(ws_dup.cell(r, c).value or '')
        if 'roe' in val.lower():
            coord = ws_dup.cell(r, c).coordinate
            print(f"Dupont {coord}: {val}")

# 3. AI Valuation Summary
if 'AI Valuation Summary' in wb.sheetnames:
    ws_sum = wb['AI Valuation Summary']
    print("\n=== AI SUMMARY ===")
    for r in range(1, 35):
        vals = [f"{ws_sum.cell(r, c).coordinate}={ascii(ws_sum.cell(r, c).value)}" for c in range(1, 8) if ws_sum.cell(r, c).value is not None]
        if vals:
            print(f"R{r}:", " | ".join(vals))
