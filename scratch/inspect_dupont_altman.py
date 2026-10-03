import openpyxl

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)

ws_dup = wb['Dupont Analysis']
print('=== DUPONT ANALYSIS SHEET ===')
for r in [49, 50, 51, 52, 60, 61, 62, 63, 64, 65, 66, 67, 68, 69, 70, 71, 72, 73, 74, 75, 76, 77, 78]:
    lbl = ws_dup.cell(r, 2).value or ws_dup.cell(r, 1).value
    val_e = ws_dup.cell(r, 5).value # Col E
    val_i = ws_dup.cell(r, 9).value # Col I
    print(f'Row {r:2d} | Col B: {str(lbl)[:30]:30s} | Col E: {str(val_e)[:30]:30s} | Col I: {str(val_i)[:30]:30s}')

ws_alt = wb["Altman's Z Score"]
print('\n=== ALTMAN Z SCORE SHEET ===')
for r in [58, 59, 60, 61, 62, 63, 64, 65, 66, 67, 68, 69, 70, 71, 72, 73, 74, 75, 76, 77, 78, 85, 86, 87, 88, 89, 90]:
    lbl = ws_alt.cell(r, 2).value or ws_alt.cell(r, 1).value
    val_e = ws_alt.cell(r, 5).value # Col E
    val_i = ws_alt.cell(r, 9).value # Col I
    print(f'Row {r:2d} | Col B: {str(lbl)[:30]:30s} | Col E: {str(val_e)[:30]:30s} | Col I: {str(val_i)[:30]:30s}')
