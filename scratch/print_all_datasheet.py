import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=False)
ws = wb['Data Sheet']

print("=== ALL ROWS IN DATA SHEET (Nestle India 2.xlsx) ===")
for r in range(1, 95):
    lbl = ws.cell(r, 1).value
    val_b = ws.cell(r, 2).value
    val_k = ws.cell(r, 11).value
    if lbl or val_b:
        print(f"Row {r:2d}: {str(lbl):32s} | B={str(val_b)[:25]:25s} | K={str(val_k)[:25]}")
