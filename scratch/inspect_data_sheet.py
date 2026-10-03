import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['Data Sheet']
print('--- ITC Model.xlsx Data Sheet top rows ---')
for r in range(1, 15):
    print(f'Row {r:2d}: A={ws[f"A{r}"].value!r:30s} | B={ws[f"B{r}"].value!r:30s}')

print('\n--- Key rows in Data Sheet ---')
for r in [55, 56, 57, 58, 59, 60, 68, 69, 70, 71, 72, 73, 74, 75]:
    print(f'Row {r:2d}: A={ws[f"A{r}"].value!r:30s} | B={ws[f"B{r}"].value!r:20s} | K={ws[f"K{r}"].value!r}')
