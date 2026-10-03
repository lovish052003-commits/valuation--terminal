import openpyxl

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)
ws = wb['Intrinsic Valuation']
rows_with_raw = {}
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        val = str(ws.cell(r, c).value or '')
        if 'raw fs' in val.lower():
            row_label = str(ws.cell(r, 2).value or ws.cell(r, 1).value or '')
            rows_with_raw.setdefault(r, {'label': row_label, 'cells': []})
            rows_with_raw[r]['cells'].append((ws.cell(r, c).coordinate, val))

for r, info in sorted(rows_with_raw.items()):
    lbl = info['label'][:35]
    print(f"Row {r:2d} [{lbl:35s}]: {len(info['cells']):2d} cells -> e.g. {info['cells'][0]}")
