import os, sys
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

def test_dcf_share_price(ticker):
    print(f"\n==========================================")
    print(f"Testing DCF Share Price for: {ticker}")
    print(f"==========================================")
    data = fetch_company_data(ticker)
    cmp = data.get('current_price')
    print(f"Fetched Screener CMP for {ticker}: Rs. {cmp}")

    val = calculate_valuation(data)
    out_file = export_valuation_model(data, val, f"### Valuation Report for {ticker}")
    print(f"Exported to: {out_file}")

    wb = openpyxl.load_workbook(out_file, data_only=False)
    print(f"Total Sheets: {len(wb.sheetnames)}")
    ws = wb['DCF']
    print(f"DCF Sheet Inspection:")
    print(f"  B42: {ws['B42'].value} | D42: {ws['D42'].value}")
    print(f"  B44: {ws['B44'].value} | D44: {ws['D44'].value}")
    print(f"  B45: {ws['B45'].value} | D45: {ws['D45'].value}")
    wb.close()

    # Also check evaluated values
    wb_eval = openpyxl.load_workbook(out_file, data_only=True)
    ws_eval = wb_eval['DCF']
    print(f"Evaluated Values in DCF:")
    print(f"  D42 (Equity Value per Share): {ws_eval['D42'].value}")
    print(f"  D44 (Share Price): {ws_eval['D44'].value}")
    print(f"  D45 (Discount/Premium): {ws_eval['D45'].value}")
    wb_eval.close()

if __name__ == '__main__':
    test_dcf_share_price('HINDUNILVR')
    test_dcf_share_price('BRITANNIA')
