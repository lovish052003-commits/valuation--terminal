import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
wb = openpyxl.load_workbook(r"C:\Users\LENOVO\Downloads\test\ADANIENT_Valuation_Model (8).xlsx", data_only=False)
wb_d = openpyxl.load_workbook(r"C:\Users\LENOVO\Downloads\test\ADANIENT_Valuation_Model (8).xlsx", data_only=True)
ws = wb['AI Valuation Summary']
ws_d = wb_d['AI Valuation Summary']
print("=== ADANIENT (8) AI Valuation Summary Rows 21-33 ===")
for r in range(21, 34):
    print(f"Row {r:2d}: A='{ws[f'A{r}'].value}' | B_form='{ws[f'B{r}'].value}' | B_eval={ws_d[f'B{r}'].value}")
