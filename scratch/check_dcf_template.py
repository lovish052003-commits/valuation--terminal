import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['DCF']
print("=== DCF Sheet Key Cells ===")
cells = ['D19', 'D20', 'D21', 'D29', 'D33', 'D34', 'D35', 'D37', 'D38', 'D39', 'D40', 'D42', 'D44']
for c in cells:
    print(f"{c:4s}: {ws[c].value}")
print("\nForecast Rows 10-16 across Cols I to M:")
for r in [8, 10, 11, 12, 13, 14, 16]:
    vals = [ws.cell(r, c).value for c in range(9, 14)]
    print(f"Row {r:2d} (Cols I..M): {vals}")
