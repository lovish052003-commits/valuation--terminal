import openpyxl

wb = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=False)

def inspect_sheet_summary(name):
    if name not in wb.sheetnames:
        print(f"Sheet {name} does not exist.")
        return
    ws = wb[name]
    print(f"\n=================== SHEET: {name} (max_row={ws.max_row}, max_col={ws.max_column}) ===================")
    for r in range(1, min(ws.max_row + 1, 60)):
        row_vals = [str(ws.cell(r, c).value or '') for c in range(1, min(ws.max_column + 1, 16))]
        if any(row_vals):
            last_idx = max([i for i, v in enumerate(row_vals) if v != ''] + [0])
            print(f"R{r:02d}: " + " | ".join(row_vals[:last_idx+1]))

for s in ['WACC', 'Comp_Valuation', 'Forecasting', 'Ratio Analysis', 'Historical FS', 'Dupont Analysis', 'Altman\'s Z Score']:
    inspect_sheet_summary(s)
