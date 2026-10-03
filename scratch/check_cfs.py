import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['Cash Flow Statement']
print('=== Cash Flow Statement Rows 4 to 34 ===')
for r in range(4, 35):
    print(f'Row {r:2d}: Particulars = "{ws.cell(r, 2).value}" | Col C = {ws.cell(r, 3).value}')
