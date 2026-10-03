import sys
import os
import re
sys.path.insert(0, os.path.abspath('.'))
import openpyxl

import screener_client
import valuation_engine
import excel_exporter

def test_company_resolution_and_classification():
    print("=== [1] Testing Company Resolution & 7-Tier Classification ===")
    
    test_cases = [
        ('LICI', 'INSURANCE', True),
        ('SBIN', 'BANK', True),
        ('BAJFINANCE', 'NBFC', True),
        ('TCS', 'NON_FINANCIAL', False),
        ('ITC', 'NON_FINANCIAL', False),
        ('HDFCBANK', 'BANK', True)
    ]
    
    for ticker, exp_type, exp_is_fin in test_cases:
        res = screener_client.resolve_company(ticker)
        assert res is not None, f"Failed to resolve {ticker}"
        c_type = screener_client.classify_company(
            res.get('sector', ''),
            res.get('industry', ''),
            res.get('company_name', '')
        )
        assert c_type == exp_type, f"Expected {exp_type} for {ticker}, got {c_type}"
        
        # Test sector key and is_financial_sector
        dummy_sd = {'company_type': c_type, 'company_name': res.get('company_name'), 'ticker': ticker}
        sec_key = excel_exporter.get_sector_key(dummy_sd)
        is_fin = excel_exporter.is_financial_sector(sec_key)
        assert is_fin == exp_is_fin, f"Expected is_financial={exp_is_fin} for {ticker} ({sec_key}), got {is_fin}"
        print(f"  [OK] {ticker:12} -> Type: {c_type:15} | Sector: {sec_key:15} | Financial: {is_fin}")

    # String normalization & disambiguation
    print("\n  Testing ambiguous query: 'TATA'")
    try:
        screener_client.resolve_company("TATA", raise_on_ambiguous=True)
        assert False, "Should have raised AmbiguousCompanyError for 'TATA'"
    except screener_client.AmbiguousCompanyError as e:
        assert len(e.candidates) > 1, "Expected multiple candidates for 'TATA'"
        print(f"  [OK] Successfully raised AmbiguousCompanyError for 'TATA' with {len(e.candidates)} candidates")


def test_market_price_validation():
    print("\n=== [2] Testing Market Price Validation ===")
    base_data = {
        'ticker': 'TESTCO',
        'company_name': 'TestCo',
        'market_cap_cr': 10000,
        'tables': {'profit-loss': {}, 'balance-sheet': {}}
    }
    # Valid price
    is_valid, msg = valuation_engine.validate_valuation_inputs({**base_data, 'current_price': 500})
    assert is_valid, f"Expected valid, got: {msg}"
    print("  [OK] Valid price accepted (500)")

    # Missing / <= 0 price
    for invalid_p in [0, -10, None]:
        is_valid, msg = valuation_engine.validate_valuation_inputs({**base_data, 'current_price': invalid_p})
        assert not is_valid, f"Expected invalid for price {invalid_p}"
        assert "Current market price unavailable for TestCo." in msg, f"Unexpected message: {msg}"
        print(f"  [OK] Invalid price ({invalid_p}) correctly rejected with message: '{msg}'")


def test_sequential_runs_and_data_isolation():
    print("\n=== [3] Sequential Valuation Runs (LICI -> TCS -> ITC -> SBIN -> BAJFINANCE -> HDFCBANK) ===")
    sequence = ['LICI', 'TCS', 'ITC', 'SBIN', 'BAJFINANCE', 'HDFCBANK']
    contexts = []

    for ticker in sequence:
        print(f"\n--- Running Valuation Pipeline for {ticker} ---")
        sd = screener_client.fetch_company_data(ticker)
        assert sd is not None, f"Failed to fetch data for {ticker}"
        
        # 1. Verify company identity & CMP
        c_name = sd.get('company_name', '').strip()
        cmp = sd.get('current_price', 0)
        assert cmp > 0, f"CMP must be > 0 for {ticker}, got {cmp}"
        c_type = sd.get('company_type', 'NON_FINANCIAL')
        sec_key = excel_exporter.get_sector_key(sd)
        is_fin = excel_exporter.is_financial_sector(sec_key)
        
        # 2. Valuation calculation
        vr = valuation_engine.calculate_valuation(sd)
        ctx = vr.get('valuation_context', {})
        assert ctx.get('ticker') == ticker, f"Context ticker mismatch: {ctx.get('ticker')}"
        assert ctx.get('cmp') == cmp, f"Context cmp mismatch: {ctx.get('cmp')} vs {cmp}"
        assert ctx.get('company_type') == c_type, f"Context company_type mismatch: {ctx.get('company_type')} vs {c_type}"
        
        # 3. Dynamic Peers validation
        peers = excel_exporter.get_effective_peers(sd, vr)
        assert len(peers) >= 5, f"Expected at least 5 peers for {ticker}, got {len(peers)}"
        peer_tickers = [p.get('ticker', '').upper() for p in peers]
        peer_names = [p.get('name', '').lower() for p in peers]
        assert ticker not in peer_tickers, f"Target {ticker} must NEVER appear in peer group! Peers: {peer_tickers}"
        assert c_name.lower() not in peer_names, f"Target company name {c_name} must NEVER appear in peer group!"
        assert len(peer_tickers) == len(set(peer_tickers)), f"Peer tickers contain duplicates for {ticker}: {peer_tickers}"

        # Check peer classification matching
        if ticker == 'LICI':
            # Peers must be insurance companies, NOT nbfcs or general banks
            assert any(t in peer_tickers for t in ['SBILIFE', 'HDFCLIFE', 'ICICIPRULI', 'GICRE', 'NIACL', 'STARHEALTH', 'MFSL', 'CANHLIFE']), \
                f"LICI peers should be insurance companies, got {peer_tickers}"

        print(f"  [OK] {ticker}: Name='{c_name}', CMP={cmp}, Type={c_type}, is_financial={is_fin}, Peers Count={len(peers)}")
        print(f"       Peers: {peer_tickers[:6]}")

        # 4. Model Export
        export_path = excel_exporter.export_valuation_model(sd, vr)
        assert os.path.exists(export_path), f"Export file missing: {export_path}"

        wb = openpyxl.load_workbook(export_path, data_only=False)

        # A. DCF!D44 Dynamic Reference Check
        ws_dcf = wb['DCF']
        dcf_d44_val = str(ws_dcf['D44'].value or '')
        assert dcf_d44_val == "='Data Sheet'!B8", f"DCF!D44 must be dynamic formula ='Data Sheet'!B8, got {dcf_d44_val}"
        
        ws_data = wb['Data Sheet']
        assert ws_data['B8'].value == cmp, f"Data Sheet!B8 must contain target current price {cmp}, got {ws_data['B8'].value}"
        assert ws_data['B1'].value == c_name, f"Data Sheet!B1 must contain target company name {c_name}, got {ws_data['B1'].value}"

        # B. Comp_Valuation Multiples Check
        ws_comp = wb['Comp_Valuation']
        assert ws_comp['B30'].value == f"{c_name} Comparable Valuation"
        if is_fin:
            # Financials: EV/Revenue and EV/EBITDA must be "N/A"
            assert ws_comp['O12'].value == "N/A", f"{ticker} (Financial): Peer 1 EV/Revenue (O12) should be N/A, got {ws_comp['O12'].value}"
            assert ws_comp['P12'].value == "N/A", f"{ticker} (Financial): Peer 1 EV/EBITDA (P12) should be N/A, got {ws_comp['P12'].value}"
            assert ws_comp['O25'].value == "N/A", f"{ticker} (Financial): Median EV/Revenue (O25) should be N/A, got {ws_comp['O25'].value}"
            assert ws_comp['P25'].value == "N/A", f"{ticker} (Financial): Median EV/EBITDA (P25) should be N/A, got {ws_comp['P25'].value}"
            assert ws_comp['O32'].value == "N/A", f"{ticker} (Financial): Implied EV (O32) should be N/A, got {ws_comp['O32'].value}"
            assert ws_comp['P32'].value == "N/A", f"{ticker} (Financial): Implied EV (P32) should be N/A, got {ws_comp['P32'].value}"
            assert ws_comp['O39'].value == "N/A", f"{ticker} (Financial): Verdict EV/Rev (O39) should be N/A, got {ws_comp['O39'].value}"
            assert ws_comp['P39'].value == "N/A", f"{ticker} (Financial): Verdict EV/EBITDA (P39) should be N/A, got {ws_comp['P39'].value}"
            # P/E must be active
            assert str(ws_comp['Q12'].value).startswith('='), f"{ticker}: Peer 1 P/E (Q12) should be a formula, got {ws_comp['Q12'].value}"
            assert str(ws_comp['Q32'].value).startswith('='), f"{ticker}: Implied Equity P/E (Q32) should be a formula, got {ws_comp['Q32'].value}"
            assert str(ws_comp['Q39'].value).startswith('='), f"{ticker}: Verdict P/E (Q39) should be a formula, got {ws_comp['Q39'].value}"
        else:
            # Non-financials: EV/Revenue, EV/EBITDA, and P/E active
            assert str(ws_comp['O12'].value).startswith('='), f"{ticker} (Non-Fin): Peer 1 EV/Revenue should be a formula"
            assert str(ws_comp['P12'].value).startswith('='), f"{ticker} (Non-Fin): Peer 1 EV/EBITDA should be a formula"
            assert str(ws_comp['Q12'].value).startswith('='), f"{ticker} (Non-Fin): Peer 1 P/E should be a formula"

        # C. DuPont & Altman Narrative Check
        ws_dup = wb['Dupont Analysis']
        ws_alt = wb["Altman's Z Score"]
        assert ws_dup['B2'].value == "='Data Sheet'!B1"
        assert ws_alt['B2'].value == "='Data Sheet'!B1"
        dup_b8 = str(ws_dup['B8'].value or '')
        alt_b8 = str(ws_alt['B8'].value or '')
        assert len(dup_b8) > 10, f"{ticker}: Dupont!B8 is empty!"
        assert len(alt_b8) > 10, f"{ticker}: Altman!B8 is empty!"
        if is_fin:
            assert ws_alt['I89'].value == "Not applicable / insufficient data", f"{ticker} Altman I89 should be 'Not applicable / insufficient data', got {ws_alt['I89'].value}"

        # D. AI Valuation Summary Check
        ws_sum = wb['AI Valuation Summary']
        assert ws_sum['A1'].value == f"{c_name} ({ticker}) - Institutional Valuation"
        assert ws_sum['A5'].value == "=DCF!D44"
        if is_fin:
            assert ws_sum['F5'].value == "Not applicable / insufficient data"
        else:
            assert ws_sum['F5'].value == "='Altman''s Z Score'!I89"

        wb.close()
        print(f"  [OK] Workbook verified for {ticker}: DCF!D44='Data Sheet'!B8, Multiples correct, DuPont/Altman dynamic.")
        contexts.append({'ticker': ticker, 'name': c_name, 'cmp': cmp, 'path': export_path})

    # 5. Strict Cross-Contamination Verification (LICI -> TCS, TCS -> ITC, etc.)
    print("\n=== [4] Cross-Company Contamination Verification ===")
    
    # Specific inspection: LICI vs TCS
    lici_ctx = next(c for c in contexts if c['ticker'] == 'LICI')
    tcs_ctx = next(c for c in contexts if c['ticker'] == 'TCS')
    itc_ctx = next(c for c in contexts if c['ticker'] == 'ITC')

    print("  Checking TCS workbook for any LICI contamination...")
    wb_tcs = openpyxl.load_workbook(tcs_ctx['path'], data_only=False)
    for sheetname in wb_tcs.sheetnames:
        ws = wb_tcs[sheetname]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and isinstance(cell, str):
                    cell_lower = cell.lower()
                    if 'life insurance' in cell_lower or 'lici' in cell_lower:
                        assert False, f"LICI contamination found in TCS sheet '{sheetname}': '{cell}'"
    wb_tcs.close()
    print("  [OK] TCS workbook has 0% LICI contamination.")

    print("  Checking ITC workbook for any TCS contamination...")
    wb_itc = openpyxl.load_workbook(itc_ctx['path'], data_only=False)
    for sheetname in wb_itc.sheetnames:
        ws = wb_itc[sheetname]
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell and isinstance(cell, str):
                    cell_lower = cell.lower()
                    if 'tata consultancy' in cell_lower or 'consultancy services' in cell_lower:
                        assert False, f"TCS contamination found in ITC sheet '{sheetname}': '{cell}'"
    wb_itc.close()
    print("  [OK] ITC workbook has 0% TCS contamination.")

    print("\n>>> ALL FINAL AUDIT TESTS PASSED WITH 100% SUCCESS! <<<")

if __name__ == '__main__':
    test_company_resolution_and_classification()
    test_market_price_validation()
    test_sequential_runs_and_data_isolation()
