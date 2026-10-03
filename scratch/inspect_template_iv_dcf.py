import openpyxl

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)

print("Template sheetnames:", wb.sheetnames)

# Inspect Intrinsic Valuation formulas
if 'Intrinsic Valuation' in wb.sheetnames:
    ws_iv = wb['Intrinsic Valuation']
    print("\n--- master_model_template.xlsx Intrinsic Valuation L35:L45 ---")
    for r in range(35, 46):
        lbl = ws_iv.cell(r, 2).value or ws_iv.cell(r, 1).value
        l_form = ws_iv.cell(r, 12).value
        print(f"Row {r:2d} | Label: {str(lbl)[:30]:30s} | L Form: {l_form}")

# Inspect DCF formulas
if 'DCF' in wb.sheetnames:
    ws_dcf = wb['DCF']
    print("\n--- master_model_template.xlsx DCF Row 8-15 ---")
    for r in range(8, 16):
        c_lbl = ws_dcf.cell(r, 3).value
        h_form = ws_dcf.cell(r, 8).value
        print(f"Row {r:2d} | C Label: {str(c_lbl)[:25]:25s} | H Form: {h_form}")

    print("\n--- master_model_template.xlsx DCF Row 32-46 ---")
    for r in range(32, 47):
        c_lbl = ws_dcf.cell(r, 3).value
        d_form = ws_dcf.cell(r, 4).value
        print(f"Row {r:2d} | C Label: {str(c_lbl)[:30]:30s} | D Form: {d_form}")
