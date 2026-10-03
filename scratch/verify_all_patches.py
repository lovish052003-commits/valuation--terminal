import openpyxl

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)

ws_dcf = wb['DCF']
print('DCF B37:', ws_dcf['B37'].value, '| D37:', ws_dcf['D37'].value)
print('DCF B45:', ws_dcf['B45'].value, '| D45:', ws_dcf['D45'].value)
print('DCF B22:', ws_dcf['B22'].value, '| D22:', ws_dcf['D22'].value)
print('DCF D21:', ws_dcf['D21'].value)
print('DCF D18:', ws_dcf['D18'].value)
print('DCF dates:', [ws_dcf.cell(row=6, column=c).value for c in range(8, 14)])

ws_iv = wb['Intrinsic Valuation']
print('IV Row 38 (H..L):', [ws_iv.cell(row=38, column=c).value for c in range(8, 13)])
print('IV Row 44 (H..L):', [ws_iv.cell(row=44, column=c).value for c in range(8, 13)])

ws_w = wb['WACC']
print('WACC J16:', ws_w['J16'].value, '| E27:', ws_w['E27'].value)

ws_rd = wb['Raw Data']
print('Raw Data T24:', ws_rd['T24'].value)
print('Raw Data Q12..Q16:', [ws_rd.cell(row=r, column=17).value for r in range(12, 17)])

print('Dupont Analysis B3:', wb['Dupont Analysis']['B3'].value)
print("Altman's Z Score B3:", wb["Altman's Z Score"]['B3'].value)
