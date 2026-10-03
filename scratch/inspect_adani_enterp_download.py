import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\LENOVO\Downloads\Adani Enterp.xlsx", data_only=True)
print("Sheet names in Adani Enterp.xlsx:", wb.sheetnames)
if 'Data Sheet' in wb.sheetnames:
    ws = wb['Data Sheet']
    print("B1:", ws['B1'].value)
    print("B16:", ws['B16'].value)
    print("B17:", ws['B17'].value)
    print("K17:", ws['K17'].value)
    print("K30:", ws['K30'].value)
