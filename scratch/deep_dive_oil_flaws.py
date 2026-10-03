import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

adanient_path = r"C:\Users\LENOVO\Downloads\test\ADANIENT_Valuation_Model (8).xlsx"
oil_path = r"C:\Users\LENOVO\Downloads\test\OIL_Valuation_Model.xlsx"

wb_a = openpyxl.load_workbook(adanient_path, data_only=False)
wb_o = openpyxl.load_workbook(oil_path, data_only=False)

wb_ad = openpyxl.load_workbook(adanient_path, data_only=True)
wb_od = openpyxl.load_workbook(oil_path, data_only=True)

print("======================================================================")
print("DEEP DIVE: FINDING ALL FLAWS AND DIVERGENCES IN OIL_VALUATION_MODEL")
print("======================================================================")

# 1. Sheet presence and ordering
print(f"ADANIENT sheets ({len(wb_a.sheetnames)}): {wb_a.sheetnames}")
print(f"OIL sheets      ({len(wb_o.sheetnames)}): {wb_o.sheetnames}")

# 2. Check for None in evaluated values for key sheets
for s_name in ['AI Valuation Summary', 'DCF', 'WACC', 'Comp_Valuation', 'Ratio Analysis', 'Dupont Analysis', "Altman's Z Score"]:
    if s_name in wb_od.sheetnames and s_name in wb_ad.sheetnames:
        ws_o = wb_o[s_name]
        ws_od = wb_od[s_name]
        ws_a = wb_a[s_name]
        ws_ad = wb_ad[s_name]
        
        # Check formulas that have eval=None in OIL but have eval in ADANIENT
        none_evals = []
        formula_diffs = []
        for r in range(1, min(ws_o.max_row + 1, 60)):
            for c in range(1, min(ws_o.max_column + 1, 25)):
                cl = openpyxl.utils.get_column_letter(c)
                fo = str(ws_o.cell(row=r, column=c).value or '')
                vo = ws_od.cell(row=r, column=c).value
                fa = str(ws_a.cell(row=r, column=c).value or '')
                va = ws_ad.cell(row=r, column=c).value
                
                if fo.startswith('=') and vo is None and va is not None:
                    none_evals.append(f"{cl}{r}: form='{fo}', ADANIENT eval={va}")
                    
                # Check for structural formula discrepancies
                if fo.startswith('=') or fa.startswith('='):
                    # Check if references differ in target row (e.g. K85 vs K69, B6 vs D40)
                    if 'K85' in fo or 'B6' == fo:
                        formula_diffs.append(f"{cl}{r}: OIL='{fo}' vs ADANIENT='{fa}'")

        print(f"\n[{s_name}]")
        print(f"  Un-evaluated formulas in OIL: {len(none_evals)}")
        if none_evals[:5]:
            for ne in none_evals[:5]:
                print(f"    {ne}")
        if formula_diffs:
            print(f"  Suspicious/Broken formula references in OIL: {len(formula_diffs)}")
            for fd in formula_diffs:
                print(f"    {fd}")

# 3. Check DCF sheet in OIL vs ADANIENT
ws_dcf_o = wb_o['DCF']
ws_dcf_a = wb_a['DCF']
print("\n--- DCF Sheet Formula Comparison (OIL vs ADANIENT) ---")
for r in range(1, 48):
    lbl_o = ws_dcf_o[f'B{r}'].value
    lbl_a = ws_dcf_a[f'B{r}'].value
    fo = ws_dcf_o[f'D{r}'].value
    fa = ws_dcf_a[f'D{r}'].value
    if fo != fa or lbl_o != lbl_a:
        print(f"Row {r:2d}:")
        print(f"  OIL:      B='{lbl_o}' | D='{fo}'")
        print(f"  ADANIENT: B='{lbl_a}' | D='{fa}'")

# 4. Check Comp_Valuation sheet in OIL vs ADANIENT
ws_comp_o = wb_o['Comp_Valuation']
ws_comp_a = wb_a['Comp_Valuation']
print("\n--- Comp_Valuation Sheet Formula Comparison (OIL vs ADANIENT) ---")
for r in range(22, 43):
    lbl_o = ws_comp_o[f'B{r}'].value
    lbl_a = ws_comp_a[f'B{r}'].value
    fo_o = ws_comp_o[f'O{r}'].value
    fa_o = ws_comp_a[f'O{r}'].value
    if lbl_o != lbl_a or fo_o != fa_o:
        print(f"Row {r:2d}:")
        print(f"  OIL:      B='{lbl_o}' | O='{fo_o}'")
        print(f"  ADANIENT: B='{lbl_a}' | O='{fa_o}'")
