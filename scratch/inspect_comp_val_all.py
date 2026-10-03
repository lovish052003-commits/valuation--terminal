import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['Comp_Valuation']
for r in range(11, 40):
    row_data = []
    for c in range(1, 20):
        v = ws.cell(r, c).value
        if v is not None:
            col_letter = openpyxl.utils.get_column_letter(c)
            row_data.append(f"{col_letter}{r}: {v}")
    if row_data:
        print(f"Row {r:2d} -> " + " | ".join(row_data))
wb.close()
