import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import screener_client
import valuation_engine
import excel_exporter
import openpyxl

print("[1/3] Fetching Screener data for ADANIENT...")
data = screener_client.fetch_company_data('ADANIENT')
print(f"Company: {data['company_name']} | Ticker: {data['ticker']}")
print(f"Sector: {data['sector']} | Industry: {data['industry']}")
print(f"52W High/Low: {data['high_52w']} / {data['low_52w']}")
print(f"About (Wiki): {data['about_company'][:80]}...")
print(f"Recent Updates count: {len(data['recent_updates'])}")
print(f"Peers DF count: {len(data['peers_df'])}")

print("\n[2/3] Calculating valuation...")
val_res = valuation_engine.calculate_valuation(data)

print("\n[3/3] Exporting to Excel model...")
excel_path = excel_exporter.export_valuation_model(data, val_res, "")
print(f"Exported to: {excel_path}")

print("\n=== VERIFYING GENERATED EXCEL WORKBOOK ===")
wb = openpyxl.load_workbook(excel_path, data_only=True)

# 1. Check Raw FS peers rows 56 to 66
if 'Raw FS' in wb.sheetnames:
    ws_raw = wb['Raw FS']
    print("\n--- Raw FS Rows 56-66 ---")
    for r in range(56, 67):
        p_name = ws_raw.cell(r, 12).value
        p_cmp = ws_raw.cell(r, 13).value
        p_mcap = ws_raw.cell(r, 44).value
        print(f"  Row {r:2d}: {repr(p_name)} | CMP: {p_cmp} | Mcap: {p_mcap}")

# 2. Check Comp_Valuation
if 'Comp_Valuation' in wb.sheetnames:
    ws_comp = wb['Comp_Valuation']
    print("\n--- Comp_Valuation Rows 12-21 ---")
    for r in range(12, 22):
        c_name = ws_comp.cell(r, 2).value
        p_val = ws_comp.cell(r, 4).value
        print(f"  Row {r:2d}: {repr(c_name)} | Share Price: {p_val}")
    print("  B30 (Title):", ws_comp['B30'].value)
    print("  O25 (Median EV/Rev):", ws_comp['O25'].value)
    print("  P25 (Median EV/EBITDA):", ws_comp['P25'].value)

# 3. Check Dupont Analysis
if 'Dupont Analysis' in wb.sheetnames:
    ws_dup = wb['Dupont Analysis']
    print("\n--- Dupont Analysis ---")
    print("  B5 (52W Range):", repr(ws_dup['B5'].value))
    print("  B8 (About Co):", repr(str(ws_dup['B8'].value)[:100]))
    for r in [37, 39, 41, 43, 45]:
        print(f"  B{r} (Update):", repr(str(ws_dup.cell(r, 2).value)[:80]))

# 4. Check Altman's Z Score
if "Altman's Z Score" in wb.sheetnames:
    ws_alt = wb["Altman's Z Score"]
    print("\n--- Altman's Z Score ---")
    print("  B5 (52W Range):", repr(ws_alt['B5'].value))
    print("  B8 (About Co):", repr(str(ws_alt['B8'].value)[:100]))
    for r in [36, 38, 40, 42, 44]:
        print(f"  B{r} (Update):", repr(str(ws_alt.cell(r, 2).value)[:80]))

print("\nALL VERIFICATIONS COMPLETE!")
