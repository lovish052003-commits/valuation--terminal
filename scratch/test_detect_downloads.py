import os, sys, gc
import win32com.client, pythoncom
import pandas as pd
from datetime import datetime
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data, clean_num

def run_test_export():
    symbol = 'NESTLEIND'
    screener_data = fetch_company_data(symbol)
    
    # Check if there is a downloaded Screener file in Downloads
    import glob
    downloads_dir = os.path.expanduser(r'~\Downloads')
    matching_files = glob.glob(os.path.join(downloads_dir, '*Nestle*India*.xlsx'))
    print("Found downloaded Screener files in Downloads:", matching_files)
    
    # Let's inspect the latest one
    if matching_files:
        latest_file = max(matching_files, key=os.path.getmtime)
        print("Using latest Screener export:", latest_file)
        wb_screener = openpyxl.load_workbook(latest_file, data_only=True)
        if 'Data Sheet' in wb_screener.sheetnames:
            ws_sd = wb_screener['Data Sheet']
            print("Row 93 from downloaded file:", [ws_sd.cell(93, c).value for c in range(2, 12)])
            print("Row 70 from downloaded file:", [ws_sd.cell(70, c).value for c in range(2, 12)])
            print("Row 67 from downloaded file:", [ws_sd.cell(67, c).value for c in range(2, 12)])
            print("Row 68 from downloaded file:", [ws_sd.cell(68, c).value for c in range(2, 12)])
            print("Row 69 from downloaded file:", [ws_sd.cell(69, c).value for c in range(2, 12)])
            print("Row 90 from downloaded file:", [ws_sd.cell(90, c).value for c in range(2, 12)])

run_test_export()
