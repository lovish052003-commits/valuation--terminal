import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
import openpyxl
import screener_client
import valuation_engine
import excel_exporter
from excel_exporter import get_effective_peers, export_valuation_model, get_sector_key

def test_vbl_peers_and_comp_valuation():
    print("=== 1. FETCHING SCREENER DATA FOR VBL ===")
    sd = screener_client.fetch_company_data("VBL")
    print(f"Company: {sd.get('company_name')} | Ticker: {sd.get('ticker')} | MCap: {sd.get('market_cap_cr')}")
    
    sec_key = get_sector_key(sd)
    print(f"Sector Key: {sec_key}")
    
    print("\n=== 2. TESTING GET_EFFECTIVE_PEERS ===")
    peers = get_effective_peers(sd)
    print(f"Total Peers Returned: {len(peers)}")
    for i, p in enumerate(peers, 1):
        print(f"  {i}. {p['name']} ({p.get('ticker')}) - MCap: Rs. {p.get('mcap'):,.1f} Cr | CMP: Rs. {p.get('cmp'):,.1f}")
        
    penny_stocks = ['valencia', 'orient beverages', 'ans industries', 'transglobe']
    found_penny = [p['name'] for p in peers if any(ps in p['name'].lower() for ps in penny_stocks)]
    assert not found_penny, f"Error: Penny stocks found in peer group: {found_penny}"
    print("[PASS] No penny stocks found in peer group!")
    
    # Check self-exclusion
    found_self = [p['name'] for p in peers if 'varun' in p['name'].lower() or p.get('ticker') == 'VBL']
    assert not found_self, f"Error: Varun Beverages found in its own peer group: {found_self}"
    print("[PASS] Target company strictly excluded from peer group!")
    
    print("\n=== 3. TESTING VALUATION ENGINE ===")
    val = valuation_engine.calculate_valuation(sd)
    print(f"Methodology: {val.get('methodology')}")
    
    print("\n=== 4. EXPORTING MODEL & VERIFYING COMP_VALUATION SHEET ===")
    out_dir = r"C:\Users\LENOVO\Downloads\Advance Financial Project\scratch\test_output"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "VBL_Test_Comp_Valuation.xlsx")
    
    # Export model (returns path)
    out_path = excel_exporter.export_valuation_model(sd, val)
    print(f"Generated Excel Model at: {out_path}")
    
    wb = openpyxl.load_workbook(out_path, data_only=False)
    wb_val = openpyxl.load_workbook(out_path, data_only=True)
    ws = wb['Comp_Valuation']
    ws_v = wb_val['Comp_Valuation']
    
    print("\n--- Comp_Valuation Rows 12 to 20 (Peers & Ratios) ---")
    for r in range(12, 12 + len(peers)):
        comp = ws.cell(r, 2).value
        ev_rev_f = ws.cell(r, 9).value
        ev_rev_v = ws_v.cell(r, 9).value
        ev_eb_f = ws.cell(r, 10).value
        ev_eb_v = ws_v.cell(r, 10).value
        print(f"Row {r:02d}: {str(comp):25s} | EV/Rev(Col I): {str(ev_rev_f):25s} (Val={ev_rev_v}) | EV/EBITDA(Col J): {str(ev_eb_f):25s} (Val={ev_eb_v})")
        
    print("\n--- Comp_Valuation Rows 23 to 28 (Summary Multiples) ---")
    for r in range(23, 29):
        metric = ws.cell(r, 2).value
        print(f"Row {r}: {metric:15s} | Col I (EV/Rev): {ws.cell(r, 9).value} (Val={ws_v.cell(r, 9).value}) | Col J (EV/EBITDA): {ws.cell(r, 10).value} (Val={ws_v.cell(r, 10).value}) | Col Q (P/E): {ws.cell(r, 17).value} (Val={ws_v.cell(r, 17).value})")
        
    print("\n--- Comp_Valuation Rows 30 to 39 (Implied Valuation Bridge) ---")
    for r in [30, 32, 33, 34, 35, 37, 39]:
        label = ws.cell(r, 2).value
        col_i = f"{ws.cell(r, 9).value} (Val={ws_v.cell(r, 9).value})"
        col_j = f"{ws.cell(r, 10).value} (Val={ws_v.cell(r, 10).value})"
        col_q = f"{ws.cell(r, 17).value} (Val={ws_v.cell(r, 17).value})"
        print(f"Row {r:02d} [{str(label):25s}]: Col I={col_i} | Col J={col_j} | Col Q={col_q}")

if __name__ == '__main__':
    test_vbl_peers_and_comp_valuation()
