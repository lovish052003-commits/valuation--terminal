import os
import sys
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
import screener_client
import valuation_engine

data = screener_client.fetch_company_data('TATASTEEL')
val = valuation_engine.calculate_valuation(data)

print("=== TERMINAL / ENGINE VALUES FOR TATASTEEL ===")
print("CMP:", val.get('current_price'))
print("Intrinsic Value:", val.get('intrinsic_value'))
print("DCF Value:", val.get('dcf_value'))
print("EV:", val.get('enterprise_value'))
print("Equity Value:", val.get('equity_value'))
print("Cash:", val.get('cash_estimate'))
print("Debt:", val.get('total_debt'))
print("Shares (Cr):", val.get('shares_cr'))
print("WACC:", val.get('wacc'))
print("Margin of Safety:", val.get('margin_of_safety'))
print("Verdict:", val.get('verdict'))

excel_path = os.path.abspath('exports/TATASTEEL_Valuation_Model.xlsx')
if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if 'DCF' in wb.sheetnames:
        ws_dcf = wb['DCF']
        print("\n=== EXCEL DCF SHEET FOR TATASTEEL ===")
        print("D44 (CMP):", ws_dcf['D44'].value)
        print("D42 (Intrinsic Value/Share):", ws_dcf['D42'].value)
        print("D35 (Enterprise Value):", ws_dcf['D35'].value)
        print("D39 (Equity Value):", ws_dcf['D39'].value)
        print("D37 (Cash):", ws_dcf['D37'].value)
        print("D38 (Debt):", ws_dcf['D38'].value)
        print("D40 (Shares Cr):", ws_dcf['D40'].value)
        print("D20 (WACC):", ws_dcf['D20'].value)
    
    if 'AI Valuation Summary' in wb.sheetnames:
        ws_sum = wb['AI Valuation Summary']
        print("\n=== EXCEL AI VALUATION SUMMARY FOR TATASTEEL ===")
        print("A5 (CMP):", ws_sum['A5'].value)
        print("B5 (Intrinsic Value):", ws_sum['B5'].value)
        print("C5 (Margin of Safety):", ws_sum['C5'].value)
        print("D5 (Verdict):", ws_sum['D5'].value)
        print("E5 (WACC):", ws_sum['E5'].value)
        print("F5 (Altman):", ws_sum['F5'].value)
        print("G5 (DuPont):", ws_sum['G5'].value)
else:
    print(f"File not found: {excel_path}")
