import os
import sys
import json
import openpyxl
import pandas as pd

sys.path.insert(0, os.path.abspath('.'))
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

def run_test(ticker):
    print(f"\n========================================================")
    print(f"RUNNING TEST FOR TICKER: {ticker}")
    print(f"========================================================")
    
    p = os.path.join('exports', '.screener_cache', f'{ticker}_data.json')
    if not os.path.exists(p):
        print(f"Cache file not found for {ticker}")
        return
        
    with open(p, 'r', encoding='utf-8') as f:
        data = json.load(f)
    if 'tables' in data:
        for tb_k, tb_v in data['tables'].items():
            if isinstance(tb_v, dict) and 'columns' in tb_v and 'data' in tb_v:
                data['tables'][tb_k] = pd.DataFrame(tb_v['data'], columns=tb_v['columns'])
                
    # 1. Run valuation engine
    val_res = calculate_valuation(data)
    print(f"Valuation Engine Results:")
    print(f"  CMP: Rs. {val_res['current_price']} | Intrinsic Value: Rs. {val_res['intrinsic_value_per_share']}")
    print(f"  Normalized ROIC: {val_res['normalized_roic']}%")
    print(f"  Expected Growth: {val_res['expected_growth_rate']}% (Source: {val_res['growth_source']})")
    print(f"  Fundamental Reinvestment Rate: {val_res['fundamental_reinvestment_rate']}%")
    print(f"  Historical Median Reinvestment Rate: {val_res['historical_median_reinvestment_rate']}%")
    print(f"  Terminal ROIC: {val_res['terminal_roic']}% | Terminal Reinvest: {val_res['terminal_reinvestment_rate']}%")
    print(f"  Forecast Reinvestment (Y1-Y5): {[x['reinvestment_rate'] for x in val_res['dcf_table']]}")
    print(f"  Consistency Check: {val_res['growth_roic_consistency']} | Confidence: {val_res['reinvestment_confidence']}")
    if val_res['reinvestment_warnings']:
        print(f"  Warnings: {val_res['reinvestment_warnings']}")
        
    # 2. Export Model to Excel
    export_path = os.path.join('exports', f'{ticker}_Fundamental_Test_Model.xlsx')
    if os.path.exists(export_path):
        try:
            os.remove(export_path)
        except:
            pass
            
    final_path = export_valuation_model(data, val_res, report_markdown="")
    print(f"Export returned path: {final_path}")
    
    # 3. Audit Excel Cells
    wb = openpyxl.load_workbook(final_path, data_only=False)
    print(f"\nAuditing Workbook Cells in {final_path}:")
    
    if 'Intrinsic Valuation' in wb.sheetnames:
        iv = wb['Intrinsic Valuation']
        print(f"  Intrinsic Valuation Sheet:")
        for r in range(67, 75):
            print(f"    Row {r}: B{r}='{iv[f'B{r}'].value}' | L{r}='{iv[f'L{r}'].value}' (num_fmt: {iv[f'L{r}'].number_format})")
            
    if 'DCF' in wb.sheetnames:
        dcf = wb['DCF']
        print(f"  DCF Sheet:")
        print(f"    D18 (Growth): {dcf['D18'].value}")
        print(f"    D19 (Terminal Growth): {dcf['D19'].value}")
        print(f"    D20 (WACC): {dcf['D20'].value}")
        print(f"    D21 (Terminal Reinvest): {dcf['D21'].value}")
        print(f"    I11 (Reinvest Y1): {dcf['I11'].value}")
        print(f"    J11 (Reinvest Y2): {dcf['J11'].value}")
        print(f"    K11 (Reinvest Y3): {dcf['K11'].value}")
        print(f"    L11 (Reinvest Y4): {dcf['L11'].value}")
        print(f"    M11 (Reinvest Y5): {dcf['M11'].value}")
        for col in ['I', 'J', 'K', 'L', 'M']:
            print(f"    {col}12 (FCFF): {dcf[f'{col}12'].value}")
            
    if 'AI Valuation Summary' in wb.sheetnames:
        ais = wb['AI Valuation Summary']
        print(f"  AI Valuation Summary Sheet (DCF Forecast Rows 22-26):")
        for r in range(22, 27):
            print(f"    Row {r}: A='{ais[f'A{r}'].value}' | D (Reinvest)='{ais[f'D{r}'].value}' | E (FCFF)='{ais[f'E{r}'].value}'")
            
    # Also evaluate via openpyxl data_only=True to check evaluated values if COM calculated them
    wb_eval = openpyxl.load_workbook(final_path, data_only=True)
    if 'Intrinsic Valuation' in wb_eval.sheetnames:
        iv_ev = wb_eval['Intrinsic Valuation']
        print(f"\n  Evaluated Values:")
        print(f"    L67 (Norm ROIC): {iv_ev['L67'].value}")
        print(f"    L68 (Growth): {iv_ev['L68'].value}")
        print(f"    L69 (Fund Reinvest): {iv_ev['L69'].value}")
        print(f"    L70 (Terminal ROIC): {iv_ev['L70'].value}")
        print(f"    L71 (Terminal Reinvest): {iv_ev['L71'].value}")
        print(f"    L72 (Consistency): {iv_ev['L72'].value}")
        print(f"    L73 (Source): {iv_ev['L73'].value}")
        print(f"    L74 (Confidence): {iv_ev['L74'].value}")
    if 'DCF' in wb_eval.sheetnames:
        dcf_ev = wb_eval['DCF']
        print(f"    DCF D21: {dcf_ev['D21'].value}")
        print(f"    DCF I11: {dcf_ev['I11'].value}")
        print(f"    DCF M11: {dcf_ev['M11'].value}")
        print(f"    DCF Intrinsic Value per share (D42): {dcf_ev['D42'].value}")
        
    wb.close()
    wb_eval.close()
    print("AUDIT FINISHED.")

if __name__ == '__main__':
    for tk in ['TATASTEEL', 'INFY']:
        run_test(tk)
