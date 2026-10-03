import openpyxl

paths = [
    r"C:\Users\LENOVO\Downloads\test\ADANIENT_Valuation_Model (8).xlsx",
    r"C:\Users\LENOVO\Downloads\Advance Financial Project\exports\SBIN_Valuation_Model.xlsx"
]

for p in paths:
    print(f"\n=======================================================")
    print(f"RAW FS ROW 56 IN: {p}")
    print(f"=======================================================")
    wb_v = openpyxl.load_workbook(p, data_only=True)
    wb_f = openpyxl.load_workbook(p, data_only=False)
    if 'Raw FS' in wb_v.sheetnames:
        ws_v = wb_v['Raw FS']
        ws_f = wb_f['Raw FS']
        # Check header row 55 or 10
        for col_idx in range(1, 55):
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            val = ws_v.cell(56, col_idx).value
            form = ws_f.cell(56, col_idx).value
            h_val = ws_v.cell(55, col_idx).value or ws_v.cell(10, col_idx).value or ws_v.cell(11, col_idx).value
            if any([val, form, h_val]):
                print(f"Col {col_letter:3s} ({col_idx:2d}) | Header: {str(h_val)[:20]:20s} | Val: {val} | Form: {form}")
