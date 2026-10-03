import openpyxl

paths = [
    r"C:\Users\LENOVO\Downloads\test\SBIN_Valuation_Model (1).xlsx",
    r"C:\Users\LENOVO\Downloads\test\SBIN_Valuation_Model (2).xlsx",
    r"C:\Users\LENOVO\Downloads\test\SBIN_Valuation_Model.xlsx"
]

for p in paths:
    print(f"\nChecking {p}:")
    try:
        wb = openpyxl.load_workbook(p, data_only=True)
        if 'Data Sheet' in wb.sheetnames:
            print("  Data Sheet B8:", wb['Data Sheet']['B8'].value)
            print("  Data Sheet B6 (shares):", wb['Data Sheet']['B6'].value)
            print("  Data Sheet B9 (mcap):", wb['Data Sheet']['B9'].value)
        if 'AI Valuation Summary' in wb.sheetnames:
            print("  AI Val Summary A5:", wb['AI Valuation Summary']['A5'].value)
            print("  AI Val Summary B5:", wb['AI Valuation Summary']['B5'].value)
            print("  AI Val Summary Row 8-11:", [wb['AI Valuation Summary'].cell(r, 2).value for r in range(8, 12)])
    except Exception as e:
        print("  Error:", e)
