import openpyxl

target_path = r'C:\Users\LENOVO\Downloads\test\SUNPHARMA_Valuation_Model (1).xlsx'
wb_f = openpyxl.load_workbook(target_path, data_only=False)
wb_d = openpyxl.load_workbook(target_path, data_only=True)

ws_f = wb_f['DCF']
ws_d = wb_d['DCF']

out_file = r'C:\Users\LENOVO\Downloads\Advance Financial Project\scratch\dcf_full_audit.txt'

with open(out_file, 'w', encoding='utf-8') as f:
    f.write("=== DCF SHEET FULL GRID (Rows 5 to 46, Cols B to M) ===\n")
    for r in range(5, 47):
        labels = ws_f.cell(row=r, column=2).value
        f.write(f"Row {r:2d} | B: {str(labels):<30}\n")
        f.write(f"   Formulas D..M: {[ws_f.cell(row=r, column=c).value for c in range(4, 14)]}\n")
        f.write(f"   Values   D..M: {[ws_d.cell(row=r, column=c).value for c in range(4, 14)]}\n")

print("DCF audit written to scratch/dcf_full_audit.txt")
