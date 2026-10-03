import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx')
ws = wb['Dupont Analysis']
print("=== DUPONT ANALYSIS ===")
for r in [8, 37, 39, 41, 43, 45]:
    val = ws.cell(r, 2).value
    print(f"B{r}: {str(val).encode('ascii', 'replace').decode('ascii')}")

ws2 = wb["Altman's Z Score"]
print("\n=== ALTMAN'S Z SCORE ===")
for r in [8, 36, 38, 40, 42, 44]:
    val = ws2.cell(r, 2).value
    print(f"B{r}: {str(val).encode('ascii', 'replace').decode('ascii')}")

print("\n--- Merged cells in Dupont Analysis near B8 and B37 ---")
for rng in ws.merged_cells.ranges:
    s = str(rng)
    if any(s.startswith(f'B{r}:') for r in [8, 37, 39, 41, 43, 45]):
        print(rng)

print("\n--- Merged cells in Altman's Z Score near B8 and B36 ---")
for rng in ws2.merged_cells.ranges:
    s = str(rng)
    if any(s.startswith(f'B{r}:') for r in [8, 36, 38, 40, 42, 44]):
        print(rng)
