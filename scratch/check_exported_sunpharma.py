import openpyxl, os

path = r'c:\Users\LENOVO\Downloads\Advance Financial Project\exports\SUNPHARMA_Valuation_Model.xlsx'
if not os.path.exists(path):
    print("File does not exist:", path)
else:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Data Sheet']
    print("=== Data Sheet Rows 16 to 34 ===")
    for r in range(16, 35):
        label = ws.cell(r, 1).value
        vals = [ws.cell(r, c).value for c in range(2, 12)]
        print(f"Row {r:2d} | {str(label):25s} | {vals}")
