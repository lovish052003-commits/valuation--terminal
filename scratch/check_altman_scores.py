import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
wb_v = openpyxl.load_workbook('ITC Model.xlsx', data_only=True)
ws = wb["Altman's Z Score"]
ws_v = wb_v["Altman's Z Score"]

for r in range(86, 92):
    for c in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I']:
        f = ws[f"{c}{r}"].value
        v = ws_v[f"{c}{r}"].value
        if f is not None or v is not None:
            print(f"{c}{r}: formula={f} | val={v}")
