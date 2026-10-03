import openpyxl, glob, os

files = glob.glob(r'c:\Users\LENOVO\Downloads\*.xlsx')
for f in files:
    fname = os.path.basename(f)
    if 'Valuation_Model' in fname:
        continue
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        if 'Data Sheet' in wb.sheetnames:
            ws = wb['Data Sheet']
            cname = ws.cell(1, 2).value
            fv = ws.cell(7, 2).value
            mcap = ws.cell(9, 2).value
            sh_cr = ws.cell(6, 2).value
            r57 = ws.cell(57, 11).value # latest eq capital
            r70 = ws.cell(70, 11).value # latest no of shares
            r72 = ws.cell(72, 11).value # latest face value
            r93 = ws.cell(93, 11).value # latest adjusted shares
            r93_hist = [ws.cell(93, c).value for c in range(2, 12)]
            print(f"\nFile: {fname} | Company: {cname}")
            print(f"  Latest Eq Capital (Row 57): {r57}, Latest Shares (Row 70): {r70}, Latest FV (Row 72): {r72}")
            print(f"  Row 93 Adjusted Shares (10 yrs): {r93_hist}")
    except Exception as e:
        pass
