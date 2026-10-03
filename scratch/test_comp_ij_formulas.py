import sys, os
sys.path.insert(0, os.path.abspath('.'))
import openpyxl
import excel_exporter
import screener_client
import valuation_engine

wb = openpyxl.load_workbook('ITC Model.xlsx')
screener_data = screener_client.fetch_company_data('SUZLON')
val_result = valuation_engine.calculate_valuation(screener_data)

excel_exporter.update_comp_valuation_sheet_openpyxl(wb, screener_data, val_result)
ws = wb['Comp_Valuation']

for r in range(12, 17):
    val_i = ws[f'I{r}'].value
    val_j = ws[f'J{r}'].value
    print(f'Row {r}: I = {val_i} | J = {val_j}')
    assert val_i == f'=IF(G{r}>0, F{r}/G{r}, "N/A")', f'Mismatch I{r}: {val_i}'
    assert val_j == f'=IF(H{r}>0, F{r}/H{r}, "N/A")', f'Mismatch J{r}: {val_j}'

print('Median I25:', ws['I25'].value)
print('Median J25:', ws['J25'].value)
assert 'MEDIAN' in str(ws['I25'].value)
assert 'MEDIAN' in str(ws['J25'].value)
print('ALL VERIFICATION CHECKS PASSED!')
