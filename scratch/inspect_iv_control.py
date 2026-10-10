import os
import openpyxl

def main():
    p = 'exports/ITC_Valuation_Model.xlsx'
    wb = openpyxl.load_workbook(p, data_only=False)
    wb_v = openpyxl.load_workbook(p, data_only=True)

    for sheetname in ['Intrinsic Valuation', 'Control']:
        if sheetname not in wb.sheetnames:
            print(f"Sheet {sheetname} not found")
            continue
        ws = wb[sheetname]
        ws_v = wb_v[sheetname]
        print(f'=== {sheetname} ===')
        for r in [8, 9, 10, 11, 12, 14, 20, 21, 22, 23, 24, 25, 26, 38, 40, 55, 64]:
            for c in ['B', 'C', 'K', 'L']:
                cell = f'{c}{r}'
                val = ws[cell].value
                val_v = ws_v[cell].value
                if val is not None or val_v is not None:
                    print(f'{cell}: formula={val} | val={val_v}')

if __name__ == '__main__':
    main()
