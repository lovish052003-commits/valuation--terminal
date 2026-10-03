import os
import sys
import openpyxl
import re

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

adanient_path = r"C:\Users\LENOVO\Downloads\test\ADANIENT_Valuation_Model (8).xlsx"
oil_path = r"C:\Users\LENOVO\Downloads\test\OIL_Valuation_Model.xlsx"

print("=================================================================")
print("COMPARING GOLD-STANDARD (ADANIENT) VS PROBLEMATIC (OIL)")
print("=================================================================")

for name, path in [("ADANIENT (Gold Standard)", adanient_path), ("OIL (Problematic)", oil_path)]:
    print(f"\n--- Inspecting {name} ---")
    if not os.path.exists(path):
        print(f"File NOT found: {path}")
        continue
    
    wb = openpyxl.load_workbook(path, data_only=False)
    wb_data = openpyxl.load_workbook(path, data_only=True)
    
    print(f"Sheet count: {len(wb.sheetnames)}")
    print(f"Sheets: {wb.sheetnames}")
    
    # 1. Error scan across all sheets
    error_tokens = ['#REF!', '#DIV/0!', '#VALUE!', '#NAME?', '#NUM!', '#NULL!']
    errors_found = {}
    for s_name in wb_data.sheetnames:
        ws_d = wb_data[s_name]
        ws_f = wb[s_name]
        sheet_errs = []
        for r in range(1, min(ws_d.max_row + 1, 150)):
            for c in range(1, min(ws_d.max_column + 1, 30)):
                col_letter = openpyxl.utils.get_column_letter(c)
                v = ws_d.cell(row=r, column=c).value
                f = ws_f.cell(row=r, column=c).value
                if isinstance(v, str):
                    for tok in error_tokens:
                        if tok in v:
                            sheet_errs.append(f"{col_letter}{r}: evaluated='{v}', formula='{f}'")
                            break
        if sheet_errs:
            errors_found[s_name] = sheet_errs
            
    if errors_found:
        print(f"ERRORS DETECTED in {len(errors_found)} sheets:")
        for s_name, errs in errors_found.items():
            print(f"  [{s_name}] ({len(errs)} errors):")
            for e in errs[:10]:
                print(f"    {e}")
            if len(errs) > 10:
                print(f"    ... and {len(errs) - 10} more")
    else:
        print("Zero evaluated formula errors detected.")

    # 2. Check AI Valuation Summary
    if 'AI Valuation Summary' in wb.sheetnames:
        ws = wb['AI Valuation Summary']
        ws_d = wb_data['AI Valuation Summary']
        print("\nAI Valuation Summary Key Cells:")
        print(f"  A1: {ws['A1'].value}")
        for cell_coord in ['A4', 'B4', 'C4', 'D4', 'E4', 'F4', 'G4',
                           'A5', 'B5', 'C5', 'D5', 'E5', 'F5', 'G5']:
            print(f"  {cell_coord}: formula={repr(ws[cell_coord].value)}, eval={repr(ws_d[cell_coord].value)}")
        print(f"  A7: {ws['A7'].value}")
        print(f"  A8: {ws['A8'].value}")
        print(f"  B8: {repr(ws['B8'].value)[:80]}...")
        print(f"  A13: {ws['A13'].value}")
        print(f"  A21: {ws['A21'].value}")
        for r in range(21, 33):
            lbl = ws[f'A{r}'].value
            val = ws[f'B{r}'].value
            val_d = ws_d[f'B{r}'].value
            if lbl or val:
                print(f"    A{r}: '{lbl}' | B{r}: formula={repr(val)}, eval={repr(val_d)}")

    # 3. Check DCF sheet bridge
    if 'DCF' in wb.sheetnames:
        ws = wb['DCF']
        ws_d = wb_data['DCF']
        print("\nDCF Sheet Valuation Bridge:")
        for r in range(30, 48):
            lbl = ws[f'B{r}'].value
            val = ws[f'D{r}'].value
            val_d = ws_d[f'D{r}'].value
            if lbl or val:
                print(f"  Row {r}: '{lbl}' -> formula={repr(val)}, eval={repr(val_d)}")

    # 4. Check Comp_Valuation
    if 'Comp_Valuation' in wb.sheetnames:
        ws = wb['Comp_Valuation']
        ws_d = wb_data['Comp_Valuation']
        print("\nComp_Valuation Peers & Summary:")
        print(f"  Target row (B12): {repr(ws['B12'].value)}")
        for r in range(12, 22):
            p_name = ws[f'B{r}'].value
            p_ev_rev = ws[f'O{r}'].value
            p_ev_ebitda = ws[f'P{r}'].value
            p_pe = ws[f'Q{r}'].value
            if p_name:
                print(f"    Row {r}: {p_name} | O={p_ev_rev} | P={p_ev_ebitda} | Q={p_pe}")
        print("  Summary statistics:")
        for r in range(22, 28):
            lbl = ws[f'B{r}'].value
            o = ws[f'O{r}'].value
            p = ws[f'P{r}'].value
            q = ws[f'Q{r}'].value
            print(f"    Row {r}: '{lbl}' | O={o} | P={p} | Q={q}")
        print("  Target valuation bridge:")
        for r in range(32, 43):
            lbl = ws[f'B{r}'].value
            o = ws[f'O{r}'].value
            p = ws[f'P{r}'].value
            q = ws[f'Q{r}'].value
            if lbl or o or p or q:
                print(f"    Row {r}: '{lbl}' | O={o} | P={p} | Q={q}")

    wb.close()
    wb_data.close()
