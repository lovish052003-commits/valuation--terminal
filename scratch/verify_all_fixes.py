import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

file_path = 'exports/HINDUNILVR_Valuation_Model.xlsx'
wb = openpyxl.load_workbook(file_path, data_only=False)

print("=" * 60)
print("1. CASH & LIQUID INVESTMENTS")
print("=" * 60)
ws_ds = wb['Data Sheet']
ws_dcf = wb['DCF']
print(f"Data Sheet K69 (Cash & Bank): {ws_ds['K69'].value}")
print(f"Data Sheet K64 (Investments): {ws_ds['K64'].value}")
print(f"DCF B37 Label: {ws_dcf['B37'].value}")
print(f"DCF D37 Formula: {ws_dcf['D37'].value}")

print("\n" + "=" * 60)
print("2. TERMINAL ROIC & REINVESTMENT RATE")
print("=" * 60)
ws_iv = wb['Intrinsic Valuation']
print(f"Intrinsic Valuation B40 Label: {ws_iv['B40'].value}")
print(f"Intrinsic Valuation L38 (Operating EBIT): {ws_iv['L38'].value}")
print(f"Intrinsic Valuation L37 (Invested Capital): {ws_iv['L37'].value}")
print(f"Intrinsic Valuation L40 (After-Tax ROIC formula): {ws_iv['L40'].value}")
print(f"DCF B22 Label: {ws_dcf['B22'].value}")
print(f"DCF D22 (Terminal ROIC formula): {ws_dcf['D22'].value}")
print(f"DCF D21 (Terminal Reinvestment Rate formula): {ws_dcf['D21'].value}")

print("\n" + "=" * 60)
print("3. REINVESTMENT VIA DELTA INVESTED CAPITAL & GROWTH")
print("=" * 60)
print(f"Intrinsic Valuation B51: {ws_iv['B51'].value}")
print(f"Intrinsic Valuation L51 (Delta IC): {ws_iv['L51'].value}")
print(f"Intrinsic Valuation L52 (Reinvestment Rate formula): {ws_iv['L52'].value}")
print(f"Intrinsic Valuation L55 (Median Reinvestment Rate): {ws_iv['L55'].value}")
print(f"Intrinsic Valuation L62 (Fundamental Growth formula): {ws_iv['L62'].value}")
print(f"Intrinsic Valuation L65 (Median Fundamental Growth): {ws_iv['L65'].value}")
print(f"DCF D18 (Forecast Growth Rate formula): {ws_dcf['D18'].value}")

print("\n" + "=" * 60)
print("4. TAX RATE")
print("=" * 60)
ws_rd = wb['Raw Data']
ws_wacc = wb['WACC']
print(f"Raw Data T24: {ws_rd['T24'].value}")
print(f"WACC E27: {ws_wacc['E27'].value}")
print(f"Intrinsic Valuation L48: {ws_iv['L48'].value}")

print("\n" + "=" * 60)
print("5. PEERS & TARGET WACC CAPITAL STRUCTURE")
print("=" * 60)
print("Raw Data Rows 12-16 (Atomically written peer data):")
for r in range(12, 17):
    p_name = ws_rd.cell(r, 15).value
    p_cmp = ws_rd.cell(r, 16).value
    p_mcap = ws_rd.cell(r, 18).value
    p_debt = ws_rd.cell(r, 19).value
    p_cash = ws_rd.cell(r, 20).value
    print(f"  Row {r}: Name={p_name} | CMP={p_cmp} | MCap={p_mcap} | Debt={p_debt} | Cash={p_cash}")

print(f"\nWACC Target Row 16 Beta: {ws_wacc['J16'].value}")
print(f"WACC C34 (Target Debt): {ws_wacc['C34'].value}")
print(f"WACC C35 (Target Equity): {ws_wacc['C35'].value}")
print(f"WACC E34 (Debt Weight): {ws_wacc['E34'].value}")
print(f"WACC E35 (Equity Weight): {ws_wacc['E35'].value}")
print(f"WACC E38 (Target D/E): {ws_wacc['E38'].value}")

print("\n" + "=" * 60)
print("6. VALUATION DATE & DISCOUNTING STUB")
print("=" * 60)
print(f"DCF Row 13 Discount Periods: {[ws_dcf.cell(13, c).value for c in range(9, 14)]}")
print(f"DCF Row 14 Discount Factors: {[ws_dcf.cell(14, c).value for c in range(9, 14)]}")
print(f"DCF D34 (PV of Terminal Value): {ws_dcf['D34'].value}")

print("\n" + "=" * 60)
print("7. SMALLER ISSUES: DCF D45, ALTMAN Z, DUPONT ROE")
print("=" * 60)
print(f"DCF B45 Label: {ws_dcf['B45'].value}")
print(f"DCF D45 Formula: {ws_dcf['D45'].value}")
ws_az = wb["Altman's Z Score"]
print(f"Altman Z Row 77 (Market Cap formula): {ws_az['I77'].value}")
print(f"Altman Z Row 78 (Total Borrowings formula): {ws_az['I78'].value}")
print(f"Data Sheet K90 (Stock Price): {ws_ds['K90'].value}")
print(f"Data Sheet Row 93 (Shares in Cr): {[ws_ds.cell(93, c).value for c in range(2, 12)]}")

print("\n" + "=" * 60)
print("8. AI VALUATION SUMMARY SHEET INTEGRITY")
print("=" * 60)
ws_sum = wb['AI Valuation Summary']
for r in range(1, 49):
    row_vals = [f"{openpyxl.utils.get_column_letter(c)}{r}: {repr(ws_sum.cell(r, c).value)}" for c in range(1, 9) if ws_sum.cell(r, c).value is not None]
    if row_vals:
        print(" | ".join(row_vals))
