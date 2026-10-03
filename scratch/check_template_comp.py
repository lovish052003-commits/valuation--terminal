import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
if 'Comp_Valuation' in wb.sheetnames:
    ws = wb['Comp_Valuation']
    print("=== ITC Model.xlsx Comp_Valuation ===")
    for r in range(9, 23):
        vals = [f"{ws.cell(r, c).coordinate}={ws.cell(r, c).value}" for c in range(2, 18) if ws.cell(r, c).value is not None]
        print(f"R{r}:", " | ".join(vals[:8]))
