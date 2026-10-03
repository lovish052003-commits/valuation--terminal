import openpyxl

for fname in ['master_model_template.xlsx', 'ITC Model.xlsx']:
    wb = openpyxl.load_workbook(fname)
    print(f"Updating {fname}...")
    if 'DCF' in wb.sheetnames:
        dcf = wb['DCF']
        dcf['H11'] = "='Intrinsic Valuation'!$L$69"
        dcf['I11'] = "='Intrinsic Valuation'!$L$69"
        dcf['J11'] = "=$I$11+($M$11-$I$11)/4*1"
        dcf['K11'] = "=$I$11+($M$11-$I$11)/4*2"
        dcf['L11'] = "=$I$11+($M$11-$I$11)/4*3"
        dcf['M11'] = "=D21"
        dcf['D18'] = "='Intrinsic Valuation'!$L$68"
        dcf['D21'] = "='Intrinsic Valuation'!$L$71"
        dcf['B45'] = "Margin of Safety / (Discount)"
        dcf['D45'] = "=(D42-D44)/D44"
        dcf['D45'].number_format = "+0.0%;-0.0%;0.0%"
        for col in ['H', 'I', 'J', 'K', 'L', 'M']:
            dcf[f'{col}12'] = f"={col}10*(1-{col}11)"
        print("  Updated DCF.")

    if 'Intrinsic Valuation' in wb.sheetnames:
        iv = wb['Intrinsic Valuation']
        iv['B67'] = 'Normalized ROIC (Sustainable)'
        iv['L67'] = '=IFERROR(MEDIAN(I40:L40), 0.12)'
        iv['L67'].number_format = '0.00%'

        iv['B68'] = 'Expected Growth Rate'
        iv['L68'] = 0.05
        iv['L68'].number_format = '0.00%'

        iv['B69'] = 'Fundamental Reinvestment Rate'
        iv['L69'] = '=IF(L67<=0, L55, L68/L67)'
        iv['L69'].number_format = '0.00%'

        iv['B70'] = 'Sustainable Terminal ROIC'
        iv['L70'] = '=MIN(0.18, MAX(0.06, 0.5*L67 + 0.5*DCF!D20))'
        iv['L70'].number_format = '0.00%'

        iv['B71'] = 'Terminal Reinvestment Rate'
        iv['L71'] = '=DCF!D19/L70'
        iv['L71'].number_format = '0.00%'

        iv['B72'] = 'Growth-ROIC Consistency Check'
        iv['L72'] = '=IF(ABS(L69*L67 - L68) <= 0.005, "PASS", "WARNING")'
        print("  Updated Intrinsic Valuation.")

    if 'WACC' in wb.sheetnames:
        wacc = wb['WACC']
        wacc['C34'] = "='Data Sheet'!K59"
        wacc['C35'] = "='Data Sheet'!K61"
        wacc['E27'] = 0.30
        print("  Updated WACC.")

    wb.save(fname)
    print(f"Successfully saved {fname}.")
