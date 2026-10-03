import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Advance Financial Project\exports\NESTLEIND_Valuation_Model.xlsx', data_only=True)
ws = wb['Data Sheet']

cols = [chr(65+i) for i in range(1, 12)] # B to L
for r in [56, 57, 67, 68, 69, 70, 72, 90, 93]:
    lbl = ws.cell(r, 1).value
    row_str = ' | '.join([f"{ws.cell(r, c).coordinate}={ws.cell(r, c).value}" for c in range(2, 12)])
    print(f"Row {r:2d} ({lbl}):")
    print("  " + row_str)
