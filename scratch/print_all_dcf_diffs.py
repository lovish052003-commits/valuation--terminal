import openpyxl

wb_f = openpyxl.load_workbook("FORCEMOT_Valuation_Model_FIXED.xlsx", data_only=False)
wb_o = openpyxl.load_workbook(r"exports\FORCEMOT_Valuation_Model.xlsx", data_only=False)

ws_f = wb_f['DCF']
ws_o = wb_o['DCF']

print("=== ALL DCF DIFFS ===")
for r in range(1, 60):
    for c in range(1, 15):
        vf = ws_f.cell(r, c).value
        vo = ws_o.cell(r, c).value
        if vf != vo:
            coord = openpyxl.utils.get_column_letter(c) + str(r)
            print(f"{coord}: ORIG={vo!r} -> FIXED={vf!r}")
