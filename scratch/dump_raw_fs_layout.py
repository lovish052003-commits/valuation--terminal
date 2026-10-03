import openpyxl

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=True)
ws = wb['Raw FS']

print(f"{'Row':<4} | {'Col B (Balance Sheet)':<30} | {'Col R (Income / Cash Flow)':<35}")
print('-' * 75)
for r in range(1, 55):
    b = str(ws.cell(row=r, column=2).value or '')
    r_val = str(ws.cell(row=r, column=18).value or '')
    if b or r_val:
        print(f"{r:<4d} | {b:<30} | {r_val:<35}")
