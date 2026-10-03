import openpyxl

wb = openpyxl.load_workbook('exports/ADANIENT_Valuation_Model.xlsx', data_only=True)
ws = wb['Dupont Analysis']
print("Dupont max_row:", ws.max_row, "max_column:", ws.max_column)
for r in range(1, min(ws.max_row + 1, 50)):
    for c in range(1, min(ws.max_column + 1, 15)):
        val = ws.cell(row=r, column=c).value
        if val is not None:
            col_letter = openpyxl.utils.get_column_letter(c)
            print(f"{col_letter}{r}: {str(val)[:40]}")
