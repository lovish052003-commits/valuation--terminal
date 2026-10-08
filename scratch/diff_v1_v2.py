import openpyxl

wb_v1 = openpyxl.load_workbook("UltraTech_fixed.xlsx", data_only=False)
wb_v2 = openpyxl.load_workbook("UltraTech_fixed_v2.xlsx", data_only=False)

def safe_str(val):
    return str(val).encode('ascii', errors='replace').decode('ascii')

for s in wb_v2.sheetnames:
    ws_1 = wb_v1[s]
    ws_2 = wb_v2[s]
    diffs = []
    max_r = max(ws_1.max_row, ws_2.max_row)
    max_c = max(ws_1.max_column, ws_2.max_column)
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            v1 = ws_1.cell(r, c).value
            v2 = ws_2.cell(r, c).value
            if v1 != v2:
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                diffs.append((coord, v1, v2))
    if diffs:
        print(f"Sheet {s} has {len(diffs)} diffs between v1 and v2:")
        for coord, v1, v2 in diffs:
            print(f"  {coord}: v1={safe_str(v1)!r} -> v2={safe_str(v2)!r}")
    else:
        print(f"Sheet {s}: IDENTICAL")
