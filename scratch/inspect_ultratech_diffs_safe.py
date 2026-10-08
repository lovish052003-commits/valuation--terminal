import openpyxl
import os

f1 = "UltraTech_fixed.xlsx"
f2 = "UltraTech_fixed_v2.xlsx"
orig = r"exports\ULTRACEMCO_Valuation_Model.xlsx"

wb_f1 = openpyxl.load_workbook(f1, data_only=False)
wb_f2 = openpyxl.load_workbook(f2, data_only=False)
wb_orig = openpyxl.load_workbook(orig, data_only=False) if os.path.exists(orig) else None

def safe_str(val):
    return str(val).encode('ascii', errors='replace').decode('ascii')

def compare_wbs(wb_a, wb_b, label_a, label_b):
    print(f"\n=================== COMPARING {label_a} vs {label_b} ===================")
    for s in wb_b.sheetnames:
        if s not in wb_a.sheetnames:
            print(f"Sheet {s} is ONLY in {label_b}!")
            continue
        ws_a = wb_a[s]
        ws_b = wb_b[s]
        diffs = []
        max_r = max(ws_a.max_row, ws_b.max_row)
        max_c = max(ws_a.max_column, ws_b.max_column)
        for r in range(1, max_r + 1):
            for c in range(1, max_c + 1):
                va = ws_a.cell(r, c).value
                vb = ws_b.cell(r, c).value
                if va != vb:
                    coord = openpyxl.utils.get_column_letter(c) + str(r)
                    diffs.append((coord, va, vb))
        if diffs:
            print(f"\nSheet {s}: {len(diffs)} differences:")
            for coord, va, vb in diffs[:40]:
                print(f"  {coord}: {label_a}={safe_str(va)!r} -> {label_b}={safe_str(vb)!r}")
            if len(diffs) > 40:
                print(f"  ... and {len(diffs) - 40} more diffs")
        else:
            print(f"Sheet {s}: IDENTICAL")

# Compare f1 vs f2 first
compare_wbs(wb_f1, wb_f2, "FIXED_V1", "FIXED_V2")

# Compare orig vs f2
if wb_orig:
    compare_wbs(wb_orig, wb_f2, "ORIG", "FIXED_V2")
