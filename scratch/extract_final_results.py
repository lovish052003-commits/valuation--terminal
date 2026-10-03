"""
extract_final_results.py
Reads and audits the final calculated workbook SUNPHARMA_Valuation_Model_Corrected.xlsx
"""

import openpyxl

TARGET_FILE = r'C:\Users\LENOVO\Downloads\test\SUNPHARMA_Valuation_Model_Corrected.xlsx'

print(f"Loading {TARGET_FILE}...")
wb_f = openpyxl.load_workbook(TARGET_FILE, data_only=False)
wb_d = openpyxl.load_workbook(TARGET_FILE, data_only=True)

print("\n1. FORMULA ERROR AUDIT:")
error_tokens = {'#REF!', '#DIV/0!', '#VALUE!', '#NAME?'}
found_errors = []

for sheet in wb_d.sheetnames:
    ws = wb_d[sheet]
    ws_f = wb_f[sheet]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r, column=c).value or '')
            if any(err in val for err in error_tokens):
                cell_ref = f"{openpyxl.utils.get_column_letter(c)}{r}"
                form_val = ws_f.cell(row=r, column=c).value
                found_errors.append((sheet, cell_ref, val, form_val))

print(f"   Total formula errors found across {len(wb_d.sheetnames)} sheets: {len(found_errors)}")
for err in found_errors:
    print(f"   -> {err[0]}!{err[1]}: Value='{err[2]}' | Formula='{err[3]}'")

# 2. KEY VALUATION METRICS EXTRACT
print("\n2. KEY VALUATION METRICS:")

# DCF
ws_dcf = wb_d['DCF']
dcf_price = ws_dcf['D42'].value
cmp_price = ws_dcf['D44'].value
wacc_val = ws_dcf['D20'].value
term_g = ws_dcf['D19'].value
ev_val = ws_dcf['D35'].value
cash_val = ws_dcf['D37'].value
debt_val = ws_dcf['D38'].value
eq_val = ws_dcf['D39'].value
shares_val = ws_dcf['D40'].value

print(f"   DCF Value per Share: Rs. {dcf_price:.2f}" if isinstance(dcf_price, (int, float)) else f"   DCF Value per Share: {dcf_price}")
print(f"   Current Market Price: Rs. {cmp_price:.2f}" if isinstance(cmp_price, (int, float)) else f"   CMP: {cmp_price}")
print(f"   WACC: {wacc_val:.2%}" if isinstance(wacc_val, float) else f"   WACC: {wacc_val}")
print(f"   Terminal Growth Rate: {term_g:.2%}" if isinstance(term_g, float) else f"   g: {term_g}")
print(f"   Operating Assets (EV): Rs. {ev_val:,.2f} Cr" if isinstance(ev_val, (int, float)) else f"   EV: {ev_val}")
print(f"   Cash: Rs. {cash_val:,.2f} Cr" if isinstance(cash_val, (int, float)) else f"   Cash: {cash_val}")
print(f"   Debt: Rs. {debt_val:,.2f} Cr" if isinstance(debt_val, (int, float)) else f"   Debt: {debt_val}")
print(f"   Net Equity Value: Rs. {eq_val:,.2f} Cr" if isinstance(eq_val, (int, float)) else f"   Net Equity Value: {eq_val}")
print(f"   Shares: {shares_val} Cr")

# Historical FS Metrics (2017 to 2025)
ws_hfs = wb_d['Historical FS']
print(f"\n3. HISTORICAL FS METRICS (Last 5 Years):")
for offset, cl in enumerate(['G', 'H', 'I', 'J', 'K']):
    yr = 2021 + offset
    sales = ws_hfs[f'{cl}7'].value
    cogs = ws_hfs[f'{cl}10'].value
    gp = ws_hfs[f'{cl}13'].value
    sga = ws_hfs[f'{cl}16'].value
    ebitda = ws_hfs[f'{cl}19'].value
    ebitda_m = ws_hfs[f'{cl}20'].value
    ebit = ws_hfs[f'{cl}25'].value
    ebit_m = ws_hfs[f'{cl}26'].value
    pbt = ws_hfs[f'{cl}31'].value
    tax = ws_hfs[f'{cl}34'].value
    eff_t = ws_hfs[f'{cl}35'].value
    np = ws_hfs[f'{cl}37'].value
    print(f"   FY{yr}: Sales={sales:,.0f} | GrossProfit={gp:,.0f} | SG&A={sga:,.0f} | EBITDA={ebitda:,.0f} ({ebitda_m:.1%}) | EBIT={ebit:,.0f} ({ebit_m:.1%}) | PBT={pbt:,.0f} | Tax={tax:,.0f} ({eff_t:.1%}) | NetProfit={np:,.0f}")

# Comp Valuation
ws_comp = wb_d['Comp_Valuation']
print(f"\n4. COMPARABLE COMPANY VALUATION:")
print(f"   - Median EV/Revenue: {ws_comp['O25'].value:.2f}x" if isinstance(ws_comp['O25'].value, float) else f"   - Median EV/Rev: {ws_comp['O25'].value}")
print(f"   - Median EV/EBITDA:  {ws_comp['P25'].value:.2f}x" if isinstance(ws_comp['P25'].value, float) else f"   - Median EV/EBITDA: {ws_comp['P25'].value}")
print(f"   - Median P/E:        {ws_comp['Q25'].value:.2f}x" if isinstance(ws_comp['Q25'].value, float) else f"   - Median P/E: {ws_comp['Q25'].value}")
print(f"   - Sun Pharma Net Debt used: Rs. {ws_comp['O33'].value:,.2f} Cr" if isinstance(ws_comp['O33'].value, (int, float)) else f"   - Net Debt: {ws_comp['O33'].value}")
print(f"   - Implied Value/Share (EV/Rev):    Rs. {ws_comp['O37'].value:.2f}" if isinstance(ws_comp['O37'].value, float) else f"   - EV/Rev Value: {ws_comp['O37'].value}")
print(f"   - Implied Value/Share (EV/EBITDA): Rs. {ws_comp['P37'].value:.2f}" if isinstance(ws_comp['P37'].value, float) else f"   - EV/EBITDA Value: {ws_comp['P37'].value}")
print(f"   - Implied Value/Share (P/E):       Rs. {ws_comp['Q37'].value:.2f}" if isinstance(ws_comp['Q37'].value, float) else f"   - P/E Value: {ws_comp['Q37'].value}")

# Sensitivity Table Check
print(f"\n5. DCF SENSITIVITY TABLE (Implied Value per Share in Rs.):")
header_g = [ws_dcf.cell(row=49, column=c).value for c in range(3, 9)]
if all(isinstance(g, (int, float)) for g in header_g):
    print(f"   WACC \\ g     " + "   ".join([f"{g:.1%}" for g in header_g]))
for r in range(50, 57):
    w_val = ws_dcf.cell(row=r, column=2).value
    row_prices = [ws_dcf.cell(row=r, column=c).value for c in range(3, 9)]
    row_str = "   ".join([f"Rs.{p:,.1f}" if isinstance(p, (int, float)) else str(p) for p in row_prices])
    w_str = f"{w_val:.1%}" if isinstance(w_val, (int, float)) else str(w_val)
    print(f"   {w_str:<10}  {row_str}")

# Beta Regression
ws_beta = wb_d['Beta-Regression']
raw_beta = ws_beta['L9'].value
adj_beta = ws_beta['L15'].value
print(f"\n6. BETA REGRESSION:")
print(f"   - Raw Levered Beta: {raw_beta:.4f}" if isinstance(raw_beta, float) else f"   - Raw Beta: {raw_beta}")
print(f"   - Adjusted Beta:     {adj_beta:.4f}" if isinstance(adj_beta, float) else f"   - Adj Beta: {adj_beta}")

# WACC Capital Structure
ws_wacc = wb_d['WACC']
debt_w = ws_wacc['D34'].value
eq_w = ws_wacc['D35'].value
total_w = ws_wacc['D36'].value
print(f"\n7. WACC CAPITAL STRUCTURE:")
print(f"   - Debt Weight:   {debt_w:.2%}" if isinstance(debt_w, float) else f"   - Debt Weight: {debt_w}")
print(f"   - Equity Weight: {eq_w:.2%}" if isinstance(eq_w, float) else f"   - Equity Weight: {eq_w}")
print(f"   - Total Weight:  {total_w:.2%}" if isinstance(total_w, float) else f"   - Total Weight: {total_w}")

print("\nAudit and extraction completed successfully!")
