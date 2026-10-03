import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\test\WIPRO_Valuation_Model.xlsx', data_only=True)
print("Sheetnames:", wb.sheetnames)
ws = wb['Data Sheet']
print("Data Sheet B1:", ws['B1'].value)
print("Data Sheet B6:", ws['B6'].value)
print("Data Sheet B8:", ws['B8'].value)
print("Data Sheet B9:", ws['B9'].value)
print("Data Sheet K17:", ws['K17'].value)
print("Data Sheet K69:", ws['K69'].value)
print("Data Sheet K70:", ws['K70'].value)
