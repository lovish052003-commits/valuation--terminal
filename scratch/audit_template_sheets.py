import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
for sname in wb.sheetnames:
    ws = wb[sname]
    sample = []
    has_formula = False
    has_itc = False
    for r in range(1, 10):
        for c in range(1, 10):
            v = ws.cell(r, c).value
            if v is not None:
                vs = str(v)
                if vs.startswith('='):
                    has_formula = True
                if 'itc' in vs.lower():
                    has_itc = True
                sample.append(f"{openpyxl.utils.get_column_letter(c)}{r}:{vs[:25]}")
    print(f"{sname:25} | Formulas: {str(has_formula):5} | Has ITC: {str(has_itc):5} | Sample: {' '.join(sample[:3])}")
