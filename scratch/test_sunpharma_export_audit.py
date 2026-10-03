import sys, os
sys.path.insert(0, os.path.abspath('.'))
import screener_client
import valuation_engine
import excel_exporter
import openpyxl

screener_data = screener_client.fetch_company_data('SUNPHARMA')
val_result = valuation_engine.calculate_valuation(screener_data)

out_path = excel_exporter.export_valuation_model(screener_data, val_result, report_markdown="")
print("Exported to:", out_path)

wb = openpyxl.load_workbook(out_path, data_only=True)
ws = wb['Data Sheet']
print("=== Data Sheet Rows 16 to 33 ===")
for r in range(16, 34):
    label = ws.cell(r, 1).value
    vals = [ws.cell(r, c).value for c in range(2, 12)]
    print(f"Row {r:2d} | {str(label):25s} | {vals}")
