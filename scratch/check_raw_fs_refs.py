import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)

for sname in ['Intrinsic Valuation', 'Comp_Valuation']:
    ws = wb[sname]
    refs = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            cell = ws.cell(r, c)
            if cell.value and isinstance(cell.value, str) and 'Raw FS' in cell.value:
                refs.append((cell.coordinate, cell.value))
    print(f"=== {sname} has {len(refs)} references to Raw FS ===")
    for coord, val in refs[:15]:
        print(f"  {coord}: {val}")
