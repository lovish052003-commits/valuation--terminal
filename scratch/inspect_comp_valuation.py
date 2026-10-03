import openpyxl

for fname in ['ITC Model.xlsx', 'exports/TATASTEEL_Valuation_Model.xlsx']:
    wb = openpyxl.load_workbook(fname, data_only=True)
    if 'Comp_Valuation' in wb.sheetnames:
        ws = wb['Comp_Valuation']
        print(f"\n=== {fname} -> Comp_Valuation ===")
        for r in range(10, 26):
            row_vals = [f"Col {openpyxl.utils.get_column_letter(c)}: {ws.cell(r, c).value}" for c in range(2, 18) if ws.cell(r, c).value is not None]
            if row_vals:
                print(f"  Row {r}:", " | ".join(row_vals[:6]))
