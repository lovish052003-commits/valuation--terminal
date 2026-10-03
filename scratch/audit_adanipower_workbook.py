import openpyxl

path = r'C:\Users\LENOVO\Downloads\test\ADANIPOWER_Valuation_Model (2).xlsx'
wb_v = openpyxl.load_workbook(path, data_only=True)
wb_f = openpyxl.load_workbook(path, data_only=False)

print('=== DATA SHEET ROW 61 & 66 ===')
ws_d = wb_f['Data Sheet']
ws_dv = wb_v['Data Sheet']
print('Row 61 Forms:', [ws_d.cell(61, c).value for c in range(2, 12)])
print('Row 61 Vals :', [ws_dv.cell(61, c).value for c in range(2, 12)])
print('Row 66 Forms:', [ws_d.cell(66, c).value for c in range(2, 12)])
print('Row 66 Vals :', [ws_dv.cell(66, c).value for c in range(2, 12)])

print("\n=== ALTMAN Z SCORE ===")
ws_a = wb_f["Altman's Z Score"]
ws_av = wb_v["Altman's Z Score"]
print('Row 60 Vals (Total Assets):', [ws_av.cell(60, c).value for c in [5, 9]])
print('Row 89 Vals (Final Score) :', [ws_av.cell(89, c).value for c in [5, 9]])
print('Row 90 Vals (Stability)   :', [ws_av.cell(90, c).value for c in [5, 9]])

print('\n=== DUPONT ANALYSIS ===')
ws_dp = wb_f['Dupont Analysis']
ws_dpv = wb_v['Dupont Analysis']
print('Row 71 Vals (Avg Assets):', [ws_dpv.cell(71, c).value for c in [5, 9]])
print('Row 72 Vals (Asset Turn):', [ws_dpv.cell(72, c).value for c in [5, 9]])
print('Row 78 Vals (ROE)       :', [ws_dpv.cell(78, c).value for c in [5, 9]])

print('\n=== AI VALUATION SUMMARY ===')
ws_s = wb_f['AI Valuation Summary']
ws_sv = wb_v['AI Valuation Summary']
print('F5 (Altman): form =', ws_s['F5'].value, ', val =', ws_sv['F5'].value)
print('G5 (DuPont): form =', ws_s['G5'].value, ', val =', ws_sv['G5'].value)

print('\n=== COMMON SIZE STATEMENT ROW 27 & 33 ===')
ws_cs = wb_f['Common Size Statement']
ws_csv = wb_v['Common Size Statement']
print('Row 27 Forms:', [ws_cs.cell(27, c).value for c in range(2, 6)])
print('Row 27 Vals :', [ws_csv.cell(27, c).value for c in range(2, 6)])
print('Row 33 Forms:', [ws_cs.cell(33, c).value for c in range(2, 6)])
print('Row 33 Vals :', [ws_csv.cell(33, c).value for c in range(2, 6)])

print('\n=== INTRINSIC VALUATION ROWS 9-19 ===')
ws_iv = wb_f['Intrinsic Valuation']
ws_ivv = wb_v['Intrinsic Valuation']
for r in [9, 10, 11, 12, 13, 16, 17, 18, 19, 21]:
    lbl = ws_iv.cell(r, 2).value
    form = ws_iv.cell(r, 12).value
    val = ws_ivv.cell(r, 12).value
    print(f'Row {r:2d} | {str(lbl):25s} | Form: {str(form):20s} | Val: {val}')
