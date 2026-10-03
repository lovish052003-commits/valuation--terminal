import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['Beta-Regression']

print("B7:", ws['B7'].value)
print("B10:", ws['B10'].value)
print("C10:", ws['C10'].value)
print("D10:", ws['D10'].value)
print("F10:", ws['F10'].value)
print("G10:", ws['G10'].value)
print("H10:", ws['H10'].value)
print("C11:", ws['C11'].value)
print("D11:", ws['D11'].value)
print("C252:", ws['C252'].value)
print("O11:", ws['O11'].value)
print("L15:", ws['L15'].value)
