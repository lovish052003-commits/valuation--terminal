import os
import openpyxl

def main():
    p = 'exports/ITC_Valuation_Model.xlsx'
    wb = openpyxl.load_workbook(p, data_only=False)
    ws = wb['Intrinsic Valuation']
    wb_v = openpyxl.load_workbook(p, data_only=True)
    ws_v = wb_v['Intrinsic Valuation']
    
    # We can also check with COM for evaluated values
    import pythoncom, win32com.client
    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx('Excel.Application')
    excel.Visible = False
    wb_com = excel.Workbooks.Open(os.path.abspath(p), UpdateLinks=0, ReadOnly=True)
    ws_com = wb_com.Sheets('Intrinsic Valuation')
    
    for r in range(40, 66):
        lbl = ws.cell(row=r, column=2).value or ''
        f_k = ws[f'K{r}'].value
        f_l = ws[f'L{r}'].value
        v_l = ws_com.Range(f'L{r}').Value
        print(f"R{r} ({lbl}): K={f_k} | L_formula={f_l} | L_val={v_l}")
        
    wb_com.Close(False)
    excel.Quit()

if __name__ == '__main__':
    main()
