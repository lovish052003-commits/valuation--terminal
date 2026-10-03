import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=False)
print("Sheet names:", wb.sheetnames)

for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"\n--- Sheet: {sname} (max_row={ws.max_row}, max_column={ws.max_column}) ---")
    for r in range(1, min(25, ws.max_row + 1)):
        row_vals = [f"{ws.cell(r, c).coordinate}={ws.cell(r, c).value}" for c in range(1, min(15, ws.max_column + 1)) if ws.cell(r, c).value is not None]
        if row_vals:
            print(f"Row {r:2d}: " + " | ".join(row_vals[:6]))
