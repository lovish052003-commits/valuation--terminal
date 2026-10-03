import openpyxl

wb_r = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=True)
wb_g = openpyxl.load_workbook('exports/NESTLEIND_Valuation_Model.xlsx', data_only=True)

ws_r = wb_r['Data Sheet']
ws_g = wb_g['Data Sheet']

cols = ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K']

print("=== Comparing Data Sheet between Nestle India Model.xlsx and Generated NESTLEIND ===")
for r in range(5, 95):
    lbl = ws_r.cell(r, 1).value
    if not lbl:
        continue
    row_diffs = []
    for c_idx, c_letter in enumerate(cols, start=2):
        vr = ws_r.cell(r, c_idx).value
        vg = ws_g.cell(r, c_idx).value
        # Compare strings or rounded floats
        if vr is None and vg is None:
            continue
        try:
            if abs(float(vr) - float(vg)) > 1.0:
                row_diffs.append((c_letter, vr, vg))
        except (ValueError, TypeError):
            if str(vr)[:10] != str(vg)[:10]:
                row_diffs.append((c_letter, vr, vg))
    if row_diffs:
        print(f"\nRow {r:02d} ({lbl}): {len(row_diffs)} column diffs")
        for c_let, vr, vg in row_diffs[:5]:
            print(f"   Col {c_let}: REF={vr} | GEN={vg}")
