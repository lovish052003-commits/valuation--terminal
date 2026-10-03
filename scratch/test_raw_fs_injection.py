import os, sys, json, re, shutil
import openpyxl
import pandas as pd

sys.path.insert(0, os.path.abspath('.'))
from screener_client import clean_num

def parse_table_payload(tbl_data):
    if tbl_data is None:
        return None
    if isinstance(tbl_data, pd.DataFrame):
        return tbl_data
    if isinstance(tbl_data, dict):
        if 'columns' in tbl_data and 'data' in tbl_data:
            return pd.DataFrame(data=tbl_data['data'], columns=tbl_data['columns'], index=tbl_data.get('index'))
        elif 'columns' in tbl_data and 'rows' in tbl_data:
            return pd.DataFrame(tbl_data['rows'])
        else:
            return pd.DataFrame.from_dict(tbl_data)
    if isinstance(tbl_data, list):
        return pd.DataFrame(tbl_data)
    return None

def inject_target_financials_to_raw_fs(wb_or_path, screener_data_or_json, valuation_result=None):
    """
    Parses the incoming JSON payload / screener_data for the target company's historical
    Income Statement, Balance Sheet, and Cash Flow data.
    Opens the 'Raw FS' tab using openpyxl.
    Maps and writes the target company's actual historical Revenue, EBIT, Net Income, and
    all core line items row-by-row into their respective cells on the Raw FS tab,
    completely overwriting the existing template data and eliminating template bleed.
    """
    is_path = isinstance(wb_or_path, str)
    if is_path:
        wb = openpyxl.load_workbook(wb_or_path)
    else:
        wb = wb_or_path

    if 'Raw FS' not in wb.sheetnames:
        if is_path:
            wb.close()
        return wb

    ws_raw = wb['Raw FS']

    # Parse screener_data if it's a json file or string
    if isinstance(screener_data_or_json, str):
        if os.path.exists(screener_data_or_json):
            with open(screener_data_or_json, 'r', encoding='utf-8') as f:
                screener_data = json.load(f)
        else:
            screener_data = json.loads(screener_data_or_json)
    else:
        screener_data = screener_data_or_json

    tables = screener_data.get('tables', {})
    schedules = screener_data.get('schedules', {})

    bs_df = parse_table_payload(tables.get('balance-sheet'))
    pl_df = parse_table_payload(tables.get('profit-loss'))
    cf_df = parse_table_payload(tables.get('cash-flow'))

    def get_sched_val(parent_cat, item_name, period):
        sch = schedules.get(parent_cat, {})
        for k, v in sch.items():
            if item_name.lower() in k.lower():
                # Fuzzy match period
                p_str = str(period).strip()
                p_clean = re.sub(r'\s*\d+m\b', '', p_str, flags=re.I).strip()
                if p_clean in v:
                    return clean_num(v[p_clean])
                p_norm = p_clean.lower().replace('-', ' ').replace('_', ' ')
                for pk, pv in v.items():
                    pk_norm = str(pk).strip().lower().replace('-', ' ').replace('_', ' ')
                    if pk_norm == p_norm:
                        return clean_num(pv)
        return 0.0

    # -------------------------------------------------------------------------
    # 1. BALANCE SHEET (Cols C to N = cols 3 to 14, Rows 3 to 44)
    # -------------------------------------------------------------------------
    # First, completely wipe existing template values in Balance Sheet area
    for r in range(3, 45):
        for c in range(3, 15):
            ws_raw.cell(row=r, column=c, value=None if r == 3 else 0.0)

    if bs_df is not None and not bs_df.empty:
        bs_period_cols = [c for c in bs_df.columns if c != 'Metric']
        # Take up to 12 periods, right-aligned to Column 14 (Col N)
        sel_bs_periods = bs_period_cols[-12:]
        bs_start_col = 14 - len(sel_bs_periods) + 1

        for offset, p_name in enumerate(sel_bs_periods):
            c_idx = bs_start_col + offset
            ws_raw.cell(row=3, column=c_idx, value=str(p_name))

        def write_bs_metric(row_idx, keywords):
            for kw in keywords:
                m = bs_df[bs_df['Metric'].str.contains(kw, case=False, na=False)]
                if not m.empty:
                    r = m.iloc[0]
                    for offset, p_name in enumerate(sel_bs_periods):
                        val = clean_num(r.get(p_name, 0.0))
                        ws_raw.cell(row=row_idx, column=bs_start_col + offset, value=val)
                    return True
            return False

        def write_bs_sched(row_idx, parent_cat, item_kw):
            has_val = False
            for offset, p_name in enumerate(sel_bs_periods):
                v = get_sched_val(parent_cat, item_kw, p_name)
                ws_raw.cell(row=row_idx, column=bs_start_col + offset, value=v)
                if v != 0:
                    has_val = True
            return has_val

        # Core Balance Sheet Line Items
        write_bs_metric(5, ['Equity Capital', 'Share Capital'])
        write_bs_metric(6, ['Reserves'])
        write_bs_metric(7, ['Borrowings'])
        write_bs_sched(8, 'Borrowings', 'Long term Borrowings')
        write_bs_sched(9, 'Borrowings', 'Short term Borrowings')
        write_bs_sched(10, 'Borrowings', 'Lease Liabilities')
        write_bs_sched(11, 'Borrowings', 'Other Borrowings')

        write_bs_metric(12, ['Other Liabilities'])
        write_bs_sched(13, 'Other Liabilities', 'Non controlling int')
        write_bs_sched(14, 'Other Liabilities', 'Trade Payables')
        write_bs_sched(15, 'Other Liabilities', 'Advance from Customers')
        write_bs_sched(16, 'Other Liabilities', 'Other liability')
        write_bs_metric(18, ['Total Liabilities'])

        # Fixed Assets Breakdown
        write_bs_metric(21, ['Fixed Assets', 'Property, Plant'])
        write_bs_sched(22, 'Fixed Assets', 'Land')
        write_bs_sched(23, 'Fixed Assets', 'Building')
        write_bs_sched(24, 'Fixed Assets', 'Plant Machinery')
        write_bs_sched(25, 'Fixed Assets', 'Equipments')
        write_bs_sched(26, 'Fixed Assets', 'Furniture')
        write_bs_sched(27, 'Fixed Assets', 'Railway sidings')
        write_bs_sched(28, 'Fixed Assets', 'Vehicles')
        write_bs_sched(29, 'Fixed Assets', 'Intangible')
        write_bs_sched(30, 'Fixed Assets', 'Other fixed assets')
        write_bs_sched(31, 'Fixed Assets', 'Gross Block')
        write_bs_sched(32, 'Fixed Assets', 'Accumulated Depreciation')
        # Net Block (Row 33)
        if not write_bs_sched(33, 'Fixed Assets', 'Net Block'):
            write_bs_metric(33, ['Fixed Assets', 'Net Block'])

        write_bs_metric(35, ['CWIP'])
        write_bs_metric(36, ['Investments'])
        write_bs_metric(38, ['Other Assets'])
        write_bs_sched(39, 'Other Assets', 'Inventories')
        write_bs_sched(40, 'Other Assets', 'Trade receivables')
        write_bs_sched(41, 'Other Assets', 'Cash Equivalents')
        write_bs_sched(42, 'Other Assets', 'Loans n Advances')
        write_bs_sched(43, 'Other Assets', 'Other asset items')
        write_bs_metric(44, ['Total Assets'])

    # -------------------------------------------------------------------------
    # 2. INCOME STATEMENT (P&L) (Cols S to AE = cols 19 to 31, Rows 3 to 15)
    # -------------------------------------------------------------------------
    # Wipe existing template data in P&L area
    for r in range(3, 16):
        for c in range(19, 32):
            ws_raw.cell(row=r, column=c, value=None if r == 3 else 0.0)

    if pl_df is not None and not pl_df.empty:
        pl_period_cols = [c for c in pl_df.columns if c != 'Metric']
        has_ttm = any('ttm' in str(c).lower() for c in pl_period_cols)
        annual_cols = [c for c in pl_period_cols if 'ttm' not in str(c).lower()]
        sel_annual = annual_cols[-12:]
        # Column 30 = Col AD (latest annual). Start col = 30 - len(sel_annual) + 1
        pl_start_col = 30 - len(sel_annual) + 1

        for offset, p_name in enumerate(sel_annual):
            ws_raw.cell(row=3, column=pl_start_col + offset, value=str(p_name))
        if has_ttm:
            ws_raw.cell(row=3, column=31, value='TTM')

        all_mapped_periods = list(sel_annual) + (['TTM'] if has_ttm else [])

        def write_pl_metric(row_idx, keywords):
            for kw in keywords:
                m = pl_df[pl_df['Metric'].str.contains(kw, case=False, na=False)]
                if not m.empty:
                    r = m.iloc[0]
                    for offset, p_name in enumerate(sel_annual):
                        val = clean_num(r.get(p_name, 0.0))
                        ws_raw.cell(row=row_idx, column=pl_start_col + offset, value=val)
                    if has_ttm:
                        # Find ttm column in r
                        ttm_col = [c for c in pl_period_cols if 'ttm' in str(c).lower()]
                        if ttm_col:
                            ws_raw.cell(row=row_idx, column=31, value=clean_num(r.get(ttm_col[0], 0.0)))
                    return True
            return False

        write_pl_metric(4, ['^Sales', 'Revenue', 'Total Revenue'])
        write_pl_metric(5, ['^Expenses', 'Total Expenses'])
        write_pl_metric(6, ['Operating Profit', 'EBITDA'])
        write_pl_metric(7, ['OPM'])
        write_pl_metric(8, ['Other Income'])
        write_pl_metric(9, ['Interest', 'Finance Cost'])
        write_pl_metric(10, ['Depreciation'])
        write_pl_metric(11, ['Profit before tax', 'PBT'])
        write_pl_metric(12, ['Tax %', 'Tax'])
        write_pl_metric(13, ['Net Profit', 'PAT'])
        write_pl_metric(14, ['EPS'])
        write_pl_metric(15, ['Dividend Payout'])

    # -------------------------------------------------------------------------
    # 3. CASH FLOW STATEMENT (Cols S to AE = cols 19 to 31, Rows 19 to 45)
    # -------------------------------------------------------------------------
    # Wipe existing template data in Cash Flow area
    for r in range(19, 46):
        for c in range(19, 32):
            ws_raw.cell(row=r, column=c, value=0.0)

    if cf_df is not None and not cf_df.empty:
        cf_period_cols = [c for c in cf_df.columns if c != 'Metric']
        has_ttm_cf = any('ttm' in str(c).lower() for c in cf_period_cols)
        annual_cf_cols = [c for c in cf_period_cols if 'ttm' not in str(c).lower()]
        sel_cf_annual = annual_cf_cols[-12:]
        cf_start_col = 30 - len(sel_cf_annual) + 1

        def write_cf_metric(row_idx, keywords):
            for kw in keywords:
                m = cf_df[cf_df['Metric'].str.contains(kw, case=False, na=False)]
                if not m.empty:
                    r = m.iloc[0]
                    for offset, p_name in enumerate(sel_cf_annual):
                        val = clean_num(r.get(p_name, 0.0))
                        ws_raw.cell(row=row_idx, column=cf_start_col + offset, value=val)
                    if has_ttm_cf:
                        ttm_c = [c for c in cf_period_cols if 'ttm' in str(c).lower()]
                        if ttm_c:
                            ws_raw.cell(row=row_idx, column=31, value=clean_num(r.get(ttm_c[0], 0.0)))
                    return True
            return False

        def write_cf_sched(row_idx, parent_cat, item_kw):
            has_val = False
            for offset, p_name in enumerate(sel_cf_annual):
                v = get_sched_val(parent_cat, item_kw, p_name)
                ws_raw.cell(row=row_idx, column=cf_start_col + offset, value=v)
                if v != 0:
                    has_val = True
            return has_val

        write_cf_metric(19, ['Cash from Operating Activity'])
        write_cf_sched(20, 'Cash from Operating Activity', 'Profit from operations')
        write_cf_sched(21, 'Cash from Operating Activity', 'Receivables')
        write_cf_sched(22, 'Cash from Operating Activity', 'Inventory')
        write_cf_sched(23, 'Cash from Operating Activity', 'Payables')
        write_cf_sched(24, 'Cash from Operating Activity', 'Working capital')
        write_cf_sched(25, 'Cash from Operating Activity', 'Direct taxes')

        write_cf_metric(26, ['Cash from Investing Activity'])
        write_cf_sched(27, 'Cash from Investing Activity', 'Fixed assets purchased')
        write_cf_sched(28, 'Cash from Investing Activity', 'Fixed assets sold')
        write_cf_sched(29, 'Cash from Investing Activity', 'Investments purchased')
        write_cf_sched(30, 'Cash from Investing Activity', 'Investments sold')
        write_cf_sched(31, 'Cash from Investing Activity', 'Interest received')
        write_cf_sched(32, 'Cash from Investing Activity', 'Dividends received')
        write_cf_sched(33, 'Cash from Investing Activity', 'subsidiaries')
        write_cf_sched(34, 'Cash from Investing Activity', 'group cos')
        write_cf_sched(35, 'Cash from Investing Activity', 'Redemp')
        write_cf_sched(36, 'Cash from Investing Activity', 'Other investing')

        write_cf_metric(37, ['Cash from Financing Activity'])
        write_cf_sched(38, 'Cash from Financing Activity', 'shares')
        write_cf_sched(39, 'Cash from Financing Activity', 'Proceeds from borrowings')
        write_cf_sched(40, 'Cash from Financing Activity', 'Repayment of borrowings')
        write_cf_sched(41, 'Cash from Financing Activity', 'Interest paid')
        write_cf_sched(42, 'Cash from Financing Activity', 'Dividends paid')
        write_cf_sched(43, 'Cash from Financing Activity', 'Financial liabilities')
        write_cf_sched(44, 'Cash from Financing Activity', 'Other financing')
        write_cf_metric(45, ['Net Cash Flow'])

    if is_path:
        wb.save(wb_or_path)
        wb.close()
    return wb

# Test with UNITEDTEA
test_out = 'scratch/test_unitedtea_injected.xlsx'
shutil.copyfile('master_model_template.xlsx', test_out)

with open('exports/.screener_cache/UNITEDTEA_data.json', 'r') as f:
    sd = json.load(f)

print("Injecting UNITEDTEA financials into Raw FS...")
inject_target_financials_to_raw_fs(test_out, sd)
print("Injection complete!")

# Verify Raw FS contents
wb_chk = openpyxl.load_workbook(test_out, data_only=True)
ws_chk = wb_chk['Raw FS']
print("\n--- Verifying Raw FS Injected Cells ---")
print("Col AD row 4 (Sales):", ws_chk['AD4'].value)
print("Col AD row 6 (Operating Profit):", ws_chk['AD6'].value)
print("Col AD row 9 (Interest):", ws_chk['AD9'].value)
print("Col AD row 11 (Profit before tax):", ws_chk['AD11'].value)
print("Col AD row 13 (Net Profit):", ws_chk['AD13'].value)
print("Col AD row 27 (Fixed assets purchased / Capex):", ws_chk['AD27'].value)
print("Col N row 5 (Equity Capital):", ws_chk['N5'].value)
print("Col N row 7 (Borrowings):", ws_chk['N7'].value)
print("Col N row 33 (Net Block):", ws_chk['N33'].value)
print("Col N row 44 (Total Assets):", ws_chk['N44'].value)
