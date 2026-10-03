import openpyxl

wb = openpyxl.load_workbook('exports/TATASTEEL_Valuation_Model.xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

def print_sheet_sample(sheet_name, cells):
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} NOT FOUND")
        return
    ws = wb[sheet_name]
    print(f"\n--- {sheet_name} ---")
    for c in cells:
        val = ws[c].value
        print(f"  {c}: {val}")

print_sheet_sample('Data Sheet', ['B1', 'B6', 'B7', 'B8', 'B9', 'B16', 'B17', 'B30', 'B56', 'B57', 'B61', 'B81', 'B82', 'K17', 'K30', 'K57', 'K61', 'K69', 'K70'])
print_sheet_sample('Profit & Loss', ['A1', 'B3', 'C3', 'B4', 'C4', 'B5', 'C5', 'B6', 'C6'])
print_sheet_sample('Balance Sheet', ['A1', 'B3', 'C3', 'B4', 'C4', 'B5', 'C5', 'B8', 'C8'])
print_sheet_sample('Cash Flow', ['A1', 'B3', 'C3', 'B4', 'C4', 'B7', 'C7'])
print_sheet_sample('Raw FS', ['L56', 'M56', 'AR56', 'L57', 'M57', 'AR57', 'L58', 'M58', 'AR58'])
print_sheet_sample('Comp_Valuation', ['C12', 'D12', 'E12', 'C13', 'D13', 'E13'])
print_sheet_sample('DCF', ['D18', 'D20', 'D33', 'D35', 'D37', 'D38', 'D39', 'D40', 'D42', 'D44'])
print_sheet_sample('AI Valuation Summary', ['A1', 'A5', 'B5', 'C5', 'D5', 'E5', 'F5', 'G5', 'B22', 'B25', 'B28', 'B30', 'B31'])
print_sheet_sample('Dupont Analysis', ['B5', 'B8', 'B37'])
print_sheet_sample("Altman's Z Score", ['B5', 'B8', 'B36'])
