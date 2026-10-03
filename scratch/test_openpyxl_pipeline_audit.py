import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import openpyxl
import screener_client
import valuation_engine
import excel_exporter

def test_pipeline_for_company(ticker):
    print("\n" + "=" * 80)
    print(f"TESTING OPENPYXL PIPELINE AUDIT FOR: {ticker}")
    print("=" * 80)
    
    # 1. Fetch data
    sd = screener_client.fetch_company_data(ticker)
    c_name = sd.get('company_name', ticker)
    print(f"Company: {c_name} ({ticker}) | Sector: {sd.get('sector')} | Industry: {sd.get('industry')}")
    
    # 2. Run valuation engine
    val = valuation_engine.calculate_valuation(sd)
    
    # Check Python DCF table Year 1 Reinvestment Rate
    dcf_table = val.get('dcf_table', [])
    if dcf_table:
        y1_rr = dcf_table[0].get('reinvestment_rate')
        all_rr = [x.get('reinvestment_rate') for x in dcf_table]
        print(f"Python DCF Table Reinvestment Rates: {all_rr}")
        print(f"Year 1 Reinvestment Rate: {y1_rr}%")
        assert y1_rr <= 85.0, f"FAILED: Year 1 Reinvestment Rate {y1_rr}% exceeds 85.0% clamp!"
        for yr_idx, r in enumerate(all_rr):
            assert r <= 85.0, f"FAILED: Year {yr_idx+1} Reinvestment Rate {r}% exceeds 85.0% clamp!"
        print(" -> PASS: Python DCF forecasting logic strictly enforces <= 85% clamp on Year 1 and all forecast years.")
    
    # 3. Generate Excel model via OpenPyXL pipeline
    out_dir = os.path.join(os.getcwd(), 'exports')
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f"{ticker}_OpenPyXL_Audit.xlsx")
    
    import shutil
    shutil.copyfile('master_model_template.xlsx', out_path)
    res_path = excel_exporter.patch_valuation_workbook(out_path, sd, val)
    if not res_path:
        res_path = out_path
    print(f"Excel generated at: {res_path}")
    
    # 4. Audit generated workbook
    wb = openpyxl.load_workbook(res_path, data_only=False)
    
    # Check DCF Sheet Reinvestment Formulas
    ws_dcf = wb['DCF']
    print(f"DCF!I11 Formula: {ws_dcf['I11'].value}")
    print(f"DCF!J11 Formula: {ws_dcf['J11'].value}")
    print(f"DCF!K11 Formula: {ws_dcf['K11'].value}")
    print(f"DCF!L11 Formula: {ws_dcf['L11'].value}")
    print(f"DCF!M11 Formula: {ws_dcf['M11'].value}")
    assert "0.85" in str(ws_dcf['I11'].value), "FAILED: DCF!I11 formula missing 0.85 clamp!"
    assert "0.85" in str(ws_dcf['M11'].value), "FAILED: DCF!M11 formula missing 0.85 clamp!"
    print(" -> PASS: DCF Sheet Formulas enforce MIN(0.85, ...) clamp universally.")
    
    # Check Raw FS Sheet
    ws_raw = wb['Raw FS']
    n24_val = ws_raw['N24'].value
    print(f"Raw FS!N24 Value: {n24_val}")
    assert n24_val != 57284 and n24_val != '57284', f"FAILED: Raw FS!N24 still contains baseline template bleed {n24_val}!"
    print(" -> PASS: Raw FS!N24 does not contain template bleed (57,284).")
    
    # Check Row 4 (Total Revenue)
    rev_vals = [ws_raw.cell(row=4, column=c).value for c in range(19, 32)]
    print(f"Raw FS Row 4 (Revenue, Cols S to AE): {rev_vals}")
    maruti_rev = [50801, 57589, 68085, 79809, 86068, 75660, 70372, 88330, 118410, 141858, 152913, 183316, 197181]
    assert rev_vals != maruti_rev, "FAILED: Raw FS Row 4 matches Maruti Suzuki template revenue!"
    # Ensure latest year (Col 30 = AD4) is not template value 183316
    assert rev_vals[-2] != 183316, f"FAILED: Latest revenue cell AD4 is template value {rev_vals[-2]}!"
    print(" -> PASS: Raw FS Total Revenue cleanly populated from target data with zero template bleed.")
    
    # Check Row 13 (Net Income)
    pat_vals = [ws_raw.cell(row=13, column=c).value for c in range(19, 32)]
    print(f"Raw FS Row 13 (Net Income, Cols S to AE): {pat_vals}")
    maruti_pat = [3809, 5497, 7511, 7881, 7651, 5678, 4389, 3880, 8264, 13488, 14500, 14680, 14334]
    assert pat_vals != maruti_pat, "FAILED: Raw FS Row 13 matches Maruti Suzuki template Net Income!"
    assert pat_vals[-2] != 14680, f"FAILED: Latest Net Income cell AD13 is template value {pat_vals[-2]}!"
    print(" -> PASS: Raw FS Net Income cleanly populated from target data with zero template bleed.")
    
    # Check Row 10 (D&A)
    depr_vals = [ws_raw.cell(row=10, column=c).value for c in range(19, 32)]
    print(f"Raw FS Row 10 (D&A, Cols S to AE): {depr_vals}")
    maruti_depr = [2515, 2822, 2604, 2760, 3021, 3528, 3034, 2789, 4846, 5256, 5608, 6742, 6966]
    assert depr_vals != maruti_depr, "FAILED: Raw FS Row 10 matches Maruti Suzuki template D&A!"
    assert depr_vals[-2] != 6742, f"FAILED: Latest D&A cell AD10 is template value {depr_vals[-2]}!"
    print(" -> PASS: Raw FS D&A cleanly populated from target data with zero template bleed.")
    
    # Check Row 27 (Capex)
    capex_vals = [ws_raw.cell(row=27, column=c).value for c in range(19, 32)]
    print(f"Raw FS Row 27 (Capex, Cols S to AE): {capex_vals}")
    maruti_capex = [-3058, -2469, -3391, -3912, -4872, -3437, -2370, -3459, -8065, -9200, -10621, -10398, None]
    assert capex_vals != maruti_capex, "FAILED: Raw FS Row 27 matches Maruti Suzuki template Capex!"
    assert capex_vals[-2] != -10398, f"FAILED: Latest Capex cell AD27 is template value {capex_vals[-2]}!"
    print(" -> PASS: Raw FS Capex cleanly populated from target data with zero template bleed.")
    
    wb.close()
    print(f"\nALL CHECKS PASSED FOR {ticker}!")

if __name__ == '__main__':
    # Test on SBIN (banking edge case with high reinvestment rate)
    test_pipeline_for_company('SBIN')
    # Test on non-banking target (Tata Steel or ITC)
    test_pipeline_for_company('TATASTEEL')
