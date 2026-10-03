import os
import sys
import zipfile
import re
import xml.etree.ElementTree as ET
import openpyxl

xlsx_path = os.path.join("exports", "RBLBANK_Valuation_Model.xlsx")
print("Inspecting:", xlsx_path)
print("File size:", os.path.getsize(xlsx_path))

# 1. Inspect zip entries
print("\n--- 1. ZIP ENTRIES & DRAWING RELATIONSHIPS ---")
with zipfile.ZipFile(xlsx_path, 'r') as z:
    names = z.namelist()
    drawing_files = [n for n in names if 'drawing' in n]
    media_files = [n for n in names if 'media' in n]
    rel_files = [n for n in names if '_rels' in n]
    print(f"Total entries: {len(names)}")
    print(f"Drawing files: {drawing_files}")
    print(f"Media files: {media_files}")
    
    # Check drawings xml syntax
    for df in drawing_files:
        if df.endswith('.xml'):
            content = z.read(df)
            try:
                ET.fromstring(content)
                print(f"Drawing XML {df}: VALID XML")
            except Exception as e:
                print(f"Drawing XML {df}: INVALID XML -> {e}")

    # Check sheet rels for drawings and media
    for rf in rel_files:
        if 'worksheets' in rf:
            content = z.read(rf)
            try:
                ET.fromstring(content)
            except Exception as e:
                print(f"Rel file {rf}: INVALID XML -> {e}")

# 2. Inspect Formulas in Comp_Valuation and other sheets with openpyxl
print("\n--- 2. OPENPYXL SHEET & FORMULA INSPECTION ---")
wb = openpyxl.load_workbook(xlsx_path, data_only=False)
ws_comp = wb['Comp_Valuation']

print("\nComp_Valuation Headers & Multiples (Rows 10-21):")
for r in range(10, 22):
    row_vals = [f"{col}{r}={ws_comp[f'{col}{r}'].value}" for col in ['B', 'C', 'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q'] if ws_comp[f'{col}{r}'].value is not None]
    print(f"Row {r}: {', '.join(row_vals)}")

print("\nComp_Valuation Stats (Rows 23-28):")
for r in range(23, 29):
    row_vals = [f"{col}{r}={ws_comp[f'{col}{r}'].value}" for col in ['B', 'I', 'J', 'O', 'P', 'Q'] if ws_comp[f'{col}{r}'].value is not None]
    print(f"Row {r}: {', '.join(row_vals)}")

print("\nComp_Valuation Target (Rows 30-39):")
for r in range(30, 40):
    row_vals = [f"{col}{r}={ws_comp[f'{col}{r}'].value}" for col in ['B', 'O', 'P', 'Q'] if ws_comp[f'{col}{r}'].value is not None]
    print(f"Row {r}: {', '.join(row_vals)}")

# 3. Check for invalid formula syntax across all sheets
print("\n--- 3. SCANNING ALL FORMULAS ACROSS ALL SHEETS ---")
invalid_formulas = []
for s_name in wb.sheetnames:
    ws = wb[s_name]
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                val = cell.value
                # Check for common corruption: double equals, dangling quotes, unmatched parens
                if val.startswith('=='):
                    invalid_formulas.append((s_name, cell.coordinate, val, "Double equals"))
                if val.count('(') != val.count(')'):
                    invalid_formulas.append((s_name, cell.coordinate, val, "Unmatched parentheses"))
                if val.count('"') % 2 != 0:
                    invalid_formulas.append((s_name, cell.coordinate, val, "Unmatched quotes"))
                if '#REF!' in val:
                    invalid_formulas.append((s_name, cell.coordinate, val, "#REF! in formula"))

if invalid_formulas:
    print(f"FOUND {len(invalid_formulas)} POTENTIALLY INVALID FORMULAS:")
    for item in invalid_formulas[:20]:
        print(f"  [{item[0]} {item[1]}] {item[3]}: {item[2]}")
else:
    print("Zero invalid formulas detected with regex/syntax check.")

wb.close()
