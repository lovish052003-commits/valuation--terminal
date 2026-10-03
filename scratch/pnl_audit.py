import openpyxl

target_path = r'C:\Users\LENOVO\Downloads\test\SUNPHARMA_Valuation_Model (1).xlsx'
wb_d = openpyxl.load_workbook(target_path, data_only=True)
wb_f = openpyxl.load_workbook(target_path, data_only=False)

out_file = r'C:\Users\LENOVO\Downloads\Advance Financial Project\scratch\pnl_audit.txt'

with open(out_file, 'w', encoding='utf-8') as f:
    if 'Profit & Loss' in wb_d.sheetnames:
        ws = wb_d['Profit & Loss']
        ws_f = wb_f['Profit & Loss']
        f.write("=== PROFIT & LOSS SHEET ===\n")
        for r in range(1, 40):
            row_vals = [ws.cell(row=r, column=c).value for c in range(1, 15)]
            if any(row_vals):
                f.write(f"Row {r:2d}: {row_vals}\n")

    if 'Raw FS' in wb_d.sheetnames:
        ws = wb_d['Raw FS']
        f.write("\n=== RAW FS P&L SECTION ===\n")
        for r in range(1, 80):
            b_v = ws.cell(row=r, column=2).value
            if b_v and any(k in str(b_v).lower() for k in ['tax', 'pbt', 'profit', 'sales', 'expenses', 'ebit', 'operating']):
                row_vals = [ws.cell(row=r, column=c).value for c in range(2, 16)]
                f.write(f"Raw FS Row {r:2d}: {row_vals}\n")

print("P&L audit written to scratch/pnl_audit.txt")
