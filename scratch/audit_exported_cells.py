import openpyxl

wb = openpyxl.load_workbook('exports/NESTLEIND_Valuation_Model.xlsx', data_only=True)
ws = wb['Data Sheet']

rows_to_check = {
    16: 'P&L Dates',
    17: 'Sales',
    28: 'PBT',
    30: 'PAT',
    56: 'BS Dates',
    57: 'Equity Cap',
    58: 'Reserves',
    59: 'Borrowings',
    60: 'Other Liab',
    61: 'Total Liab',
    62: 'Net Block',
    65: 'Other Assets',
    67: 'Receivables',
    68: 'Inventory',
    69: 'Cash & Bank',
    70: 'Shares',
    72: 'Face Value',
    81: 'CF Dates',
    82: 'CFO',
    83: 'CFI',
    84: 'CFF',
    85: 'Net Cash Flow',
    90: 'Price'
}

cols = list(range(2, 12))
header = f"{'Row':<5} | {'Label':<15} | " + " | ".join([f"Col {openpyxl.utils.get_column_letter(c)}" for c in cols])
print(header)
print("-" * len(header))
for r, label in rows_to_check.items():
    vals = [ws.cell(r, c).value for c in cols]
    formatted = []
    for v in vals:
        if v is None:
            formatted.append("    None")
        elif hasattr(v, 'strftime'):
            formatted.append(f"{v.strftime('%b-%y'):>8}")
        elif isinstance(v, (int, float)):
            formatted.append(f"{float(v):8.1f}")
        else:
            formatted.append(f"{str(v)[:8]:>8}")
    print(f"{r:<5} | {label:<15} | " + " | ".join(formatted))
