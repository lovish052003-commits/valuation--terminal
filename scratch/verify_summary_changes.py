import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

test_files = [
    'master_model_template.xlsx',
    'Tata Steel Final Model.xlsx',
    'JSWSTEEL_Valuation_Model_Fixed.xlsx'
]

expected = {
    'A5': '=DCF!D45',
    'C38': '=DCF!D45',
    'B5': '=DCF!D43',
    'C37': '=DCF!D43',
    'C46': '=DCF!D43',
    'C18': '=DCF!D41',
    'C36': '=DCF!D41',
    'C35': '=DCF!D40',
    'B35': 'Equity Value (after minority interest)',
    'C5': '=IFERROR((B5-A5)/A5,"n/a")',
    'D5': '=IFERROR((B5-A5)/B5,"n/a")',
    'C39': '=IFERROR((C37-C38)/C38,"n/a")',
    'C40': '=IFERROR((C37-C38)/C37,"n/a")',
}

all_passed = True

for fname in test_files:
    print(f"\n================ VERIFYING {fname} ================")
    wb = openpyxl.load_workbook(fname, data_only=False)
    if 'AI Valuation Summary' not in wb.sheetnames:
        print(f"Skipping {fname}: No AI Valuation Summary")
        continue
    ws = wb['AI Valuation Summary']
    
    file_passed = True
    for cell_ref, exp_val in expected.items():
        actual = str(ws[cell_ref].value or '')
        match = (actual.strip().upper() == exp_val.strip().upper())
        status = "PASS" if match else "FAIL"
        if not match:
            file_passed = False
            all_passed = False
        print(f"  [{status}] {cell_ref:4s}: Expected '{exp_val}' | Got '{actual}'")
        
    if 'DCF' in wb.sheetnames:
        ws_dcf = wb['DCF']
        print(f"  DCF Key Rows:")
        print(f"    DCF 39: B='{ws_dcf['B39'].value}' | D='{ws_dcf['D39'].value}'")
        print(f"    DCF 40: B='{ws_dcf['B40'].value}' | D='{ws_dcf['D40'].value}'")
        print(f"    DCF 41: B='{ws_dcf['B41'].value}' | D='{ws_dcf['D41'].value}'")
        print(f"    DCF 43: B='{ws_dcf['B43'].value}' | D='{ws_dcf['D43'].value}'")
        print(f"    DCF 45: B='{ws_dcf['B45'].value}' | D='{ws_dcf['D45'].value}'")
        print(f"    DCF 46: B='{ws_dcf['B46'].value}' | D='{ws_dcf['D46'].value}'")

if all_passed:
    print("\n>>> ALL UNIVERSAL VERIFICATION CHECKS PASSED PERFECTLY! <<<")
else:
    print("\n>>> SOME CHECKS FAILED! <<<")
