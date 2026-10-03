import openpyxl

wb = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=True)
print("Sheet names:", wb.sheetnames)

for s in wb.sheetnames:
    ws = wb[s]
    print(f"\n=================== SHEET: {s} ===================")
    for r in range(1, min(25, ws.max_row + 1)):
        vals = [ws.cell(r, c).value for c in range(1, min(10, ws.max_column + 1))]
        if any(v is not None for v in vals):
            print(f"Row {r:2d}: {[str(v)[:15] if v is not None else '' for v in vals]}")
