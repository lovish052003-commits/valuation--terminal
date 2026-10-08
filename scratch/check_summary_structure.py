import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

files = [
    'master_model_template.xlsx',
    'Tata Steel Final Model.xlsx',
    'JSWSTEEL_Valuation_Model_Fixed.xlsx',
    'SUNPHARMA_Valuation_Model_Corrected.xlsx',
    'HDFCBANK_Valuation_Model.xlsx',
    'SBIN_Valuation_Model.xlsx'
]

for fname in files:
    try:
        wb = openpyxl.load_workbook(fname, data_only=False)
        if 'AI Valuation Summary' in wb.sheetnames:
            ws = wb['AI Valuation Summary']
            print(f"=== {fname} ===")
            for r in [4, 5, 18, 34, 35, 36, 37, 38, 39, 40, 46]:
                row_str = f"Row {r:2d}: "
                for col in ['A', 'B', 'C', 'D', 'E']:
                    val = ws[f"{col}{r}"].value
                    if val is not None:
                        row_str += f"{col}{r}={repr(val)} | "
                print(row_str)
            if 'DCF' in wb.sheetnames:
                ws_dcf = wb['DCF']
                print("  DCF Rows 38-46:")
                for dr in range(38, 47):
                    print(f"    DCF {dr}: B={repr(ws_dcf[f'B{dr}'].value)} | D={repr(ws_dcf[f'D{dr}'].value)}")
    except Exception as e:
        print(f"{fname}: {e}")
