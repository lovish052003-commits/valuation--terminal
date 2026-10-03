import openpyxl

wb = openpyxl.load_workbook('Nestle India Model.xlsx', data_only=False)

def inspect_sheet(sheet_name, max_row=50, max_col=15):
    if sheet_name not in wb.sheetnames:
        print(f"Sheet {sheet_name} not found!")
        return
    ws = wb[sheet_name]
    print(f"\n================ SHEET: {sheet_name} ================")
    for r in range(1, min(ws.max_row + 1, max_row + 1)):
        row_vals = []
        has_content = False
        for c in range(1, min(ws.max_column + 1, max_col + 1)):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip() != "":
                has_content = True
            row_vals.append(str(v) if v is not None else "")
        if has_content:
            # print up to last non-empty col
            last_idx = max([i for i, val in enumerate(row_vals) if val != ""] + [0])
            print(f"Row {r:2d}: " + " | ".join(row_vals[:last_idx+1]))

print("Inspecting key valuation sheets:")
inspect_sheet('DCF', max_row=50, max_col=15)
inspect_sheet('WACC', max_row=40, max_col=10)
inspect_sheet('Forecasting', max_row=45, max_col=15)
inspect_sheet('Data Sheet', max_row=40, max_col=15)
