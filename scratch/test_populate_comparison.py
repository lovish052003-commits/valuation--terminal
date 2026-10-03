import sys, os
sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data, clean_num
import openpyxl

wb_ref = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=True)
ws_ref = wb_ref['Data Sheet']

wb_test = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws_test = wb_test['Data Sheet']

# Let's inspect what happens when we populate all rows
print("Reference vs Test before population:")
print("R93 Ref:", [ws_ref.cell(93, c).value for c in range(2, 12)])
print("R93 Old:", [ws_test.cell(93, c).value for c in range(2, 12)])
print("R70 Ref:", [ws_ref.cell(70, c).value for c in range(2, 12)])
print("R70 Old:", [ws_test.cell(70, c).value for c in range(2, 12)])
print("R67 Ref:", [ws_ref.cell(67, c).value for c in range(2, 12)])
print("R67 Old:", [ws_test.cell(67, c).value for c in range(2, 12)])
print("R90 Ref:", [ws_ref.cell(90, c).value for c in range(2, 12)])
print("R90 Old:", [ws_test.cell(90, c).value for c in range(2, 12)])
