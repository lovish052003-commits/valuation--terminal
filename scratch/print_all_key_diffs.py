import openpyxl

wb_f = openpyxl.load_workbook("FORCEMOT_Valuation_Model_FIXED.xlsx", data_only=False)
wb_o = openpyxl.load_workbook(r"exports\FORCEMOT_Valuation_Model.xlsx", data_only=False)

sheets = ['Control', 'Checks', 'AI Valuation Summary', 'DCF', 'WACC', 'Beta-Regression', 'Ratio Analysis']

for s in sheets:
    print(f"\n=================== ALL DIFFS IN {s} ===================")
    ws_f = wb_f[s]
    ws_o = wb_o[s]
    diffs = []
    max_r = max(ws_f.max_row, ws_o.max_row)
    max_c = max(ws_f.max_column, ws_o.max_column)
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            vf = ws_f.cell(r, c).value
            vo = ws_o.cell(r, c).value
            if vf != vo:
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                diffs.append((coord, vo, vf))
    for coord, vo, vf in diffs:
        print(f"{coord}: ORIG={vo!r} -> FIXED={vf!r}")
