import os
import openpyxl

def main():
    p = 'exports/ITC_Valuation_Model.xlsx'
    if not os.path.exists(p):
        p = 'ITC Model.xlsx'
    
    wb = openpyxl.load_workbook(p, data_only=False)
    ws = wb['DCF']
    cols = ['H', 'I', 'J', 'K', 'L', 'M']
    print(f'File: {p}')
    for r in range(6, 36):
        lbl = ws.cell(row=r, column=2).value or ws.cell(row=r, column=1).value or ''
        f_vals = [f"{c}{r}: {ws[f'{c}{r}'].value}" for c in cols if ws[f'{c}{r}'].value is not None]
        if f_vals:
            print(f'R{r} ({lbl}): ' + ' | '.join(f_vals[:3]))

if __name__ == '__main__':
    main()
