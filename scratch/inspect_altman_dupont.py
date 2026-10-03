import openpyxl

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)

print("=== ALTMAN'S Z SCORE ===")
ws_alt = wb["Altman's Z Score"]
for r in range(65, 92):
    lbl = ws_alt.cell(r, 2).value
    f_h = ws_alt.cell(r, 8).value
    f_i = ws_alt.cell(r, 9).value
    if lbl or f_h or f_i:
        print(f"Row {r}: B='{lbl}' | H='{f_h}' | I='{f_i}'")

print("\n=== DUPONT ANALYSIS ===")
ws_dup = wb["Dupont Analysis"]
for r in range(65, 82):
    lbl = ws_dup.cell(r, 2).value
    f_h = ws_dup.cell(r, 8).value
    f_i = ws_dup.cell(r, 9).value
    if lbl or f_h or f_i:
        print(f"Row {r}: B='{lbl}' | H='{f_h}' | I='{f_i}'")
