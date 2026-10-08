import openpyxl

fixed_path = r"FORCEMOT_Valuation_Model_FIXED.xlsx"
orig_path = r"exports\FORCEMOT_Valuation_Model.xlsx"

wb_f = openpyxl.load_workbook(fixed_path, data_only=False)
wb_o = openpyxl.load_workbook(orig_path, data_only=False)

sheets_to_check = ['Control', 'Checks', 'AI Valuation Summary', 'DCF', 'WACC', 'Beta-Regression', 'Ratio Analysis', 'Forecasting', 'Raw Data']

for s in sheets_to_check:
    print(f"\n{'='*30} SHEET: {s} {'='*30}")
    if s not in wb_f.sheetnames or s not in wb_o.sheetnames:
        print(f"Sheet {s} missing in one of them!")
        continue
    ws_f = wb_f[s]
    ws_o = wb_o[s]
    
    diff_cells = []
    max_r = max(ws_f.max_row, ws_o.max_row)
    max_c = max(ws_f.max_column, ws_o.max_column)
    
    for r in range(1, max_r + 1):
        for c in range(1, max_c + 1):
            vf = ws_f.cell(r, c).value
            vo = ws_o.cell(r, c).value
            if vf != vo:
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                diff_cells.append((coord, vo, vf))
                
    print(f"Total differences in {s}: {len(diff_cells)}")
    for coord, vo, vf in diff_cells[:40]:
        print(f"  {coord}:")
        print(f"     ORIG : {vo}")
        print(f"     FIXED: {vf}")
