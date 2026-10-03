import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
import openpyxl
from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model, get_effective_peers

print("=== 1. Fetching ONGC Data ===")
sd = fetch_company_data("ONGC")
c_name = sd.get("company_name", "")
ticker = sd.get("ticker", "")
print(f"Company: {c_name} ({ticker})")
print(f"Sector Peers Count: {len(sd.get('sector_peers', []))}")

print("\n=== 2. Running Universal Valuation ===")
val = calculate_valuation(sd)
print(f"Target Ticker: {val.get('ticker')}")
print(f"DCF Intrinsic Value: {val.get('intrinsic_value_per_share')}")
print(f"Target Levered Beta: {val.get('levered_beta')}")
print(f"WACC: {val.get('wacc')}")

peers = get_effective_peers(sd, val)
print(f"Effective Peers Count: {len(peers)}")
for p in peers[:8]:
    print(f"  - {p.get('name')} ({p.get('ticker')}): MCap={p.get('mcap')}, PE={p.get('pe')}, Beta={p.get('beta')}")

# Verify target is NOT in peers
for p in peers:
    assert "ONGC" not in p.get('ticker', '').upper(), f"Target found in peers: {p}"
    assert "OIL & NATURAL GAS" not in p.get('name', '').upper(), f"Target found in peers: {p}"
print("VERIFIED: Target ONGC strictly excluded from peers!")

print("\n=== 3. Exporting Excel Model ===")
excel_path = export_valuation_model(sd, val)
print(f"Excel Export Path: {excel_path}")
assert os.path.exists(excel_path), "Export file does not exist!"

print("\n=== 4. Inspecting WACC Sheet in Exported Workbook ===")
wb = openpyxl.load_workbook(excel_path, data_only=False)
ws_wacc = wb['WACC']

print("WACC Rows 14-18 (Peer Comps):")
for r in range(14, 19):
    b = ws_wacc.cell(r, 2).value
    j = ws_wacc.cell(r, 10).value
    k = ws_wacc.cell(r, 11).value
    print(f"  Row {r}: Name={b} | Levered Beta={j} | Unlevered Beta={k}")
    # Verify target not in row 14-18
    assert j != "='Beta-Regression'!L15", f"Row {r} still has legacy target beta formula: {j}"

print("\nWACC Capital Structure & Calculation Formulas:")
print(f"  C34 (Target Debt): {ws_wacc['C34'].value}")
print(f"  C35 (Target MCap): {ws_wacc['C35'].value}")
print(f"  E27 (Target Tax Rate): {ws_wacc['E27'].value}")
print(f"  K33 (Comps Median Unlev Beta): {ws_wacc['K33'].value}")
print(f"  K34 (Target D/E): {ws_wacc['K34'].value}")
print(f"  K35 (Tax Rate): {ws_wacc['K35'].value}")
print(f"  K36 (Target Re-levered Beta): {ws_wacc['K36'].value}")
print(f"  K28 (Levered Beta): {ws_wacc['K28'].value}")
print(f"  K29 (Cost of Equity Ke): {ws_wacc['K29'].value}")
print(f"  K40 (Cost of Equity): {ws_wacc['K40'].value}")
print(f"  K41 (Equity Weight We): {ws_wacc['K41'].value}")
print(f"  K43 (Post-Tax Kd): {ws_wacc['K43'].value}")
print(f"  K44 (Debt Weight Wd): {ws_wacc['K44'].value}")
print(f"  K46 (WACC): {ws_wacc['K46'].value}")

assert ws_wacc['K34'].value == "=D38", f"K34 should be =D38, got {ws_wacc['K34'].value}"
assert ws_wacc['K41'].value == "=D35", f"K41 should be =D35, got {ws_wacc['K41'].value}"
assert ws_wacc['K44'].value == "=D34", f"K44 should be =D34, got {ws_wacc['K44'].value}"
assert ws_wacc['K43'].value == "=E28", f"K43 should be =E28, got {ws_wacc['K43'].value}"
assert ws_wacc['K46'].value == "=(K40*K41)+(K43*K44)", f"K46 should be =(K40*K41)+(K43*K44), got {ws_wacc['K46'].value}"

print("\n=== 5. Inspecting Comp_Valuation Sheet ===")
ws_comp = wb['Comp_Valuation']
for r in range(12, 17):
    b = ws_comp.cell(r, 2).value
    c = ws_comp.cell(r, 3).value
    o = ws_comp.cell(r, 15).value
    p = ws_comp.cell(r, 16).value
    q = ws_comp.cell(r, 17).value
    print(f"  Row {r}: Name={b} ({c}) | EV/Rev={o} | EV/EBITDA={p} | P/E={q}")
    assert c != "ONGC", f"Target ONGC in Comp_Valuation row {r}!"

print("\n=== 6. Inspecting AI Valuation Summary ===")
ws_ai = wb['AI Valuation Summary']
for r in range(8, 21):
    b = ws_ai.cell(r, 2).value
    c = ws_ai.cell(r, 3).value
    print(f"  Row {r}: {b} = {c}")

print("\nALL VERIFICATIONS PASSED SUCCESSFULLY!")
