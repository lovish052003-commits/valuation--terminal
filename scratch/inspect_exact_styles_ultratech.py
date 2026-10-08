import openpyxl

wb = openpyxl.load_workbook("UltraTech_fixed_v2.xlsx", data_only=False)

def print_row_details(ws, rows, cols):
    print(f"\n--- {ws.title} Details ---")
    for r in rows:
        row_str = []
        for c in cols:
            cell = ws.cell(r, c)
            val = cell.value
            if val is not None:
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                row_str.append(f"{coord}: val={val!r}, num_fmt={cell.number_format!r}, font={cell.font.name} {cell.font.size} bold={cell.font.bold}")
        if row_str:
            print(" | ".join(row_str).encode('ascii', errors='replace').decode('ascii'))

print_row_details(wb['AI Valuation Summary'], [4, 5, 10, 12, 40, 41, 42], range(1, 11))
print_row_details(wb['Control'], range(31, 38), range(1, 6))
print_row_details(wb['Checks'], [13] + list(range(17, 23)), range(1, 7))
print_row_details(wb['Intrinsic Valuation'], [55, 65] + list(range(67, 72)), range(1, 14))
print_row_details(wb['DCF'], [58, 60, 61, 64, 65, 66, 67, 74, 75, 76, 77, 78, 79, 80, 81, 82, 83], range(1, 15))
