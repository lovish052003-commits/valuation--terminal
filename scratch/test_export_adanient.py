import sys
sys.path.insert(0, '.')
import screener_client
import valuation_engine
import excel_exporter

print("Fetching ADANIENT...")
sd = screener_client.fetch_company_data("ADANIENT")
print("Calculating valuation...")
val = valuation_engine.calculate_valuation(sd)
print("Exporting valuation model...")
path = excel_exporter.export_valuation_model(sd, val)
print("Exported to:", path)

import openpyxl
wb = openpyxl.load_workbook(path, data_only=True)
print("\nVerifying exported file:", path)
print("Data Sheet B1:", wb['Data Sheet']['B1'].value)
print("Data Sheet K17:", wb['Data Sheet']['K17'].value)
print("Data Sheet K30:", wb['Data Sheet']['K30'].value)
print("DCF D42 (IV):", wb['DCF']['D42'].value)
print("DCF D44 (CMP):", wb['DCF']['D44'].value)
print("AI Valuation Summary exists:", 'AI Valuation Summary' in wb.sheetnames)
if 'AI Valuation Summary' in wb.sheetnames:
    print("AI Summary A1:", wb['AI Valuation Summary']['A1'].value)
    print("AI Summary B5:", wb['AI Valuation Summary']['B5'].value)
