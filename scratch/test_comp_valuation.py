import os, sys
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

def test_comp_val():
    print("Testing Comp_Valuation Sheet for HINDUNILVR...")
    data = fetch_company_data('HINDUNILVR')
    val = calculate_valuation(data)
    out_file = export_valuation_model(data, val, "### HUL Valuation Report")
    print(f"Exported to: {out_file}")

    wb = openpyxl.load_workbook(out_file, data_only=False)
    ws = wb['Comp_Valuation']
    print(f"\n--- Comp_Valuation Formulas ---")
    print(f"B30 (Title): {ws['B30'].value}")
    print(f"Row 26 (Peer Avg): O26={ws['O26'].value} | P26={ws['P26'].value} | Q26={ws['Q26'].value}")
    print(f"Row 32 (Implied EV): O32={ws['O32'].value} | P32={ws['P32'].value} | Q32={ws['Q32'].value}")
    print(f"Row 33 (Net Debt): O33={ws['O33'].value} | P33={ws['P33'].value} | Q33={ws['Q33'].value}")
    print(f"Row 35 (Shares): O35={ws['O35'].value} | P35={ws['P35'].value} | Q35={ws['Q35'].value}")
    print(f"Row 37 (Implied Value/Sh): O37={ws['O37'].value} | P37={ws['P37'].value} | Q37={ws['Q37'].value}")
    print(f"Row 39 (Verdict): O39={ws['O39'].value} | P39={ws['P39'].value} | Q39={ws['Q39'].value}")
    wb.close()

    wb_eval = openpyxl.load_workbook(out_file, data_only=True)
    ws_eval = wb_eval['Comp_Valuation']
    print(f"\n--- Comp_Valuation Evaluated Values ---")
    print(f"Target Company (Row 12): Name={ws_eval['B12'].value} | CMP=Rs. {ws_eval['D12'].value} | Shares={ws_eval['E12'].value} Cr")
    print(f"Peer Avg (Row 26): EV/Rev={ws_eval['O26'].value} | EV/EBITDA={ws_eval['P26'].value} | P/E={ws_eval['Q26'].value}")
    print(f"Implied EV (Row 32): EV/Rev=Rs. {ws_eval['O32'].value} Cr | EV/EBITDA=Rs. {ws_eval['P32'].value} Cr | P/E=Rs. {ws_eval['Q32'].value} Cr")
    print(f"Net Debt (Row 33): Rs. {ws_eval['O33'].value} Cr")
    print(f"Implied Market Value (Row 34): EV/Rev=Rs. {ws_eval['O34'].value} Cr | EV/EBITDA=Rs. {ws_eval['P34'].value} Cr | P/E=Rs. {ws_eval['Q34'].value} Cr")
    print(f"Shares Outstanding (Row 35): {ws_eval['O35'].value} Cr")
    print(f"Implied Value per Share (Row 37): EV/Rev=Rs. {ws_eval['O37'].value} | EV/EBITDA=Rs. {ws_eval['P37'].value} | P/E=Rs. {ws_eval['Q37'].value}")
    print(f"Verdict (Row 39): EV/Rev={ws_eval['O39'].value} | EV/EBITDA={ws_eval['P39'].value} | P/E={ws_eval['Q39'].value}")
    wb_eval.close()

if __name__ == '__main__':
    test_comp_val()
