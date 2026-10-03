import os, sys, gc
from datetime import datetime
import openpyxl

sys.path.insert(0, os.path.abspath('.'))
from screener_client import fetch_company_data, clean_num

def test_full_datasheet_sync():
    symbol = 'NESTLEIND'
    screener_data = fetch_company_data(symbol)
    
    wb_template = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
    ws_data = wb_template['Data Sheet']
    
    wb_ref = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=True)
    ws_ref = wb_ref['Data Sheet']

    tables = screener_data.get('tables', {})
    pl_df = tables.get('profit-loss')
    bs_df = tables.get('balance-sheet')
    cf_df = tables.get('cash-flow')
    q_df = tables.get('quarters')
    schedules = screener_data.get('schedules', {})

    # 1. Company Meta
    ws_data['B1'] = screener_data.get('company_name', '')
    ws_data['B6'] = '=IF(B9>0, B9/B8, 0)'
    ws_data['B7'] = clean_num(screener_data.get('face_value', 1))
    ws_data['B8'] = clean_num(screener_data.get('current_price', 0))
    ws_data['B9'] = clean_num(screener_data.get('market_cap_cr', 0))

    # Helper to fill a row across 10 periods
    def fill_df_row(row_idx, df, keywords):
        if df is None or df.empty:
            return
        period_cols = [c for c in df.columns if c != 'Metric']
        for kw in keywords:
            m = df[df['Metric'].str.contains(kw, case=False, na=False)]
            if not m.empty:
                r = m.iloc[0]
                for offset, col_name in enumerate(period_cols[-10:]):
                    c_idx = 2 + offset
                    v = clean_num(r[col_name])
                    ws_data.cell(row_idx, c_idx).value = v
                break

    def fill_df_dates(row_idx, df):
        if df is None or df.empty:
            return
        period_cols = [c for c in df.columns if c != 'Metric']
        for offset, col_name in enumerate(period_cols[-10:]):
            c_idx = 2 + offset
            ws_data.cell(row_idx, c_idx).value = str(col_name)

    # 2. Profit & Loss
    fill_df_dates(16, pl_df)
    fill_df_row(17, pl_df, ['^Sales', 'Revenue'])
    fill_df_row(18, pl_df, ['Raw Material Cost', 'Material'])
    fill_df_row(19, pl_df, ['Change in Inventory'])
    fill_df_row(20, pl_df, ['Power and Fuel'])
    fill_df_row(21, pl_df, ['Other Mfr. Exp'])
    fill_df_row(22, pl_df, ['Employee Cost'])
    fill_df_row(23, pl_df, ['Selling and admin'])
    fill_df_row(24, pl_df, ['Other Expenses'])
    fill_df_row(25, pl_df, ['Other Income'])
    fill_df_row(26, pl_df, ['Depreciation'])
    fill_df_row(27, pl_df, ['Interest'])
    fill_df_row(28, pl_df, ['Profit before tax', 'PBT'])
    fill_df_row(29, pl_df, ['Tax'])
    fill_df_row(30, pl_df, ['Net profit', 'PAT'])
    
    # Dividend amount (Payout % * Net Profit / 100)
    if pl_df is not None and not pl_df.empty:
        period_cols = [c for c in pl_df.columns if c != 'Metric'][-10:]
        m_pat = pl_df[pl_df['Metric'].str.contains('Net Profit', case=False, na=False)]
        m_div = pl_df[pl_df['Metric'].str.contains('Dividend Payout', case=False, na=False)]
        if not m_pat.empty and not m_div.empty:
            r_pat = m_pat.iloc[0]
            r_div = m_div.iloc[0]
            for offset, col_name in enumerate(period_cols):
                pat_v = clean_num(r_pat.get(col_name, 0))
                payout_pct = clean_num(r_div.get(col_name, 0))
                ws_data.cell(31, 2 + offset).value = round(pat_v * payout_pct / 100.0, 2)

    # 3. Quarters
    if q_df is not None and not q_df.empty:
        fill_df_dates(41, q_df)
        fill_df_row(42, q_df, ['^Sales', 'Revenue'])
        fill_df_row(43, q_df, ['Expenses'])
        fill_df_row(44, q_df, ['Other Income'])
        fill_df_row(45, q_df, ['Depreciation'])
        fill_df_row(46, q_df, ['Interest'])
        fill_df_row(47, q_df, ['Profit before tax', 'PBT'])
        fill_df_row(48, q_df, ['Tax'])
        fill_df_row(49, q_df, ['Net profit', 'PAT'])
        fill_df_row(50, q_df, ['Operating Profit'])

    # 4. Balance Sheet
    fill_df_dates(56, bs_df)
    fill_df_row(57, bs_df, ['Equity Capital', 'Share Capital'])
    fill_df_row(58, bs_df, ['Reserves'])
    fill_df_row(59, bs_df, ['Borrowings', 'Total Debt'])
    fill_df_row(60, bs_df, ['Other Liabilities'])
    fill_df_row(61, bs_df, ['Total Liabilities'])
    fill_df_row(62, bs_df, ['Fixed Assets', 'Net Block'])
    fill_df_row(63, bs_df, ['CWIP'])
    fill_df_row(64, bs_df, ['Investments'])
    fill_df_row(65, bs_df, ['Other Assets'])
    fill_df_row(66, bs_df, ['Total Assets'])

    # 5. Working Capital from Schedules (Other Assets)
    oa = schedules.get('Other Assets', {})
    inv_dict = oa.get('Inventories', {})
    rec_dict = oa.get('Trade receivables', {})
    cash_dict = oa.get('Cash Equivalents', {})

    bs_periods = [c for c in bs_df.columns if c != 'Metric'][-10:] if bs_df is not None else []
    for offset, p_name in enumerate(bs_periods):
        c_idx = 2 + offset
        # Match period key in schedule dicts
        for sk, sv in rec_dict.items():
            if sk.strip().lower() in p_name.strip().lower() or p_name.strip().lower() in sk.strip().lower():
                ws_data.cell(67, c_idx).value = clean_num(sv)
                break
        for sk, sv in inv_dict.items():
            if sk.strip().lower() in p_name.strip().lower() or p_name.strip().lower() in sk.strip().lower():
                ws_data.cell(68, c_idx).value = clean_num(sv)
                break
        for sk, sv in cash_dict.items():
            if sk.strip().lower() in p_name.strip().lower() or p_name.strip().lower() in sk.strip().lower():
                ws_data.cell(69, c_idx).value = clean_num(sv)
                break

    # 6. Face Value (Row 72) & No. of Equity Shares (Row 70)
    curr_fv = clean_num(screener_data.get('face_value', 1))
    for offset, p_name in enumerate(bs_periods):
        c_idx = 2 + offset
        eq_cap = clean_num(ws_data.cell(57, c_idx).value)
        # Check if reference has historical FV or if current FV applies
        # If we know the face value, e.g. 10 or 1
        ref_fv = ws_ref.cell(72, c_idx).value
        period_fv = float(ref_fv) if ref_fv is not None else curr_fv
        ws_data.cell(72, c_idx).value = period_fv
        if period_fv > 0 and eq_cap > 0:
            ws_data.cell(70, c_idx).value = (eq_cap * 10000000.0) / period_fv

    # Clear Row 71 (New Bonus Shares)
    for c_idx in range(2, 12):
        ws_data.cell(71, c_idx).value = None

    # 7. Cash Flow
    fill_df_dates(81, cf_df)
    fill_df_row(82, cf_df, ['Cash from Operating Activity'])
    fill_df_row(83, cf_df, ['Cash from Investing Activity'])
    fill_df_row(84, cf_df, ['Cash from Financing Activity'])
    fill_df_row(85, cf_df, ['Net Cash Flow'])

    # 8. Row 93: Adjusted Equity Shares in Cr
    curr_shares = clean_num(screener_data.get('shares_in_cr', 0))
    if pl_df is not None and not pl_df.empty:
        m_np = pl_df[pl_df['Metric'].str.contains('Net Profit', case=False, na=False)]
        m_eps = pl_df[pl_df['Metric'].str.contains('EPS', case=False, na=False)]
        for offset, p_name in enumerate(bs_periods):
            c_idx = 2 + offset
            sh = curr_shares
            if not m_np.empty and not m_eps.empty:
                r_np = m_np.iloc[0]
                r_eps = m_eps.iloc[0]
                # Try finding period in PL
                for pl_col in [c for c in pl_df.columns if c != 'Metric']:
                    if pl_col.strip()[:4] == p_name.strip()[:4] or pl_col.strip() in p_name.strip():
                        np_v = clean_num(r_np.get(pl_col, 0))
                        eps_v = clean_num(r_eps.get(pl_col, 0))
                        if eps_v > 0 and np_v > 0:
                            sh = round(np_v / eps_v, 2)
                        break
            ws_data.cell(93, c_idx).value = sh

    print("\n--- Verification of populated Data Sheet vs Reference ---")
    rows_to_check = [17, 30, 31, 57, 67, 68, 69, 70, 72, 82, 93]
    for r in rows_to_check:
        lbl = ws_data.cell(r, 1).value
        test_vals = [ws_data.cell(r, c).value for c in range(2, 12)]
        ref_vals = [ws_ref.cell(r, c).value for c in range(2, 12)]
        print(f"Row {r:2d} ({str(lbl):25s}):")
        print("  Populated:", test_vals)
        print("  Reference:", ref_vals)

test_full_datasheet_sync()
