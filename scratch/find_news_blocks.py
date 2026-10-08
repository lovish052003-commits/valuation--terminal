import openpyxl

wb = openpyxl.load_workbook('master_model_template.xlsx', data_only=False)
for name in ["Altman's Z Score", 'Dupont Analysis']:
    if name in wb.sheetnames:
        ws = wb[name]
        print(f"=== {name} ===")
        print("  B5:", ws['B5'].value)
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                val = str(ws.cell(r, c).value or '')
                if any(k in val.lower() for k in ['news', 'economic times', 'wikipedia', 'recent developments', 'about', 'filing']):
                    print(f"  Row {r} Col {c}: {val[:80]}")
