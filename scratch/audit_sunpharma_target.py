import openpyxl
import os

target_path = r'C:\Users\LENOVO\Downloads\test\SUNPHARMA_Valuation_Model (1).xlsx'
wb = openpyxl.load_workbook(target_path, data_only=False)

print("=== SHEET NAMES ===")
print(wb.sheetnames)

print("\n=== 1. HISTORICAL FS AUDIT ===")
if 'Historical FS' in wb.sheetnames:
    ws = wb['Historical FS']
    # Print rows around Tax Rate and EBITDA
    for r in range(1, 80):
        lbl = ws.cell(row=r, column=2).value
        if lbl:
            lbl_str = str(lbl).strip()
            # print matching rows
            if any(k in lbl_str.lower() for k in ['tax', 'ebitda', 'ebit', 'gross profit', 'selling', 'operating profit', 'depreciation', 'pbt']):
                c_val = ws.cell(row=r, column=3).value
                d_val = ws.cell(row=r, column=4).value
                print(f"Row {r:2d} | Label: {lbl_str:<35} | Col C: {str(c_val):<25} | Col D: {str(d_val):<25}")

print("\n=== 2. DCF AUDIT ===")
if 'DCF' in wb.sheetnames:
    ws = wb['DCF']
    for r in range(1, 55):
        b_val = ws.cell(row=r, column=2).value
        c_val = ws.cell(row=r, column=3).value
        d_val = ws.cell(row=r, column=4).value
        if b_val or c_val or d_val:
            print(f"Row {r:2d} | B: {str(b_val):<30} | C: {str(c_val):<25} | D: {str(d_val):<25}")

print("\n=== 3. WACC AUDIT ===")
if 'WACC' in wb.sheetnames:
    ws = wb['WACC']
    for r in range(1, 50):
        b_val = ws.cell(row=r, column=2).value
        c_val = ws.cell(row=r, column=3).value
        d_val = ws.cell(row=r, column=4).value
        k_val = ws.cell(row=r, column=11).value
        if b_val or c_val or k_val:
            print(f"Row {r:2d} | B: {str(b_val):<25} | C: {str(c_val):<20} | K: {str(k_val):<25}")

print("\n=== 4. BETA REGRESSION AUDIT (Around row 256) ===")
if 'Beta-Regression' in wb.sheetnames:
    ws = wb['Beta-Regression']
    for r in range(250, 265):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 10)]
        if any(row_vals):
            print(f"Row {r:3d}: {row_vals}")
    # Also check regression statistics cells
    for r in range(1, 20):
        print(f"Stats Row {r:2d} | G: {ws.cell(row=r, column=7).value} | H: {ws.cell(row=r, column=8).value} | I: {ws.cell(row=r, column=9).value}")

print("\n=== 5. COMP_VALUATION AUDIT ===")
if 'Comp_Valuation' in wb.sheetnames:
    ws = wb['Comp_Valuation']
    for r in range(10, 25):
        row_vals = {openpyxl.utils.get_column_letter(c): ws.cell(row=r, column=c).value for c in range(2, 18)}
        print(f"Row {r:2d}: {row_vals}")
    for r in range(23, 40):
        b_val = ws.cell(row=r, column=2).value
        o_val = ws.cell(row=r, column=15).value
        p_val = ws.cell(row=r, column=16).value
        q_val = ws.cell(row=r, column=17).value
        print(f"Row {r:2d} | B: {str(b_val):<30} | O: {str(o_val):<20} | P: {str(p_val):<20} | Q: {str(q_val):<20}")

print("\n=== 6. AI VALUATION SUMMARY AUDIT ===")
if 'AI Valuation Summary' in wb.sheetnames:
    ws = wb['AI Valuation Summary']
    for r in range(1, 40):
        b_val = ws.cell(row=r, column=2).value
        c_val = ws.cell(row=r, column=3).value
        d_val = ws.cell(row=r, column=4).value
        if b_val or c_val:
            print(f"Row {r:2d} | B: {str(b_val):<35} | C: {str(c_val):<30} | D: {str(d_val):<15}")
