import openpyxl, sys

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
if 'AI Valuation Summary' in wb.sheetnames:
    ws = wb['AI Valuation Summary']
    for r in range(1, 45):
        b = str(ws[f'B{r}'].value or '').encode('ascii', errors='replace').decode('ascii')
        c = str(ws[f'C{r}'].value or '').encode('ascii', errors='replace').decode('ascii')
        d = str(ws[f'D{r}'].value or '').encode('ascii', errors='replace').decode('ascii')
        if b or c or d:
            print(f"Row {r:2d}: B={b:<35} C={c:<35} D={d}")
else:
    print("AI Valuation Summary not found")
