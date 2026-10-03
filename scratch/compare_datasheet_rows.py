import openpyxl

wb_itc = openpyxl.load_workbook('ITC Model.xlsx', data_only=True)
wb_nestle = openpyxl.load_workbook(r'C:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=True)

ws_itc = wb_itc['Data Sheet']
ws_nestle = wb_nestle['Data Sheet']

print(f"{'Row':4s} | {'Label':30s} | {'ITC (Old)':18s} | {'Nestle (Target)':18s} | {'Diff?'}")
print("-" * 85)

for r in range(1, 96):
    lbl = ws_itc.cell(r, 1).value or ws_nestle.cell(r, 1).value
    itc_b = ws_itc.cell(r, 2).value
    nes_b = ws_nestle.cell(r, 2).value
    
    if lbl or itc_b or nes_b:
        is_diff = str(itc_b) != str(nes_b)
        print(f"R{r:02d}  | {str(lbl):30s} | {str(itc_b)[:18]:18s} | {str(nes_b)[:18]:18s} | {is_diff}")
