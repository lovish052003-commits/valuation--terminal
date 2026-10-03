import openpyxl

wb = openpyxl.load_workbook('ITC Model.xlsx', data_only=False)
ws = wb['Raw FS']

print("--- Raw FS Rows 55-67 in ITC Model.xlsx ---")
for r in range(55, 68):
    l_val = ws.cell(r, 12).value  # Col L (Company)
    m_val = ws.cell(r, 13).value  # Col M (CMP)
    n_val = ws.cell(r, 14).value  # Col N (Shares)
    ar_val = ws.cell(r, 44).value # Col AR (Mcap)
    as_val = ws.cell(r, 45).value # Col AS (Net Debt)
    at_val = ws.cell(r, 46).value # Col AT (EV)
    au_val = ws.cell(r, 47).value # Col AU (Sales)
    av_val = ws.cell(r, 48).value # Col AV (EBITDA)
    aw_val = ws.cell(r, 49).value # Col AW (PAT)
    print(f"Row {r}: L={l_val} | M={m_val} | N={n_val} | AR={ar_val} | AU={au_val} | AV={av_val}")
