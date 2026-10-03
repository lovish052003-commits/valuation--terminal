import openpyxl

wb = openpyxl.load_workbook(r'c:\Users\LENOVO\Downloads\Nestle India (2).xlsx', data_only=False)
print("Sheet count:", len(wb.sheetnames))
for s in wb.sheetnames:
    ws = wb[s]
    print(f"Sheet '{s}': state={ws.sheet_state}, max_row={ws.max_row}, max_col={ws.max_column}")
