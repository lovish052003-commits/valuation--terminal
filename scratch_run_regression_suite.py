import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

from screener_client import fetch_company_data
from valuation_engine import calculate_valuation
from excel_exporter import export_valuation_model

def test_company(ticker, is_expected_bank=False):
    print(f"\n==================================================================")
    print(f"TESTING REGRESSION COMPANY: {ticker} (Expected Bank: {is_expected_bank})")
    print(f"==================================================================")
    
    data = fetch_company_data(ticker)
    val_res = calculate_valuation(data)
    
    c_type = val_res.get('company_type', 'Operating Company')
    print(f"1. Company Classification: '{c_type}'")
    if is_expected_bank:
        assert 'bank' in str(c_type).lower(), f"Expected Bank, got {c_type}"
    else:
        assert 'bank' not in str(c_type).lower(), f"Expected Non-Bank, got {c_type}"
    
    out_file = export_valuation_model(data, val_res)
    print(f"2. Export Completed: {out_file}")
    
    # Audit formulas and cached values
    wb_f = openpyxl.load_workbook(out_file, data_only=False)
    wb_v = openpyxl.load_workbook(out_file, data_only=True)
    
    errs = []
    for sname in wb_f.sheetnames:
        ws_f = wb_f[sname]
        ws_v = wb_v[sname]
        for r in range(1, min(ws_f.max_row+1, 100)):
            for c in range(1, min(ws_f.max_column+1, 30)):
                v_val = str(ws_v.cell(r, c).value or '')
                f_val = str(ws_f.cell(r, c).value or '')
                for tok in ['#VALUE!', '#REF!', '#DIV/0!', '#NAME?']:
                    if tok in v_val or tok in f_val:
                        errs.append((sname, openpyxl.utils.get_column_letter(c) + str(r), f_val, v_val))
    
    print(f"3. Error Scan across all sheets: {len(errs)} errors found.")
    if errs:
        for e in errs[:5]:
            print("   ERR:", e)
        return False, len(errs)
    
    # Verify Shares and Market Cap
    ds_v = wb_v['Data Sheet']
    cmp_v = ds_v['B8'].value
    shares_v = ds_v['B6'].value
    mcap_v = ds_v['B9'].value
    raw_shares_v = ds_v['K70'].value
    print(f"4. Market Data Reconciliation: CMP={cmp_v}, Shares={shares_v} Cr, MktCap={mcap_v} Cr, Raw K70/1e7={float(raw_shares_v or 0)/1e7} Cr")
    
    # Verify AI Summary
    ai_v = wb_v['AI Valuation Summary']
    print(f"5. AI Summary: IV={ai_v['B5'].value}, Gap={ai_v['D5'].value}, Discount Rate={ai_v['E5'].value}")
    print(f"   Pillar 3: {ai_v['B10'].value}")
    print(f"   Pillar 4: {ai_v['B11'].value}")
    
    return True, 0

if __name__ == '__main__':
    companies = [
        ('SBIN', True),
        ('ADANIPOWER', False),
        ('OIL', False),
        ('TCS', False),
        ('TATASTEEL', False),
        ('SUNPHARMA', False),
        ('ADANIENT', False)
    ]
    
    summary = {}
    for sym, is_b in companies:
        try:
            ok, n_err = test_company(sym, is_expected_bank=is_b)
            summary[sym] = "PASS (0 errors)" if ok else f"FAIL ({n_err} errors)"
        except Exception as ex:
            print(f"❌ Exception on {sym}: {ex}")
            summary[sym] = f"ERROR: {ex}"
            
    print("\n==================================================================")
    print("UNIVERSAL ENGINE CROSS-SECTOR REGRESSION RESULTS:")
    print("==================================================================")
    for sym, res in summary.items():
        print(f"  {sym:12} : {res}")
    print("==================================================================")
