import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

path = r'C:\Users\LENOVO\Downloads\test\HDFCBANK_Valuation_Model (1).xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
wb_f = openpyxl.load_workbook(path, data_only=False)

print("Scanning for Undervalued / Overvalued in:", path)
for sname in wb.sheetnames:
    ws = wb[sname]
    ws_f = wb_f[sname]
    for r in range(1, ws.max_row+1):
        for c in range(1, ws.max_column+1):
            val = str(ws.cell(r, c).value or '')
            fval = str(ws_f.cell(r, c).value or '')
            if 'undervalued' in val.lower() or 'overvalued' in val.lower() or 'undervalued' in fval.lower() or 'overvalued' in fval.lower():
                print(f"{sname}!{openpyxl.utils.get_column_letter(c)}{r}: val='{val}', formula='{fval}'")
